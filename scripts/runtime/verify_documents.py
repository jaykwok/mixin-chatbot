"""Offline parser smoke test. Outputs remain under backup/tmp for inspection."""
from pathlib import Path
from tempfile import mkdtemp
import json
import subprocess
import sys
import pandas as pd
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

temporary = Path("backup/tmp")
temporary.mkdir(parents=True, exist_ok=True)
root = Path(mkdtemp(prefix="document-check-", dir=temporary))
text = "量子产品资料回归"
document = Document()
document.add_paragraph(text)
document.add_table(rows=1, cols=1).cell(0, 0).text = "中间表格"
document.add_paragraph("段落末尾")
document.save(root / "sample.docx")
assert Document(root / "sample.docx").paragraphs[0].text == text
presentation = Presentation()
presentation.slides.add_slide(presentation.slide_layouts[0]).shapes.title.text = text
presentation.save(root / "sample.pptx")
assert Presentation(root / "sample.pptx").slides[0].shapes.title.text == text
book = Workbook()
book.active["A1"] = text
book.active["A3"] = "=1+2"
book.save(root / "sample.xlsx")
assert load_workbook(root / "sample.xlsx").active["A1"].value == text
pd.DataFrame({"资料": [text]}).to_excel(root / "dataframe.xlsx", engine="xlsxwriter", index=False)
assert pd.read_excel(root / "dataframe.xlsx").iloc[0, 0] == text
pdf = PdfWriter()
pdf.add_blank_page(width=72, height=72)
with (root / "sample.pdf").open("wb") as handle:
    pdf.write(handle)
assert len(PdfReader(root / "sample.pdf").pages) == 1
Image.new("RGB", (16, 16), "white").save(root / "sample.png")
with Image.open(root / "sample.png") as image:
    assert image.size == (16, 16)

# Exercise the production extractor, including ordering, source locations and bounds.
extractor = Path(__file__).with_name("extract_document.py").resolve()
def extract(extension, limit=250000):
    target = root / (extension + "-" + str(limit) + ".txt")
    run = subprocess.run([sys.executable, str(extractor), str(root / ("sample." + extension)),
                          "." + extension, str(target), str(limit)], check=True, capture_output=True, text=True, encoding="utf-8")
    metadata = json.loads(run.stdout)
    content = target.read_text(encoding="utf-8")
    assert metadata["characters"] == len(content)
    assert metadata["units"] > 0
    return content, metadata

doc_text, _ = extract("docx")
assert doc_text.index(text) < doc_text.index("中间表格") < doc_text.index("段落末尾")
slide_text, _ = extract("pptx")
assert "## 幻灯片 1" in slide_text and text in slide_text
sheet_text, _ = extract("xlsx")
assert "## Sheet:" in sheet_text and "3: =1+2" in sheet_text and text in sheet_text
pdf_text, _ = extract("pdf")
assert "## 页 1" in pdf_text and "OCR" in pdf_text
document.add_paragraph("有界提取" * 1000)
document.save(root / "sample.docx")
bounded, metadata = extract("docx", 1000)
assert len(bounded) == 1000 and metadata["truncated"] is True
result = {"formats": ["docx", "pptx", "xlsx", "pdf", "png"], "production_extractor": "passed", "status": "passed"}
(root / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({**result, "directory": str(root)}, ensure_ascii=False))
