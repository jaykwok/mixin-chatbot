"""Bounded text extraction. Outputs UTF-8 text and a small JSON result on stdout."""
import json
import sys
import zipfile

source, extension, destination, limit_text = sys.argv[1:]
limit = int(limit_text)
units = 0
truncated = False
written = 0

# Reject oversized OOXML expansions before libraries load the archive.
if extension != ".pdf":
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if len(entries) > 20000 or sum(e.file_size for e in entries) > 256 * 1024 * 1024:
            raise ValueError("文档解压大小或条目数超过解析上限")

class LimitReached(Exception):
    pass

with open(destination, "w", encoding="utf-8", newline="\n") as output:
    def emit(text):
        global written, truncated
        value = str(text).replace("\x00", "") + "\n"
        remaining = limit - written
        output.write(value[:remaining])
        written += min(len(value), remaining)
        if len(value) > remaining:
            truncated = True
            raise LimitReached()

    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(source)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF 已加密，无法读取")
            for number, page in enumerate(reader.pages):
                if number >= 1000:
                    truncated = True
                    break
                units += 1
                emit("\n## 页 " + str(number + 1))
                emit(page.extract_text() or "[此页无可提取文本，可能需要 OCR]")
        elif extension == ".docx":
            from docx import Document
            from docx.table import Table
            document = Document(source)
            for number, block in enumerate(document.iter_inner_content()):
                units += 1
                if isinstance(block, Table):
                    emit("\n## 表格（段落位置 " + str(number + 1) + "）")
                    for row in block.rows:
                        emit(" | ".join(cell.text for cell in row.cells))
                else:
                    emit(block.text)
        elif extension == ".pptx":
            from pptx import Presentation
            presentation = Presentation(source)
            for number, slide in enumerate(presentation.slides):
                if number >= 1000:
                    truncated = True
                    break
                units += 1
                emit("\n## 幻灯片 " + str(number + 1))
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        emit(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            emit(" | ".join(cell.text for cell in row.cells))
        elif extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            try:
                emit("[公式按原文显示，未执行计算；行号对应原工作表]")
                cells = 0
                for sheet_number, sheet in enumerate(workbook.worksheets):
                    if sheet_number >= 100:
                        truncated = True
                        break
                    units += 1
                    emit("\n## Sheet: " + sheet.title)
                    # Cap dimensions as well as emitted text, including blank forged dimensions.
                    for row_number, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 1, 100000), max_col=min(sheet.max_column or 1, 256)), 1):
                        cells += len(row)
                        if cells > 100000:
                            truncated = True
                            raise LimitReached()
                        if any(cell.value is not None for cell in row):
                            emit(str(row_number) + ": " + " | ".join("" if cell.value is None else str(cell.value) for cell in row))
                    if (sheet.max_column or 0) > 256 or (sheet.max_row or 0) > 100000:
                        truncated = True
            finally:
                workbook.close()
        else:
            raise ValueError("不支持此文档格式")
    except LimitReached:
        truncated = True

print(json.dumps({"units": units, "truncated": truncated, "characters": written}, ensure_ascii=False))
