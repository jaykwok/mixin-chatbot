"""Render the Markdown block model onto a template's styles. Imported by document_ops.py."""
import math
import os
import re
import unicodedata
from pathlib import Path

from lxml import etree

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def runs_text(runs):
    return "".join(r.get("text", "") for r in runs)


def char_units(text, size_pt):
    """Approximate rendered width in points: CJK ≈ 1em, Latin ≈ 0.55em."""
    width = 0.0
    for char in text:
        if unicodedata.east_asian_width(char) in ("W", "F"):
            width += size_pt
        elif char == " ":
            width += size_pt * 0.3
        else:
            width += size_pt * 0.55
    return width


def parse_width(value, available_emu):
    if not value:
        return None
    match = re.search(r"(\d+(?:\.\d+)?)\s*(cm|in|%|pt|mm)", str(value))
    if not match:
        return None
    number, unit = float(match.group(1)), match.group(2)
    emu = {"cm": 360000, "mm": 36000, "in": EMU_PER_INCH, "pt": EMU_PER_PT}.get(unit)
    result = int(number * emu) if emu else int(available_emu * min(number, 100) / 100)
    return max(EMU_PER_INCH // 2, min(result, available_emu))


def image_size(path):
    from PIL import Image
    with Image.open(path) as picture:
        return picture.size


FLOWCHART_DIR = Path(__file__).resolve().parents[1] / "skills" / "document-work" / "scripts"


def flowchart_module():
    """The skill ships flowchart.py so the model can also run it directly; builds reuse the same engine."""
    import importlib
    import sys
    folder = str(FLOWCHART_DIR)
    if folder not in sys.path:
        sys.path.insert(0, folder)
    return importlib.import_module("flowchart")


def is_diagram(block):
    return block["type"] == "code" and (block.get("lang") or "") in {"mermaid", "flowchart", "flow", "graph"}


def cut_runs(runs, cut):
    """Split a run list at character offset `cut`, keeping run formatting on both sides."""
    head, tail, consumed = [], [], 0
    for run in runs:
        value = run.get("text", "")
        if consumed + len(value) <= cut:
            head.append(run)
        elif consumed >= cut:
            tail.append(run)
        else:
            head.append({**run, "text": value[:cut - consumed]})
            tail.append({**run, "text": value[cut - consumed:]})
        consumed += len(value)
    return head, tail


def strip_runs(runs):
    runs = [dict(r) for r in runs if r.get("text")]
    if runs:
        runs[0]["text"] = runs[0]["text"].lstrip()
        runs[-1]["text"] = runs[-1]["text"].rstrip()
    return [r for r in runs if r["text"]]


STAGE_LABEL = re.compile(
    r"^\s*(第\s*[一二三四五六七八九十\d]+\s*(阶段|期|步|周|月|季度|年|天|轮|批)|(阶段|步骤|里程碑)\s*[一二三四五六七八九十\d]+"
    r"|(step|phase|stage|sprint|week|day|milestone|m|q|t\+?)\s*\d+|\d{4}\s*年(\s*\d{1,2}\s*月)?|\d{1,2}\s*月(\s*\d{1,2}\s*日)?"
    r"|(19|20)\d{2}[-./](0?[1-9]|1[0-2])([-./](0?[1-9]|[12]\d|3[01]))?(?![\d.])|(20\d{2})?\s*[qh][1-4]|近期|中期|远期|短期|长期|当前|现状|未来)", re.I)
LABEL_SEPARATORS = ("：", ":", "——", "—", "–", " - ", "｜", "|")


def split_label(runs):
    """“阶段一：说明” → (label runs, description runs); no short label → (runs, [])."""
    text = runs_text(runs)
    best = None
    for separator in LABEL_SEPARATORS:
        position = text.find(separator)
        if 0 < position <= 24 and (best is None or position < best[0]):
            best = (position, separator)
    if best is None:
        return strip_runs(runs), []
    position, separator = best
    label, rest = cut_runs(runs, position)
    return strip_runs(label), strip_runs(cut_runs(rest, len(separator))[1])


CHAIN_SEPARATOR = re.compile(r"\s*(?:→|➡|⇒|⟶|->|=>|>>)\s*")
STAT_LABEL = re.compile(r"^[+\-≥≤>≈~]?\d[\d.,]*\s*[^\s：:，,；;]{0,6}$")
# A decimal can spell a valid year/month ("2000.01"); an explicit measurement unit resolves the ambiguity.
MEASURED_STAT_LABEL = re.compile(
    r"[+\-≥≤>≈~]?\d[\d,]*(?:\.\d+)?\s*(?:"
    r"[kmgtpe]?i?b(?:ps|/s)?|bytes?|[nuµμm]?s|min|h|[kmgt]?hz|[km]?w(?:h)?"
    r"|usd|eur|gbp|cny|rmb|jpy|hkd|[%％‰℃]|°c"
    r"|[万亿]?(?:元|美元|港元|人|台|个|项|次)|倍|天|个月|个国家)", re.I)
LAYOUT_RANGES = {"cards": (2, 8), "timeline": (2, 8), "flow": (2, 8), "layers": (2, 6), "pyramid": (2, 6), "cycle": (3, 8), "stats": (2, 5)}
LAYOUT_NAMES = {"cards": "卡片", "timeline": "时间轴", "flow": "流程图", "layers": "分层图", "pyramid": "金字塔", "cycle": "循环图", "stats": "数字指标"}


def is_stage_label(label):
    """Dates and event prefixes are stages, unless the entire label is a value with a measurement unit."""
    return bool(STAGE_LABEL.match(label)) and not MEASURED_STAT_LABEL.fullmatch(label.strip())


def is_glyph(char):
    code = ord(char)
    return (0x1F000 <= code <= 0x1FAFF or 0x2600 <= code <= 0x27BF or code in (0xFE0F, 0x200D)
            or (code >= 0x2190 and unicodedata.category(char) in ("So", "Sm", "Sk")))


def split_glyph(runs):
    """A leading emoji or symbol becomes the segment icon: “🔍 资产识别” → (“资产识别”, “🔍”)."""
    text = runs_text(runs)
    count = 0
    while count < len(text) and count < 4 and is_glyph(text[count]):
        count += 1
    if not count or not text[count:].strip():
        return runs, None
    return strip_runs(cut_runs(runs, count)[1]), text[:count].strip()


def split_chain(runs):
    """“调研 → 设计 → 部署” → three short run lists; None unless 3–8 short pieces."""
    text = runs_text(runs)
    matches = list(CHAIN_SEPARATOR.finditer(text))
    if not 2 <= len(matches) <= 7:
        return None
    pieces, position, rest = [], 0, runs
    for match in matches:
        head, rest = cut_runs(rest, match.start() - position)
        rest = cut_runs(rest, match.end() - match.start())[1]
        position = match.end()
        pieces.append(strip_runs(head))
    pieces.append(strip_runs(rest))
    if any(not piece or len(runs_text(piece)) > 14 for piece in pieces):
        return None
    return pieces


# ---------------------------------------------------------------- Word

class WordBuilder:
    def __init__(self, request, warnings):
        from docx import Document
        from docx.shared import Pt
        self.warnings = warnings
        self.assets = request.get("assets", {})
        template = request.get("template")
        self.document = Document(template) if template else Document()
        self.templated = bool(template)
        self.output_dir = Path(request["output"]).parent
        self.body = self.document.element.body
        if template:
            self.inherit_header_footer()
            for child in list(self.body):
                if not child.tag.endswith("}sectPr"):
                    self.body.remove(child)
        else:
            self.default_styles()
        self.styles = {s.name for s in self.document.styles}
        section = self.document.sections[-1]
        self.content_width = section.page_width - section.left_margin - section.right_margin
        self.Pt = Pt

    def inherit_header_footer(self):
        """Only the body-level sectPr survives the clearing; a later section that inherits its header or footer
        from an earlier section would otherwise lose them, so copy the effective references onto it first."""
        import copy
        from docx.oxml.ns import qn
        props = [section._sectPr for section in self.document.sections]
        if len(props) < 2:
            return
        last = props[-1]
        references = (qn("w:headerReference"), qn("w:footerReference"))
        present = {(child.tag, child.get(qn("w:type"))) for child in last if child.tag in references}
        for earlier in reversed(props[:-1]):
            for child in earlier:
                key = (child.tag, child.get(qn("w:type")))
                if child.tag in references and key not in present:
                    last.insert(0, copy.deepcopy(child))
                    present.add(key)

    def default_styles(self):
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml.ns import qn
        section = self.document.sections[0]
        section.top_margin = section.bottom_margin = Cm(2.2)
        section.left_margin = section.right_margin = Cm(2.4)
        for name, size, color in (("Normal", 11, None), ("Title", 24, "16324F"), ("Subtitle", 13, "4A5A6A"),
                                  ("Heading 1", 16, "16324F"), ("Heading 2", 13.5, "16324F"), ("Heading 3", 12, "16324F"), ("Caption", 9, "5A6A7A")):
            style = self.document.styles[name]
            style.font.name = "微软雅黑"
            style.font.size = Pt(size)
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
            if color:
                style.font.color.rgb = RGBColor.from_string(color)
            if name.startswith("Heading"):
                style.paragraph_format.keep_with_next = True
                style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        normal = self.document.styles["Normal"]
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25

    def style(self, *candidates):
        return next((c for c in candidates if c in self.styles), None)

    def paragraph(self, runs, style=None, **fmt):
        paragraph = self.document.add_paragraph(style=style)
        for key, value in fmt.items():
            setattr(paragraph.paragraph_format, key, value)
        self.fill(paragraph, runs)
        return paragraph

    def fill(self, paragraph, runs, size=None, bold=None):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import RGBColor
        for item in runs:
            text = item.get("text", "")
            if not text:
                continue
            if item.get("link"):
                part = paragraph.part
                rid = part.relate_to(item["link"], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), rid)
                run = paragraph.add_run()
                paragraph._p.remove(run._r)
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)
                run.font.underline = True
                run.font.color.rgb = RGBColor.from_string("0563C1")
            else:
                run = paragraph.add_run()
            pieces = text.split("\n")
            for index, piece in enumerate(pieces):
                if index:
                    run.add_break()
                run.add_text(piece)
            if item.get("bold") or bold:
                run.bold = True
            if item.get("italic"):
                run.italic = True
            if item.get("code"):
                run.font.name = "Consolas"
                run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
            if size:
                run.font.size = size

    def heading(self, level, runs):
        style = self.style(f"Heading {level}", "Heading 3", "Heading 2", "Heading 1")
        if style:
            self.paragraph(runs, style=style)
            return
        paragraph = self.paragraph(runs, keep_with_next=True, space_before=self.Pt(10))
        for run in paragraph.runs:
            run.bold = True
            run.font.size = self.Pt(18 - 2 * min(level, 3))

    def list(self, items):
        from docx.shared import Cm
        ordered_counter = {}
        for item in items:
            level = item["level"]
            suffix = "" if level == 0 else f" {level + 1}"
            style = self.style(("List Number" if item["ordered"] else "List Bullet") + suffix,
                               "List Number" if item["ordered"] else "List Bullet")
            if style:
                self.paragraph(item["runs"], style=style, left_indent=Cm(0.75 * (level + 1)) if suffix else None)
                continue
            if item["ordered"]:
                ordered_counter[level] = ordered_counter.get(level, 0) + 1
                for deeper in [k for k in ordered_counter if k > level]:
                    ordered_counter.pop(deeper)
                marker = f"{ordered_counter[level]}. "
            else:
                marker = ("• ", "– ", "· ")[min(level, 2)]
            self.paragraph([{"text": marker}] + item["runs"], left_indent=Cm(0.75 * (level + 1)), first_line_indent=Cm(-0.5))

    def table(self, header, rows, align):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        columns = max(len(header), *(len(r) for r in rows)) if rows else len(header)
        table = self.document.add_table(rows=1, cols=columns)
        style = self.style("Table Grid", "Light Shading Accent 1", "Light List")
        if style:
            table.style = style
        alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
        def write(cells, runs_list, is_header):
            for index, cell in enumerate(cells):
                paragraph = cell.paragraphs[0]
                if index < len(runs_list):
                    self.fill(paragraph, runs_list[index], size=self.Pt(10), bold=is_header)
                if index < len(align) and align[index] in alignment:
                    paragraph.alignment = alignment[align[index]]
                paragraph.paragraph_format.space_after = self.Pt(2)
        write(table.rows[0].cells, header, True)
        header_props = table.rows[0]._tr.get_or_add_trPr()
        header_props.append(OxmlElement("w:tblHeader"))
        for cell in table.rows[0].cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EEF5")
            shade.set(qn("w:val"), "clear")
            cell._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            write(table.add_row().cells, row, False)
        self.document.add_paragraph()

    def image(self, path, alt, width):
        from docx.shared import Emu
        real = self.assets.get(path)
        if not real or Path(real).suffix.lower() not in IMAGE_TYPES:
            self.warnings.append(f"图片未插入（不存在或格式不支持）：{path}")
            self.paragraph([{"text": f"[图片：{alt or path}]", "italic": True}])
            return
        wanted = parse_width(width, self.content_width)
        if not wanted:
            px_w, px_h = image_size(real)
            natural = int(px_w / 96 * EMU_PER_INCH)
            wanted = min(natural, self.content_width)
        paragraph = self.document.add_paragraph()
        paragraph.alignment = 1
        paragraph.paragraph_format.keep_with_next = bool(alt)
        try:
            paragraph.add_run().add_picture(real, width=Emu(wanted))
        except Exception as error:  # noqa: BLE001 - report as a warning, keep the document usable
            self.warnings.append(f"图片插入失败：{path}（{error}）")
            self.fill(paragraph, [{"text": f"[图片：{alt or path}]", "italic": True}])
            return
        if alt:
            caption = self.paragraph([{"text": alt}], style=self.style("Caption"))
            caption.alignment = 1
            if not self.style("Caption"):
                for run in caption.runs:
                    run.font.size = self.Pt(9)
                    run.italic = True

    def quote(self, paragraphs):
        from docx.shared import Cm
        style = self.style("Quote", "Intense Quote")
        for runs in paragraphs:
            paragraph = self.paragraph(runs, style=style)
            if not style:
                paragraph.paragraph_format.left_indent = Cm(1)
                for run in paragraph.runs:
                    run.italic = True

    def diagram(self, text):
        """A ```mermaid block becomes a flowchart picture; without a CJK font the source stays as code."""
        import tempfile
        from docx.shared import Emu
        flowchart = flowchart_module()
        path = None
        try:
            try:
                graph = flowchart.parse(text)
                lay = flowchart.layout(graph, graph.direction or "TB", 12)
                font = flowchart.find_cjk_font()
                handle, path = tempfile.mkstemp(prefix="flow-", suffix=".png", dir=str(self.output_dir))
                os.close(handle)
                flowchart.render_png(lay, path, 1600, font)
            except (ValueError, LookupError) as error:
                self.warnings.append(f"流程图未生成，已保留源码：{error}")
                self.code(text)
                return
            self.warnings.extend(f"流程图：{w}" for w in lay["warnings"])
            section = self.document.sections[-1]
            content_height = section.page_height - section.top_margin - section.bottom_margin
            bbox_w, bbox_h = lay["bbox"][2], lay["bbox"][3]
            width = min(int(bbox_w * EMU_PER_PT * 1.1), self.content_width, int(content_height * 0.75 * bbox_w / max(bbox_h, 1)))
            paragraph = self.document.add_paragraph()
            paragraph.alignment = 1
            paragraph.add_run().add_picture(path, width=Emu(width))
        finally:
            if path:
                Path(path).unlink(missing_ok=True)

    def code(self, text):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        for line in text.split("\n"):
            paragraph = self.paragraph([{"text": line or " ", "code": True}], space_after=self.Pt(0))
            for run in paragraph.runs:
                run.font.size = self.Pt(9)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F2F4F7")
            shade.set(qn("w:val"), "clear")
            paragraph._p.get_or_add_pPr().append(shade)
        self.document.add_paragraph()

    def build(self, request):
        title, subtitle = request.get("title"), request.get("subtitle")
        if title:
            self.paragraph([{"text": title}], style=self.style("Title"))
            if not self.style("Title"):
                for run in self.document.paragraphs[-1].runs:
                    run.bold = True
                    run.font.size = self.Pt(22)
        if subtitle:
            self.paragraph([{"text": subtitle}], style=self.style("Subtitle"))
        for block in request["blocks"]:
            kind = block["type"]
            if kind == "heading":
                self.heading(block["level"], block["runs"])
            elif kind == "paragraph":
                self.paragraph(block["runs"])
            elif kind == "list":
                self.list(block["items"])
            elif kind == "table":
                self.table(block["header"], block["rows"], block.get("align", []))
            elif kind == "image":
                self.image(block["path"], block.get("alt", ""), block.get("width"))
            elif kind == "quote":
                self.quote(block["paragraphs"])
            elif kind == "code":
                if is_diagram(block):
                    self.diagram(block["text"])
                else:
                    self.code(block["text"])
            elif kind == "pagebreak":
                self.document.add_page_break()
            elif kind == "note":
                self.warnings.append("Word 不使用备注指令，已忽略：" + block["text"][:40])
            # "layout" directives only shape slides; Word keeps headings and lists as written.
        self.document.save(request["output"])
        return {"headings": sum(1 for b in request["blocks"] if b["type"] == "heading"), "templated": self.templated}


# ---------------------------------------------------------------- PowerPoint

class TitleStyle:
    """Title geometry and font inferred from a template's own slides."""

    def __init__(self, width, height):
        self.left, self.top = int(width * 0.07), int(height * 0.05)
        self.width, self.height = int(width * 0.86), int(height * 0.12)
        self.font, self.size, self.bold = None, 28, True
        self.color = None  # ("rgb", "RRGGBB") | ("theme", int)
        self.align = None
        self.samples = 0
        self.placeholder = False

    def describe(self):
        return {"left": round(self.left / EMU_PER_INCH, 2), "top": round(self.top / EMU_PER_INCH, 2),
                "width": round(self.width / EMU_PER_INCH, 2), "height": round(self.height / EMU_PER_INCH, 2),
                "font": self.font, "size": self.size, "bold": self.bold, "color": self.color, "samples": self.samples}


def first_run(shape):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                return paragraph, run
    return None, None


def run_color(run):
    try:
        color = run.font.color
        if color is None or color.type is None:
            return None
        if color.type == 1:
            return ("rgb", str(color.rgb))
        return ("theme", int(color.theme_color))
    except (AttributeError, ValueError):
        return None


def infer_title_style(presentation, pages=None):
    from collections import Counter
    width, height = presentation.slide_width, presentation.slide_height
    style = TitleStyle(width, height)
    candidates = []
    for index, slide in enumerate(presentation.slides, 1):
        if pages and index not in pages:
            continue
        best = None
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.top is None or shape.top > height * 0.22 or shape.width < width * 0.25 or shape.height > height * 0.25:
                continue
            paragraph, run = first_run(shape)
            if run is None:
                continue
            size = run.font.size.pt if run.font.size else (paragraph.font.size.pt if paragraph.font.size else None)
            is_title = shape.is_placeholder and shape.placeholder_format.type in (1, 3)  # TITLE, CENTER_TITLE
            score = (1 if is_title else 0, size or 0, -shape.top)
            if best is None or score > best[0]:
                best = (score, shape, run, paragraph, size, is_title)
        if best:
            candidates.append(best)
    if not candidates:
        return style
    style.samples = len(candidates)
    key = Counter((round(c[1].top / 45720), round(c[1].height / 45720), c[4]) for c in candidates).most_common(1)[0][0]
    chosen = next(c for c in candidates if (round(c[1].top / 45720), round(c[1].height / 45720), c[4]) == key)
    _, shape, run, paragraph, size, is_title = chosen
    style.left, style.top, style.width, style.height = shape.left, shape.top, shape.width, shape.height
    style.width = max(style.width, int(width * 0.6))
    style.font = run.font.name or paragraph.font.name
    style.size = int(size) if size else (32 if is_title else 24)
    style.bold = run.font.bold if run.font.bold is not None else True
    style.color = run_color(run)
    style.align = paragraph.alignment
    style.placeholder = is_title
    return style


def layout_for_content(presentation, pages=None):
    """Prefer the layout used by the template's own content slides, else the emptiest layout."""
    from collections import Counter
    layouts = [l for m in presentation.slide_masters for l in m.slide_layouts]
    usage = Counter()
    for index, slide in enumerate(presentation.slides, 1):
        if pages and index not in pages:
            continue
        if index > 1 and any(s.has_text_frame and s.text_frame.text.strip() for s in slide.shapes):
            usage[slide.slide_layout.part.partname] += 1
    if usage:
        wanted = usage.most_common(1)[0][0]
        return next(l for l in layouts if l.part.partname == wanted)
    def score(layout):
        name = (layout.name or "").lower()
        placeholders = [p.placeholder_format.type for p in layout.placeholders]
        body = sum(1 for t in placeholders if t in (2, 7))  # BODY, OBJECT
        return (0 if "title" in name and "only" in name or "仅标题" in name else 1, body, len(layout.shapes))
    return min(layouts, key=score) if layouts else presentation.slide_layouts[6 if len(presentation.slide_layouts) > 6 else -1]


def cover_layout(presentation):
    for layout in (l for m in presentation.slide_masters for l in m.slide_layouts):
        name = (layout.name or "")
        if "标题幻灯片" in name or name.lower() in ("title slide", "title"):
            return layout
    return None


def set_color(font, color):
    if not color:
        return
    from pptx.dml.color import RGBColor
    if color[0] == "rgb":
        font.color.rgb = RGBColor.from_string(color[1])
    else:
        font.color.theme_color = color[1]


def apply_bullet(paragraph, level, ordered, number):
    """Non-placeholder text boxes carry no bullets; write DrawingML bullets directly."""
    props = paragraph._p.get_or_add_pPr()
    for child in list(props):
        if child.tag in ("{%s}buNone" % NS_A, "{%s}buChar" % NS_A, "{%s}buAutoNum" % NS_A, "{%s}buFont" % NS_A):
            props.remove(child)
    indent = 285750 + level * 285750
    props.set("marL", str(indent))
    props.set("indent", str(-285750))
    props.set("lvl", str(level))
    if ordered:
        bullet = etree.SubElement(props, "{%s}buAutoNum" % NS_A)
        bullet.set("type", "arabicPeriod")
        if number > 1:
            bullet.set("startAt", str(number))
    else:
        font = etree.SubElement(props, "{%s}buFont" % NS_A)
        font.set("typeface", "Arial")
        bullet = etree.SubElement(props, "{%s}buChar" % NS_A)
        bullet.set("char", ("•", "–", "·")[min(level, 2)])


class SlideBuilder:
    def __init__(self, request, warnings):
        from pptx import Presentation
        from pptx.util import Inches
        self.warnings = warnings
        self.assets = request.get("assets", {})
        template = request.get("template")
        self.presentation = Presentation(template) if template else Presentation()
        if not template:
            self.presentation.slide_width, self.presentation.slide_height = Inches(13.333333), Inches(7.5)
        self.width, self.height = self.presentation.slide_width, self.presentation.slide_height
        template_pages = list(range(1, len(self.presentation.slides) + 1))
        sample_pages = set(request.get("styleFrom") or template_pages)
        self.title_style = infer_title_style(self.presentation, sample_pages)
        self.layout = layout_for_content(self.presentation, sample_pages)
        self.cover = cover_layout(self.presentation)
        self.body_font = self.title_style.font
        margin = int(self.width * 0.06)
        top = self.title_style.top + self.title_style.height + int(self.height * 0.03)
        bottom = int(self.height * 0.92)
        self.region = (margin, top, self.width - 2 * margin, max(bottom - top, int(self.height * 0.4)))
        self.attention = []
        self.layouts = []
        self.generated = []

    # -- slide bookkeeping -------------------------------------------------
    def remove_slides(self, keep):
        from pptx.opc.package import XmlPart
        from pptx.opc.packuri import PackURI
        slides = self.presentation.slides
        id_list = slides._sldIdLst
        entries = list(id_list)
        page_of_part = {slides[index - 1].part: index for index in range(1, len(entries) + 1)}
        dropped = {part for part, index in page_of_part.items() if index not in keep}
        for index, entry in enumerate(entries, 1):
            if index not in keep:
                id_list.remove(entry)
        # Any part still reachable without passing through a dropped slide (kept slides, master, layouts, their notes,
        # custom shows) may link to a dropped slide; the link would keep the stale part alive and its name would
        # collide with a new slide, so cut them all.
        presentation_part = self.presentation.part
        reachable, stack = [], [presentation_part]
        while stack:
            part = stack.pop()
            if part in reachable or part in dropped:
                continue
            reachable.append(part)
            if isinstance(part, XmlPart):
                stack.extend(rel.target_part for rel in part.rels.values() if not rel.is_external)
        elsewhere = 0
        for part in reachable:
            if not isinstance(part, XmlPart):
                continue
            removed = self.unlink(part, dropped)
            if part is presentation_part:
                continue  # its own slide-list relationships are not user-visible links
            if removed and part in page_of_part:
                self.warnings.append(f"模板第 {page_of_part[part]} 页有 {removed} 处指向未保留页面的链接，已移除")
            elif removed:
                elsewhere += removed
        if elsewhere:
            self.warnings.append(f"模板母版、版式或备注中有 {elsewhere} 处指向未保留页面的链接，已移除")
        # python-pptx names new slides by count, so surviving parts must be renumbered first; a slide part still
        # reachable through an unexpected reference is moved to numbers new pages never reach.
        current = []
        for position, slide in enumerate(self.presentation.slides, 1):
            slide.part.partname = PackURI(f"/ppt/slides/slide{position}.xml")
            current.append(slide.part)
        stray = [part for part in presentation_part.package.iter_parts()
                 if re.fullmatch(r"/ppt/slides/slide\d+\.xml", str(part.partname)) and part not in current]
        for offset, part in enumerate(stray, 10001):
            part.partname = PackURI(f"/ppt/slides/slide{offset}.xml")
        return {page: entries[page - 1] for page in keep}

    def unlink(self, part, targets):
        """Drop every relationship from `part` to one of `targets`, removing the XML elements that used it."""
        removed = 0
        for rId, relationship in list(part.rels.items()):
            if relationship.is_external or relationship.target_part not in targets:
                continue
            for element in part._element.xpath(f'//*[@r:id="{rId}"]'):
                parent = element.getparent()
                parent.remove(element)
                # An emptied custom-show slide list is invalid, so drop the whole show with it.
                if parent.tag.endswith("}sldLst") and len(parent) == 0:
                    show = parent.getparent()
                    show_list = show.getparent()
                    show_list.remove(show)
                    if len(show_list) == 0:
                        show_list.getparent().remove(show_list)
            part.rels.pop(rId)
            removed += 1
        return removed

    def reorder(self, sequence):
        id_list = self.presentation.slides._sldIdLst
        for entry in list(id_list):
            id_list.remove(entry)
        for entry in sequence:
            id_list.append(entry)

    def new_slide(self, layout=None, keep=(1, 3, 13, 15, 16)):
        """Keep title, number, footer and date placeholders; body placeholders would show prompt text."""
        slide = self.presentation.slides.add_slide(layout or self.layout)
        for shape in list(slide.placeholders):
            if shape.placeholder_format.type not in keep:
                shape._element.getparent().remove(shape._element)
        return slide

    def subtitle_placeholder(self, slide):
        return next((s for s in slide.placeholders if s.placeholder_format.type == 4), None)

    # -- text helpers ------------------------------------------------------
    def style_run(self, run, size, bold=None, color=None, font=None):
        from pptx.util import Pt
        run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if font or self.body_font:
            run.font.name = font or self.body_font
        set_color(run.font, color)

    def add_runs(self, paragraph, runs, size, bold=None, color=None):
        for item in runs:
            text = item.get("text", "")
            if not text:
                continue
            # python-pptx escapes control characters in run text; explicit line breaks keep "\n" visible.
            for index, piece in enumerate(text.split("\n")):
                if index:
                    paragraph.add_line_break()
                if not piece:
                    continue
                run = paragraph.add_run()
                run.text = piece
                self.style_run(run, size, bold=True if item.get("bold") else bold, color=color, font="Consolas" if item.get("code") else None)
                if item.get("italic"):
                    run.font.italic = True
                if item.get("link"):
                    run.hyperlink.address = item["link"]

    def add_title(self, slide, text, cover=False, subtitle=None):
        from pptx.util import Pt
        style = self.title_style
        shape = slide.shapes.title
        size = style.size if not cover else max(style.size, 36)
        available = (style.width if not cover else int(self.width * 0.8)) / EMU_PER_PT
        while char_units(text, size) > available * 0.96 and size > 18:
            size -= 2
        if shape is not None:
            frame = shape.text_frame
            frame.clear()
            paragraph = frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = text
            if size != style.size and not cover:
                run.font.size = Pt(size)
            holder = self.subtitle_placeholder(slide)
            if cover and not subtitle and holder is not None:
                holder._element.getparent().remove(holder._element)
            if cover and subtitle:
                if holder is not None:
                    holder.text_frame.clear()
                    holder.text_frame.paragraphs[0].add_run().text = subtitle
                else:
                    sub = slide.shapes.add_textbox(shape.left, shape.top + shape.height, shape.width, int(self.height * 0.1)).text_frame
                    sub.word_wrap = True
                    sub.paragraphs[0].alignment = 2
                    run = sub.paragraphs[0].add_run()
                    run.text = subtitle
                    self.style_run(run, 20, bold=False, color=style.color, font=style.font)
            return
        for holder in list(slide.placeholders):
            if holder.placeholder_format.type == 4:
                holder._element.getparent().remove(holder._element)
        if cover:
            left, top, width, height = int(self.width * 0.1), int(self.height * 0.34), int(self.width * 0.8), int(self.height * 0.18)
        else:
            left, top, width, height = style.left, style.top, style.width, style.height
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = 3  # MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        if style.align is not None and not cover:
            paragraph.alignment = style.align
        if cover:
            paragraph.alignment = 2  # center
        run = paragraph.add_run()
        run.text = text
        self.style_run(run, size, bold=style.bold, color=style.color, font=style.font)
        if cover and subtitle:
            sub = slide.shapes.add_textbox(left, top + height, width, int(self.height * 0.1)).text_frame
            sub.word_wrap = True
            sub.paragraphs[0].alignment = 2
            run = sub.paragraphs[0].add_run()
            run.text = subtitle
            self.style_run(run, 20, bold=False, color=style.color, font=style.font)

    def estimate_height(self, lines, size):
        """Height in EMU for text at size pt within the current width; lines are (runs, level, kind, bold)."""
        total = 0.0
        for runs, level, kind, _ in lines:
            text = runs_text(runs)
            available = self.text_width_pt - (level * 22.5 if level else 0)
            line_size = size + 2 if kind == "subheading" else size
            # Explicit line breaks are written as real breaks, so each piece wraps on its own.
            count = sum(max(1, math.ceil(char_units(piece, line_size) / max(available, 1))) for piece in text.split("\n"))
            total += count * line_size * 1.22 + line_size * 0.45
        return int(total * EMU_PER_PT)

    def split_runs(self, runs, capacity, available):
        """Cut a run list after roughly `capacity` wrapped lines of `available` em each, preferring a line break or sentence punctuation."""
        text = runs_text(runs)
        cut, used, width = 0, 1, 0.0
        for index, char in enumerate(text):
            if char == "\n":
                used, width = used + 1, 0.0
            else:
                width += char_units(char, 1)
                if width > available:
                    used, width = used + 1, char_units(char, 1)
            if used > capacity:
                break
            cut = index + 1
        window = text[max(0, cut - 80):cut]
        for mark in "\n。！？；.!?;，,、":
            position = window.rfind(mark)
            if position >= 0:
                cut = cut - len(window) + position + 1
                break
        if cut <= 0 or cut >= len(text):
            return runs, []
        return cut_runs(runs, cut)

    def add_text_block(self, slide, lines, left, top, width, height):
        """Fill a region with paragraphs; returns the lines that did not fit."""
        from pptx.util import Pt
        self.text_width_pt = width / EMU_PER_PT - 14
        size = 18
        while size > 14 and self.estimate_height(lines, size) > height:
            size -= 1
        fitting = lines
        overflow = []
        if self.estimate_height(lines, size) > height:
            fitting = []
            for index, line in enumerate(lines):
                if self.estimate_height(fitting + [line], size) <= height:
                    fitting.append(line)
                    continue
                # A paragraph longer than the remaining space is split so the slide is not left empty.
                remaining = height - self.estimate_height(fitting, size)
                capacity = int(remaining / EMU_PER_PT / (size * 1.22)) - 1
                runs, level, kind, bold = line
                if capacity >= 2 and kind in ("text", "quote"):
                    head, tail = self.split_runs(runs, capacity, (self.text_width_pt - level * 22.5) / size)
                    if tail and self.estimate_height(fitting + [(head, level, kind, bold)], size) <= height:
                        fitting.append((head, level, kind, bold))
                        overflow = [(tail, level, kind, bold)] + lines[index + 1:]
                        break
                if not fitting:
                    fitting.append(line)
                    overflow = lines[index + 1:]
                else:
                    overflow = lines[index:]
                break
        box = slide.shapes.add_textbox(left, top, width, height)
        box.text_frame.word_wrap = True
        self.write_lines(box.text_frame, fitting, size)
        if size < 16:
            self.attention.append({"slide": slide, "reason": f"正文字号缩至 {size}pt，请检查可读性"})
        return overflow

    def write_lines(self, frame, lines, size, first=True):
        """Append (runs, level, kind, bold) lines as paragraphs with bullets, numbering and quote indents."""
        from pptx.util import Pt
        counters = {}
        for runs, level, kind, bold in lines:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.space_after = Pt(size * 0.45)
            if kind in ("bullet", "number"):
                counters[level] = counters.get(level, 0) + 1 if kind == "number" else 0
                for deeper in [k for k in counters if k > level]:
                    counters.pop(deeper)
                apply_bullet(paragraph, level, kind == "number", counters.get(level, 1))
            elif kind == "quote":
                props = paragraph._p.get_or_add_pPr()
                props.set("marL", "285750")
            self.add_runs(paragraph, runs, size if kind != "subheading" else size + 2, bold=True if (bold or kind == "subheading") else None)

    # -- cards and timelines ----------------------------------------------
    def block_lines(self, block):
        kind = block["type"]
        if kind == "heading":
            return [(block["runs"], 0, "subheading", True)]
        if kind == "paragraph":
            return [(block["runs"], 0, "text", False)]
        if kind == "list":
            return [(item["runs"], item["level"], "number" if item["ordered"] else "bullet", False) for item in block["items"]]
        if kind == "quote":
            return [(runs, 0, "quote", False) for runs in block["paragraphs"]]
        if kind == "code":
            return [([{"text": line or " ", "code": True}], 0, "text", False) for line in block["text"].split("\n")]
        return []

    def segments(self, group):
        """Split one slide's blocks into intro lines plus labelled segments (### headings, list items or an arrow chain)."""
        blocks = [b for b in group if b["type"] != "layout"]
        if any(b["type"] in ("table", "code") for b in blocks) or not blocks:
            return None
        if any(b["type"] == "heading" for b in blocks):
            intro, segments, current = [], [], None
            for block in blocks:
                if block["type"] == "heading":
                    label, glyph = split_glyph(strip_runs(block["runs"]))
                    current = {"label": label, "glyph": glyph, "icon": None, "lines": []}
                    segments.append(current)
                elif block["type"] == "image":
                    # One image inside a ### segment is that segment's icon, wherever the Markdown put it.
                    if current is None or current["icon"]:
                        return None
                    current["icon"] = block["path"]
                elif current is None:
                    intro.extend(self.block_lines(block))
                else:
                    current["lines"].extend(self.block_lines(block))
            return "headings", intro, segments
        if any(b["type"] == "image" for b in blocks):
            return None
        intro = [line for block in blocks[:-1] for line in self.block_lines(block)]
        if any(b["type"] != "paragraph" for b in blocks[:-1]):
            return None
        if blocks[-1]["type"] == "paragraph":
            chain = split_chain(blocks[-1]["runs"])
            return ("chain", intro, [self.segment(runs) for runs in chain]) if chain else None
        if blocks[-1]["type"] != "list":
            return None
        segments, current = [], None
        for item in blocks[-1]["items"]:
            if item["level"] == 0 or current is None:
                current = self.segment(item["runs"])
                segments.append(current)
            else:
                current["lines"].append((item["runs"], 0, "bullet", False))
        return "list", intro, segments

    def segment(self, runs):
        label, description = split_label(runs)
        label, glyph = split_glyph(label)
        return {"label": label, "glyph": glyph, "icon": None, "labelled": bool(description),
                "lines": [(description, 0, "text", False)] if description else []}

    def choose_layout(self, slide, group):
        """An explicit directive wins; otherwise recognise arrow chains, short ### groups, layer names, figures and stage lists.
        A page that looks like a diagram candidate but matches no rule is reported, never silently left as a list."""
        explicit = next((b["mode"] for b in group if b["type"] == "layout"), None)
        if explicit == "plain":
            return None, None
        parsed = self.segments(group)
        if parsed is None:
            return explicit, None
        origin, intro, segments = parsed
        count = len(segments)
        intro_chars = sum(len(runs_text(line[0])) for line in intro)
        pictures = any(segment.get("icon") for segment in segments)
        if explicit in LAYOUT_RANGES:
            if pictures and explicit in ("layers", "pyramid", "cycle", "stats"):
                self.warnings.append(f"“{explicit}”图示不放图片图标，该页已按普通版式排版；改用 cards / flow / timeline 或去掉图片")
                return None, None
            low, high = LAYOUT_RANGES[explicit]
            return explicit, (intro, segments) if low <= count <= high and intro_chars <= 300 else None
        body_chars = [sum(len(runs_text(line[0])) for line in s["lines"]) for s in segments]
        labels = [runs_text(s["label"]) for s in segments]
        short = all(c <= 220 and len(s["lines"]) <= 6 for c, s in zip(body_chars, segments))
        if intro_chars <= 160:
            if origin == "chain":
                return "flow", (intro, segments)
            if origin == "headings" and 2 <= count <= 6 and short:
                if all(label.endswith("层") for label in labels) and not pictures:
                    return "layers", (intro, segments)
                return "cards", (intro, segments)
            # Resolve date-like decimals by their unit before choosing either layout.
            if origin == "list" and 2 <= count <= 5 and all(STAT_LABEL.match(label) and not is_stage_label(label) for label in labels) and all(c <= 60 for c in body_chars):
                return "stats", (intro, segments)
            if origin == "list" and 3 <= count <= 6 and all(
                    is_stage_label(label) and len(label) <= 24 and c <= 120 for label, c in zip(labels, body_chars)):
                return "timeline", (intro, segments)
        candidate = (origin == "headings" and 2 <= count <= 6) or (origin == "list" and 2 <= count <= 8 and all(s.get("labelled") for s in segments))
        if candidate:
            if intro_chars > 160:
                why = f"图示前的引导文字有 {intro_chars} 字，超过 160 字"
            elif origin == "headings":
                why = f"{count} 个 ### 小节中有的超过 6 行或 220 字"
            else:
                why = f"{count} 项“标签：说明”列表的标签不是阶段或数字模式"
            self.attention.append({"slide": slide, "reason": why + "，未自动排成图示，已按普通版式排版；需要卡片、时间轴、指标等请在 ## 标题后加 <!-- cards -->、<!-- timeline -->、<!-- stats --> 等注释并精简文字"})
        return None, None

    def accent(self):
        color = self.title_style.color
        return color if color and color[0] == "rgb" and color[1].upper() not in ("000000", "FFFFFF") else ("theme", 5)

    def plain_shape(self, shape, fill=None, line=None, line_pt=0.75):
        """Drop python-pptx's theme style reference so fill, line and text colours are explicit."""
        from pptx.dml.color import RGBColor
        from pptx.util import Pt
        style = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            shape._element.remove(style)
        if fill:
            shape.fill.solid()
            if fill[0] == "theme":
                shape.fill.fore_color.theme_color = fill[1]
            else:
                shape.fill.fore_color.rgb = RGBColor.from_string(fill[1])
        elif hasattr(shape, "fill"):
            shape.fill.background()
        if line:
            shape.line.width = Pt(line_pt)
            if line[0] == "theme":
                shape.line.color.theme_color = line[1]
            else:
                shape.line.color.rgb = RGBColor.from_string(line[1])
        else:
            shape.line.fill.background()

    def arrow_head(self, connector):
        line = connector.line._get_or_add_ln()
        tail = etree.SubElement(line, "{%s}tailEnd" % NS_A)
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")

    def centered_text(self, shape, pad, middle=True):
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = 3 if middle else 1
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = pad
        frame.paragraphs[0].alignment = 2
        return frame

    def center_paragraphs(self, frame):
        for paragraph in frame.paragraphs:
            paragraph.alignment = 2

    def has_icons(self, segments):
        return any(s.get("icon") or s.get("glyph") for s in segments)

    def icon_badge(self, slide, segment, cx, top, size_emu):
        """Draw an image icon or a glyph badge centred on cx; returns the height used."""
        if segment.get("icon"):
            real = self.assets.get(segment["icon"])
            if real and Path(real).suffix.lower() in IMAGE_TYPES:
                px_w, px_h = image_size(real)
                scale = size_emu / max(px_w, px_h)
                draw_w, draw_h = max(1, int(px_w * scale)), max(1, int(px_h * scale))
                picture = slide.shapes.add_picture(real, cx - draw_w // 2, top + (size_emu - draw_h) // 2, draw_w, draw_h)
                picture.name = "Icon"
                return size_emu
            self.warnings.append(f"图标未插入（不存在或格式不支持）：{segment['icon']}")
        if segment.get("glyph"):
            from pptx.enum.shapes import MSO_SHAPE
            badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - size_emu // 2, top, size_emu, size_emu)
            badge.name = "Icon " + segment["glyph"]
            self.plain_shape(badge, fill=("rgb", "E3ECF8"))
            frame = self.centered_text(badge, 0)
            run = frame.paragraphs[0].add_run()
            run.text = segment["glyph"]
            self.style_run(run, max(12, int(size_emu / EMU_PER_PT * 0.5)), bold=False, color=self.accent(), font="Segoe UI Emoji")
            return size_emu
        return 0

    def intro_geometry(self, intro):
        """Region left for a diagram once the intro lines take the top; the intro is drawn later."""
        left, top, width, height = self.region
        if not intro:
            return left, top, width, height, 0
        used = min(int(height * 0.3), self.estimate_text_height(intro, width))
        gap = int(self.height * 0.02)
        return left, top + used + gap, width, height - used - gap, used

    def draw_intro(self, slide, intro, used):
        if intro:
            left, top, width, _ = self.region
            if self.add_text_block(slide, intro, left, top, width, used):
                self.warnings.append("图示页的引导文字过长，超出部分未显示")

    def note_layout(self, slide, mode, count, detail=""):
        self.layouts.append({"slide": slide, "mode": mode, "segments": count})
        self.attention.append({"slide": slide, "reason": f"已排成 {count} 段{LAYOUT_NAMES[mode]}{detail}，请检查图示文字"})

    def render_cards(self, slide, intro, segments):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        left, top, width, height, used = self.intro_geometry(intro)
        count = len(segments)
        columns = count if count <= 3 else (2 if count == 4 else (3 if count <= 6 else 4))
        rows = math.ceil(count / columns)
        gap = int(self.width * 0.02)
        card_w, card_h = (width - gap * (columns - 1)) // columns, (height - gap * (rows - 1)) // rows
        inset = Pt(10)
        icon_h = int(self.height * 0.08) if self.has_icons(segments) else 0
        icon_gap = int(self.height * 0.012) if icon_h else 0
        self.text_width_pt = (card_w - 2 * inset) / EMU_PER_PT - 4
        size = None
        for candidate in (16, 15, 14):
            needed = max(self.estimate_height([(s["label"], 0, "subheading", True)] + s["lines"], candidate) for s in segments)
            if needed + 2 * inset + icon_h + icon_gap <= card_h:
                size = candidate
                break
        if size is None:
            return False
        self.draw_intro(slide, intro, used)
        # Short cards keep a compact height instead of stretching to the bottom of the slide.
        if rows == 1:
            card_h = min(card_h, max(needed + 2 * inset + icon_h + icon_gap + int(self.height * 0.05), int(height * 0.5)))
        for index, segment in enumerate(segments):
            row, column = divmod(index, columns)
            x, y = left + column * (card_w + gap), top + row * (card_h + gap)
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
            shape.adjustments[0] = 0.06
            shape.name = "Card " + runs_text(segment["label"])[:32]
            self.plain_shape(shape, fill=("rgb", "F5F8FC"), line=("rgb", "D6DEE8"))
            frame = self.centered_text(shape, inset, middle=False)
            if icon_h:
                self.icon_badge(slide, segment, x + card_w // 2, y + inset, icon_h)
                frame.margin_top = inset + icon_h + icon_gap
            paragraph = frame.paragraphs[0]
            paragraph.space_after = Pt(size * 0.5)
            self.add_runs(paragraph, segment["label"], size + 3, bold=True, color=self.accent())
            self.write_lines(frame, segment["lines"], size, first=False)
        self.note_layout(slide, "cards", count, f"（{columns} 栏）")
        return True

    def render_timeline(self, slide, intro, segments):
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        left, top, width, height, used = self.intro_geometry(intro)
        count = len(segments)
        rows = 1 if count <= 6 else 2
        per_row = math.ceil(count / rows)
        row_h = height // rows
        column_w = width // per_row
        pad = int(column_w * 0.06)
        diameter = max(int(self.height * 0.055), min(int(self.height * 0.08), column_w // 3))
        label_h = int(row_h * 0.26)
        line_y_offset = label_h + diameter // 2 + int(self.height * 0.01)
        description_h = row_h - line_y_offset - diameter // 2 - int(self.height * 0.015)
        self.text_width_pt = (column_w - 2 * pad) / EMU_PER_PT - 4
        size = None
        for candidate in (16, 15, 14):
            if all(self.estimate_height(s["lines"], candidate) <= description_h and
                   self.estimate_height([(s["label"], 0, "text", True)], candidate + 2) <= label_h for s in segments):
                size = candidate
                break
        if size is None:
            return False
        self.draw_intro(slide, intro, used)
        accent = self.accent()
        # A single row floats slightly above the vertical centre of the free area.
        deepest = max(self.estimate_height(s["lines"], size) for s in segments)
        block_h = line_y_offset + diameter // 2 + int(self.height * 0.015) + deepest
        shift = max(0, (row_h - block_h) // 3) if rows == 1 else 0
        for row in range(rows):
            items = segments[row * per_row:(row + 1) * per_row]
            row_top = top + row * row_h + shift
            line_y = row_top + line_y_offset
            if len(items) > 1:
                x1, x2 = left + column_w // 2, left + column_w * (len(items) - 1) + column_w // 2
                connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, line_y, x2, line_y)
                connector.name = "Timeline line"
                self.plain_shape(connector, line=accent, line_pt=2.25)
            for index, segment in enumerate(items):
                number = row * per_row + index + 1
                cx = left + column_w * index + column_w // 2
                if segment.get("icon") and self.assets.get(segment["icon"]):
                    badge = int(diameter * 1.2)
                    self.icon_badge(slide, segment, cx, line_y - badge // 2, badge)
                else:
                    node = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - diameter // 2, line_y - diameter // 2, diameter, diameter)
                    node.name = f"Timeline node {number}"
                    self.plain_shape(node, fill=accent)
                    paragraph = self.centered_text(node, 0).paragraphs[0]
                    if segment.get("glyph"):
                        run = paragraph.add_run()
                        run.text = segment["glyph"]
                        self.style_run(run, 14, bold=False, color=("theme", 14), font="Segoe UI Emoji")
                    else:
                        self.add_runs(paragraph, [{"text": str(number)}], 14, bold=True, color=("theme", 14))
                label = slide.shapes.add_textbox(left + column_w * index + pad, row_top, column_w - 2 * pad, label_h)
                label.text_frame.word_wrap = True
                label.text_frame.vertical_anchor = 4  # bottom
                label.text_frame.paragraphs[0].alignment = 2
                self.add_runs(label.text_frame.paragraphs[0], segment["label"], size + 2, bold=True, color=accent)
                if segment["lines"]:
                    box = slide.shapes.add_textbox(left + column_w * index + pad, line_y + diameter // 2 + int(self.height * 0.015), column_w - 2 * pad, description_h)
                    box.text_frame.word_wrap = True
                    self.write_lines(box.text_frame, segment["lines"], size)
        self.note_layout(slide, "timeline", count)
        return True

    def render_flow(self, slide, intro, segments):
        """Chevron process steps with optional icons above and descriptions below."""
        from pptx.enum.shapes import MSO_SHAPE
        left, top, width, height, used = self.intro_geometry(intro)
        count = len(segments)
        rows = 1 if count <= 5 else 2
        per_row = math.ceil(count / rows)
        row_h = height // rows
        icon_h = int(self.height * 0.08) if self.has_icons(segments) else 0
        overlap = int(self.width * 0.006)
        step_w = (width + overlap * (per_row - 1)) // per_row
        chevron_h = max(int(self.height * 0.11), min(int(self.height * 0.17), int(row_h * 0.32)))
        pad = int(step_w * 0.08)
        point = int(chevron_h * 0.5)  # chevron tips eat horizontal room for text
        label_pt = (step_w - 2 * point - pad) / EMU_PER_PT
        self.text_width_pt = (step_w - 2 * pad) / EMU_PER_PT - 4
        description_h = row_h - icon_h - chevron_h - int(self.height * 0.04)
        size = None
        for candidate in (16, 15, 14):
            labels_fit = all(char_units(runs_text(s["label"]), candidate) <= label_pt * 2 for s in segments)
            if labels_fit and all(self.estimate_height(s["lines"], candidate) <= description_h for s in segments):
                size = candidate
                break
        if size is None:
            return False
        self.draw_intro(slide, intro, used)
        accent = self.accent()
        deepest = max([self.estimate_height(s["lines"], size) for s in segments] + [0])
        block_h = icon_h + chevron_h + int(self.height * 0.02) + deepest
        shift = max(0, (row_h - block_h) // 3) if rows == 1 else 0
        for row in range(rows):
            items = segments[row * per_row:(row + 1) * per_row]
            row_top = top + row * row_h + shift
            for index, segment in enumerate(items):
                number = row * per_row + index + 1
                x = left + index * (step_w - overlap)
                y = row_top
                if icon_h:
                    self.icon_badge(slide, segment, x + step_w // 2, y, icon_h)
                    y += icon_h + int(self.height * 0.012)
                shape = slide.shapes.add_shape(MSO_SHAPE.PENTAGON if index == 0 else MSO_SHAPE.CHEVRON, x, y, step_w, chevron_h)
                shape.name = f"Flow step {number}"
                self.plain_shape(shape, fill=accent)
                frame = self.centered_text(shape, 0)
                frame.margin_left = point if index else pad
                frame.margin_right = point
                self.add_runs(frame.paragraphs[0], segment["label"], size, bold=True, color=("theme", 14))
                if segment["lines"]:
                    box = slide.shapes.add_textbox(x + pad, y + chevron_h + int(self.height * 0.02), step_w - 2 * pad, description_h)
                    box.text_frame.word_wrap = True
                    self.write_lines(box.text_frame, segment["lines"], size)
                    self.center_paragraphs(box.text_frame)
        self.note_layout(slide, "flow", count)
        return True

    def render_layers(self, slide, intro, segments, pyramid=False):
        """Stacked bands (architecture layers) or a centred pyramid; list items inside a band become small boxes."""
        from pptx.enum.shapes import MSO_SHAPE
        left, top, width, height, used = self.intro_geometry(intro)
        count = len(segments)
        gap = int(self.height * 0.015)
        band_h = (height - gap * (count - 1)) // count
        label_w = 0 if pyramid else int(width * 0.16)
        pad = int(self.width * 0.01)
        size, plans = None, []
        for candidate in (16, 15, 14):
            plans = []
            for index, segment in enumerate(segments):
                band_w = int(width * (0.45 + 0.55 * index / max(1, count - 1))) if pyramid else width
                body_w = band_w - label_w - 2 * pad
                items = [line for line in segment["lines"] if line[2] in ("bullet", "number")]
                texts = [line for line in segment["lines"] if line[2] not in ("bullet", "number")]
                label_pt = ((band_w * 0.8) if pyramid else label_w) / EMU_PER_PT - 10
                if (items and texts) or char_units(runs_text(segment["label"]), candidate + 2) > label_pt * 2:
                    break
                if items and not pyramid:
                    per_row = min(len(items), 6)
                    rows = math.ceil(len(items) / per_row)
                    box_w = (body_w - pad * (per_row - 1)) // per_row
                    box_h = (band_h - 2 * pad - pad * (rows - 1)) // rows
                    self.text_width_pt = box_w / EMU_PER_PT - 10
                    if rows > 2 or any(self.estimate_height([(line[0], 0, "text", False)], candidate) > box_h for line in items):
                        break
                    plans.append(("items", items, per_row, box_w, box_h, band_w))
                else:
                    lines = [(line[0], 0, "text", False) for line in items] if items else texts
                    self.text_width_pt = body_w / EMU_PER_PT - 10
                    budget = band_h - 2 * pad - (self.estimate_height([(segment["label"], 0, "subheading", True)], candidate) if pyramid else 0)
                    if self.estimate_height(lines, candidate) > budget:
                        break
                    plans.append(("text", lines, band_w))
            else:
                size = candidate
                break
        if size is None:
            return False
        self.draw_intro(slide, intro, used)
        accent = self.accent()
        tints = ("EAF1FB", "F4F7FB")
        for index, (segment, plan) in enumerate(zip(segments, plans)):
            band_w = plan[-1]
            x, y = left + (width - band_w) // 2, top + index * (band_h + gap)
            band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, band_w, band_h)
            band.adjustments[0] = 0.08
            band.name = ("Pyramid " if pyramid else "Layer ") + runs_text(segment["label"])[:32]
            self.plain_shape(band, fill=("rgb", tints[index % 2]), line=("rgb", "D6DEE8"))
            if pyramid:
                frame = self.centered_text(band, pad)
                self.add_runs(frame.paragraphs[0], segment["label"], size + 2, bold=True, color=accent)
                self.write_lines(frame, plan[1], size, first=False)
                self.center_paragraphs(frame)
                continue
            label = slide.shapes.add_textbox(x + pad, y, label_w, band_h)
            self.add_runs(self.centered_text(label, pad).paragraphs[0], segment["label"], size + 2, bold=True, color=accent)
            body_x = x + label_w + pad
            if plan[0] == "items":
                _, items, per_row, box_w, box_h, _ = plan
                for position, line in enumerate(items):
                    row, column = divmod(position, per_row)
                    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, body_x + column * (box_w + pad), y + pad + row * (box_h + pad), box_w, box_h)
                    box.adjustments[0] = 0.12
                    box.name = "Layer item"
                    self.plain_shape(box, fill=("rgb", "FFFFFF"), line=("rgb", "C9D6E6"))
                    self.add_runs(self.centered_text(box, int(pad * 0.6)).paragraphs[0], line[0], size)
            else:
                box = slide.shapes.add_textbox(body_x, y + pad, band_w - label_w - 2 * pad, band_h - 2 * pad)
                box.text_frame.word_wrap = True
                box.text_frame.vertical_anchor = 3
                self.write_lines(box.text_frame, plan[1], size)
        self.note_layout(slide, "pyramid" if pyramid else "layers", count)
        return True

    def render_pyramid(self, slide, intro, segments):
        return self.render_layers(slide, intro, segments, pyramid=True)

    def render_cycle(self, slide, intro, segments):
        """Nodes on an ellipse joined by arrows from each node to the next."""
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        left, top, width, height, used = self.intro_geometry(intro)
        count = len(segments)
        node_w = int(width * (0.26 if count <= 4 else 0.21))
        node_h = int(height * (0.3 if count <= 4 else 0.24))
        self.text_width_pt = node_w / EMU_PER_PT - 14
        size = None
        for candidate in (16, 15, 14):
            if all(self.estimate_height([(s["label"], 0, "subheading", True)] + s["lines"], candidate) <= node_h - int(self.height * 0.02) for s in segments):
                size = candidate
                break
        if size is None:
            return False
        self.draw_intro(slide, intro, used)
        accent = self.accent()
        cx, cy = left + width // 2, top + height // 2
        rx, ry = (width - node_w) // 2, (height - node_h) // 2
        centers = []
        for index in range(count):
            angle = -math.pi / 2 + 2 * math.pi * index / count
            centers.append((int(cx + rx * math.cos(angle)), int(cy + ry * math.sin(angle))))

        def edge(origin, target):
            dx, dy = target[0] - origin[0], target[1] - origin[1]
            length = math.hypot(dx, dy) or 1
            ux, uy = dx / length, dy / length
            reach = min(node_w / 2 / abs(ux) if ux else float("inf"), node_h / 2 / abs(uy) if uy else float("inf")) * 1.12
            return int(origin[0] + ux * reach), int(origin[1] + uy * reach)

        for index, origin in enumerate(centers):
            target = centers[(index + 1) % count]
            start, end = edge(origin, target), edge(target, origin)
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, start[0], start[1], end[0], end[1])
            connector.name = f"Cycle arrow {index + 1}"
            self.plain_shape(connector, line=accent, line_pt=2)
            self.arrow_head(connector)
        for index, (segment, (x, y)) in enumerate(zip(segments, centers)):
            node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x - node_w // 2, y - node_h // 2, node_w, node_h)
            node.adjustments[0] = 0.15
            node.name = f"Cycle node {index + 1}"
            self.plain_shape(node, fill=("rgb", "F5F8FC"), line=accent, line_pt=1.25)
            frame = self.centered_text(node, int(self.width * 0.008))
            paragraph = frame.paragraphs[0]
            if segment.get("glyph"):
                run = paragraph.add_run()
                run.text = segment["glyph"] + " "
                self.style_run(run, size + 2, bold=False, color=accent, font="Segoe UI Emoji")
            self.add_runs(paragraph, segment["label"], size + 2, bold=True, color=accent)
            self.write_lines(frame, segment["lines"], size, first=False)
            self.center_paragraphs(frame)
        self.note_layout(slide, "cycle", count)
        return True

    def render_stats(self, slide, intro, segments):
        """Big figures with a caption under each, side by side."""
        left, top, width, height, used = self.intro_geometry(intro)
        count = len(segments)
        column_w = width // count
        pad = int(column_w * 0.06)
        number_h = int(self.height * 0.17)
        self.text_width_pt = (column_w - 2 * pad) / EMU_PER_PT - 4
        captions = []
        for segment in segments:
            lines = list(segment["lines"])
            if lines and lines[0][2] == "text":
                lines[0] = (lines[0][0], 0, "subheading", True)
            captions.append(lines)
        size = None
        for candidate in (16, 15, 14):
            if all(self.estimate_height(lines, candidate) <= height - number_h - int(self.height * 0.06) for lines in captions):
                size = candidate
                break
        if size is None:
            return False
        self.draw_intro(slide, intro, used)
        accent = self.accent()
        deepest = max([self.estimate_height(lines, size) for lines in captions] + [0])
        shift = max(0, (height - number_h - deepest) // 3)
        for index, (segment, lines) in enumerate(zip(segments, captions)):
            x = left + index * column_w
            figure = slide.shapes.add_textbox(x + pad, top + shift, column_w - 2 * pad, number_h)
            figure.name = "Stat " + runs_text(segment["label"])[:32]
            frame = figure.text_frame
            frame.word_wrap = True
            frame.vertical_anchor = 4
            frame.paragraphs[0].alignment = 2
            number_size = 44 if count <= 3 else 36
            # Bold figures run wider than the average estimate; keep a margin so a number never breaks across lines.
            while char_units(runs_text(segment["label"]), number_size) > self.text_width_pt * 0.82 and number_size > 20:
                number_size -= 2
            if segment.get("glyph"):
                run = frame.paragraphs[0].add_run()
                run.text = segment["glyph"] + " "
                self.style_run(run, number_size, bold=False, color=accent, font="Segoe UI Emoji")
            self.add_runs(frame.paragraphs[0], segment["label"], number_size, bold=True, color=accent)
            if lines:
                box = slide.shapes.add_textbox(x + pad, top + shift + number_h + int(self.height * 0.01), column_w - 2 * pad, deepest + int(self.height * 0.02))
                box.text_frame.word_wrap = True
                self.write_lines(box.text_frame, lines, size)
                self.center_paragraphs(box.text_frame)
        self.note_layout(slide, "stats", count)
        return True


    def page_of(self, slide):
        return list(self.presentation.slides).index(slide) + 1

    def add_table(self, slide, block, left, top, width, height):
        """Draw as many rows as fit; returns a continuation block carrying the fixed column widths, or None."""
        from pptx.util import Pt
        header, rows, align = block["header"], block["rows"], block.get("align", [])
        columns = max(len(header), *(len(r) for r in rows)) if rows else len(header)
        columns = max(columns, 1)
        size = 12 if columns <= 4 else 10
        # Column widths are computed once from the whole table and reused by every continuation slide, so column
        # edges do not jump between pages; each row's height then follows its tallest wrapped cell.
        widths = block.get("widths")
        if not widths or len(widths) != columns:
            weights = []
            for column in range(columns):
                cells = [header[column] if column < len(header) else []] + [r[column] if column < len(r) else [] for r in rows]
                weights.append(max(4.0, min(24.0, max(char_units(runs_text(c), 1) for c in cells))))
            total = sum(weights)
            widths = [int(width * weight / total) for weight in weights]

        def row_height(cells):
            lines = 1
            for column, runs in enumerate(cells[:columns]):
                available = max(widths[column] / EMU_PER_PT - 14.4, 8.0)
                lines = max(lines, sum(max(1, math.ceil(char_units(piece, size) / available)) for piece in runs_text(runs).split("\n")))
            return int((lines * size * 1.2 + 12) * EMU_PER_PT)

        heights = [row_height(header)] + [row_height(r) for r in rows]
        count, used = 0, heights[0]
        while count < len(rows) and used + heights[count + 1] <= height:
            used += heights[count + 1]
            count += 1
        if count == 0 and rows:
            count = 1
            self.attention.append({"slide": slide, "reason": "表格单行文字过多，超出页面，请精简单元格内容"})
        chunk, rest = rows[:count], rows[count:]
        shape = slide.shapes.add_table(len(chunk) + 1, columns, left, top, width, sum(heights[:count + 1]))
        table = shape.table
        for column, column_width in enumerate(widths):
            table.columns[column].width = column_width
        for index, row_h in enumerate(heights[:count + 1]):
            table.rows[index].height = row_h
        alignment = {"center": 2, "right": 3}
        def write(row_index, runs_list, is_header):
            for column in range(columns):
                cell = table.cell(row_index, column)
                frame = cell.text_frame
                frame.word_wrap = True
                paragraph = frame.paragraphs[0]
                if column < len(runs_list):
                    self.add_runs(paragraph, runs_list[column], size, bold=True if is_header else None,
                                  color=("theme", 14) if is_header else None)  # BACKGROUND_1 on accent fill
                if column < len(align) and align[column] in alignment:
                    paragraph.alignment = alignment[align[column]]
                cell.margin_top = cell.margin_bottom = Pt(3)
                if is_header:
                    cell.fill.solid()
                    cell.fill.fore_color.theme_color = 5  # ACCENT_1
        write(0, header, True)
        for index, row in enumerate(chunk, 1):
            write(index, row, False)
        return {**block, "rows": rest, "widths": widths} if rest else None

    def add_diagram(self, slide, text, left, top, width, height):
        """A ```mermaid block becomes native flowchart shapes in the given region; parse errors keep the source."""
        flowchart = flowchart_module()
        try:
            summary = flowchart.render_on_slide(slide, self.presentation, text, (left, top, width, height), self.accent(), self.body_font)
        except ValueError as error:
            self.warnings.append(f"流程图未生成，已按代码显示：{error}")
            lines = [([{"text": line or " ", "code": True}], 0, "text", False) for line in text.split("\n")]
            self.add_text_block(slide, lines, left, top, width, height)
            return
        self.warnings.extend(f"流程图：{w}" for w in summary["warnings"])
        self.layouts.append({"slide": slide, "mode": "flowchart", "segments": summary["nodes"]})
        self.attention.append({"slide": slide, "reason": f"已生成 {summary['nodes']} 节点流程图（{summary['direction']}，{summary['fontSize']}pt），请检查连线与文字"})

    def add_image(self, slide, path, alt, left, top, width, height):
        real = self.assets.get(path)
        if not real or Path(real).suffix.lower() not in IMAGE_TYPES:
            self.warnings.append(f"图片未插入（不存在或格式不支持）：{path}")
            return
        px_w, px_h = image_size(real)
        scale = min(width / px_w, height / px_h)
        draw_w, draw_h = int(px_w * scale), int(px_h * scale)
        picture = slide.shapes.add_picture(real, left + (width - draw_w) // 2, top + (height - draw_h) // 2, draw_w, draw_h)
        if alt:
            picture.name = alt[:64]

    # -- slide planning ---------------------------------------------------
    def build(self, request):
        keep = [p for p in request.get("sequence", []) if isinstance(p, int) and p > 0]
        pages = len(self.presentation.slides)
        for page in keep:
            if page > pages:
                raise ValueError(f"模板页码不存在：{page}")
        kept_entries = self.remove_slides(set(keep))
        title, subtitle = request.get("title"), request.get("subtitle")
        sections = self.split_sections(request["blocks"], title or "概述")
        produced = []
        if title and request.get("cover", True):
            slide = self.new_slide(self.cover or self.layout, keep=(1, 3, 4, 13, 15, 16))
            self.add_title(slide, title, cover=True, subtitle=subtitle)
            produced.append(slide)
        for section in sections:
            produced.extend(self.render_section(section))
        id_list = self.presentation.slides._sldIdLst
        new_entries = [entry for entry in id_list if entry not in kept_entries.values()]
        sequence = []
        inserted = False
        for item in request.get("sequence", []):
            if isinstance(item, int) and item > 0:
                sequence.append(kept_entries[item])
            elif not inserted:
                sequence.extend(new_entries)
                inserted = True
        if not inserted:
            sequence.extend(new_entries)
        self.reorder(sequence)
        self.presentation.save(request["output"])
        generated = [self.page_of(s) for s in produced]
        attention = [{"page": self.page_of(item["slide"]), "reason": item["reason"]} for item in self.attention]
        layouts = [{"page": self.page_of(item["slide"]), "mode": item["mode"], "segments": item["segments"]} for item in self.layouts]
        return {"generatedPages": generated, "keptPages": [self.page_of_entry(kept_entries[p]) for p in keep],
                "attention": attention, "layouts": layouts, "titleStyle": self.title_style.describe(), "layout": self.layout.name}

    def page_of_entry(self, entry):
        return list(self.presentation.slides._sldIdLst).index(entry) + 1

    def split_sections(self, blocks, default_title):
        sections, current = [], None
        for block in blocks:
            if block["type"] == "heading" and block["level"] <= 2:
                current = {"title": runs_text(block["runs"]), "blocks": [], "notes": []}
                sections.append(current)
                continue
            if current is None:
                current = {"title": default_title, "blocks": [], "notes": []}
                sections.append(current)
            if block["type"] == "note":
                current["notes"].append(block["text"])
            else:
                current["blocks"].append(block)
        return [s for s in sections if s["blocks"] or s["notes"] or len(sections) == 1]

    def render_section(self, section):
        """One section may span several slides: text overflow, tables and images each continue."""
        slides = []
        queue = list(section["blocks"])
        continuation = 0
        while queue or not slides:
            slide = self.new_slide()
            slides.append(slide)
            self.add_title(slide, section["title"] + ("（续）" if continuation else ""))
            if continuation:
                self.attention.append({"slide": slide, "reason": f"“{section['title']}”内容较多，已自动续页"})
            continuation += 1
            queue = self.fill_slide(slide, queue)
            if section["notes"] and len(slides) == 1:
                slide.notes_slide.notes_text_frame.text = "\n".join(section["notes"])
        return slides

    def fill_slide(self, slide, queue):
        left, top, width, height = self.region
        group = []
        for block in queue:
            if block["type"] == "pagebreak":
                break
            group.append(block)
        mode, plan = self.choose_layout(slide, group)
        if plan is not None:
            renderers = {"cards": self.render_cards, "timeline": self.render_timeline, "flow": self.render_flow, "layers": self.render_layers,
                         "pyramid": self.render_pyramid, "cycle": self.render_cycle, "stats": self.render_stats}
            if renderers[mode](slide, *plan):
                return queue[len(group) + 1:]
            self.attention.append({"slide": slide, "reason": LAYOUT_NAMES[mode] + "内容过长，已按普通版式排版"})
        elif mode in LAYOUT_RANGES:
            low, high = LAYOUT_RANGES[mode]
            self.warnings.append(f"“{mode}”指令所在页需要 {low}–{high} 个分段（### 小标题、列表项或箭头串）且不含表格、代码，已按普通版式排版")
        text_lines = []
        image = None
        table = None
        diagram = None
        consumed = 0
        for block in queue:
            kind = block["type"]
            if kind == "pagebreak":
                consumed += 1
                break
            if kind == "layout":
                consumed += 1
                continue
            if kind == "code" and is_diagram(block):
                if diagram or table or image:
                    break
                diagram = block
                consumed += 1
                continue
            if kind == "image":
                if image or table or diagram:
                    break
                image = block
            elif kind == "table":
                if table or image or diagram:
                    break
                table = block
            elif kind == "heading":
                text_lines.append((block["runs"], 0, "subheading", True))
            elif kind == "paragraph":
                text_lines.append((block["runs"], 0, "text", False))
            elif kind == "list":
                for item in block["items"]:
                    text_lines.append((item["runs"], item["level"], "number" if item["ordered"] else "bullet", False))
            elif kind == "quote":
                for runs in block["paragraphs"]:
                    text_lines.append((runs, 0, "quote", False))
            elif kind == "code":
                for line in block["text"].split("\n"):
                    text_lines.append(([{"text": line or " ", "code": True}], 0, "text", False))
            consumed += 1
        rest = queue[consumed:]
        if diagram:
            text_height = 0
            overflow = []
            if text_lines:
                text_height = min(int(height * 0.3), self.estimate_text_height(text_lines, width))
                overflow = self.add_text_block(slide, text_lines, left, top, width, text_height)
                text_height += int(self.height * 0.02)
            self.add_diagram(slide, diagram["text"], left, top + text_height, width, height - text_height)
            if overflow:
                rest = self.lines_to_blocks(overflow) + rest
            return rest
        if image and text_lines:
            real = self.assets.get(image["path"])
            pixels = image_size(real) if real and Path(real).suffix.lower() in IMAGE_TYPES else None
            needed = self.estimate_text_height(text_lines, width)
            gap = int(self.height * 0.02)
            if pixels and pixels[0] >= pixels[1] * 1.5 and needed <= int(height * 0.35):
                # A wide figure with a short caption-like text is stacked so the figure keeps the full width.
                overflow = self.add_text_block(slide, text_lines, left, top, width, needed)
                self.add_image(slide, image["path"], image.get("alt", ""), left, top + needed + gap, width, height - needed - gap)
            else:
                text_width = int(width * 0.52)
                overflow = self.add_text_block(slide, text_lines, left, top, text_width, height)
                self.add_image(slide, image["path"], image.get("alt", ""), left + text_width + int(width * 0.04), top, width - text_width - int(width * 0.04), height)
        elif image:
            self.add_image(slide, image["path"], image.get("alt", ""), left, top, width, height)
            overflow = []
        elif table and text_lines:
            text_height = min(int(height * 0.35), self.estimate_text_height(text_lines, width))
            overflow = self.add_text_block(slide, text_lines, left, top, width, text_height)
            rest_table = self.add_table(slide, table, left, top + text_height + int(self.height * 0.02), width, height - text_height - int(self.height * 0.02))
            if rest_table:
                rest = [rest_table] + rest
        elif table:
            rest_table = self.add_table(slide, table, left, top, width, height)
            overflow = []
            if rest_table:
                rest = [rest_table] + rest
        elif text_lines:
            overflow = self.add_text_block(slide, text_lines, left, top, width, height)
        else:
            overflow = []
        if overflow:
            rest = self.lines_to_blocks(overflow) + rest
        return rest

    def estimate_text_height(self, lines, width):
        self.text_width_pt = width / EMU_PER_PT - 14
        return self.estimate_height(lines, 16) + int(0.2 * EMU_PER_INCH)

    def lines_to_blocks(self, lines):
        blocks = []
        for runs, level, kind, _ in lines:
            if kind in ("bullet", "number"):
                if blocks and blocks[-1]["type"] == "list":
                    blocks[-1]["items"].append({"runs": runs, "level": level, "ordered": kind == "number"})
                else:
                    blocks.append({"type": "list", "items": [{"runs": runs, "level": level, "ordered": kind == "number"}]})
            elif kind == "quote":
                blocks.append({"type": "quote", "paragraphs": [runs]})
            elif kind == "subheading":
                blocks.append({"type": "heading", "level": 3, "runs": runs})
            else:
                blocks.append({"type": "paragraph", "runs": runs})
        return blocks


def build(request):
    warnings = []
    if request["format"] == "docx":
        result = WordBuilder(request, warnings).build(request)
    else:
        result = SlideBuilder(request, warnings).build(request)
    result["warnings"] = warnings
    return result


# ---------------------------------------------------------------- outlines

def slide_outline(presentation):
    width, height = presentation.slide_width, presentation.slide_height
    outline = []
    for index, slide in enumerate(presentation.slides, 1):
        title = None
        best = None
        chars = 0
        pictures = tables = charts = groups = 0
        for shape in slide.shapes:
            kind = shape.shape_type
            if kind == 13:
                pictures += 1
            elif kind == 6:
                groups += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                tables += 1
            if getattr(shape, "has_chart", False) and shape.has_chart:
                charts += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                text = shape.text_frame.text.strip()
                chars += len(text)
                if shape.top is not None and shape.top <= height * 0.22 and shape.width >= width * 0.25:
                    _, run = first_run(shape)
                    size = run.font.size.pt if run is not None and run.font.size else 0
                    is_title = shape.is_placeholder and shape.placeholder_format.type in (1, 3)
                    score = (1 if is_title else 0, size, -shape.top)
                    if best is None or score > best[0]:
                        best = (score, text.split("\n")[0][:80])
        if best:
            title = best[1]
        outline.append({"page": index, "layout": slide.slide_layout.name, "title": title, "chars": chars,
                        "pictures": pictures, "tables": tables, "charts": charts, "groups": groups,
                        "hidden": slide._element.get("show") in ("0", "false")})
    layouts = [{"name": l.name, "placeholders": [str(p.placeholder_format.type).split(" ")[0].replace("PP_PLACEHOLDER.", "") for p in l.placeholders]}
               for m in presentation.slide_masters for l in m.slide_layouts]
    return {"slides": outline, "layouts": layouts, "titleStyle": infer_title_style(presentation).describe(),
            "size": {"widthInches": round(width / EMU_PER_INCH, 2), "heightInches": round(height / EMU_PER_INCH, 2)}}


def word_outline(document):
    headings = []
    for index, block in enumerate(document.element.body, 1):
        if block.tag != "{%s}p" % NS_W:
            continue
        style_node = block.find(".//{%s}pStyle" % NS_W)
        style_id = style_node.get("{%s}val" % NS_W) if style_node is not None else ""
        match = re.fullmatch(r"(?:Heading|heading|标题)\s*(\d?)", style_id or "")
        if not match and style_id not in ("Title", "Subtitle"):
            continue
        text = "".join(block.xpath(".//w:t/text()")).strip()
        if not text:
            continue
        level = int(match.group(1)) if match and match.group(1) else 0
        headings.append({"block": index, "level": level, "text": text[:120]})
    names = {s.name for s in document.styles}
    styles = [n for n in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Table Grid", "Caption", "Quote") if n in names]
    section = document.sections[-1]
    return {"headings": headings[:400], "tables": len(document.tables), "inlineImages": len(document.inline_shapes),
            "styles": styles, "page": {"widthCm": round(section.page_width / 360000, 1), "heightCm": round(section.page_height / 360000, 1)},
            "header": "".join(p.text for p in section.header.paragraphs).strip()[:80] if section.header else ""}


def outline(path, kind):
    if kind == "pptx":
        from pptx import Presentation
        return slide_outline(Presentation(path))
    from docx import Document
    return word_outline(Document(path))
