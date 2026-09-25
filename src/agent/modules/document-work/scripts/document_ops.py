"""Local Office operations. Requests and outputs are scoped by the module's tools.ts."""
import hashlib
import json
import math
import os
import posixpath
import re
import shutil
import subprocess
import sys
import zipfile
from contextlib import closing
from pathlib import Path
from urllib.parse import unquote

from lxml import etree

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
MAX_FILE = 128 * 1024 * 1024


def xml(data):
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    root = etree.fromstring(data, parser)
    if root.getroottree().docinfo.doctype:
        raise ValueError("不支持包含 DTD 的 Office XML")
    return root


def digest(path):
    with open(path, "rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def relationship_target(name, target):
    target = unquote(target.split("#")[0])
    owner_dir = posixpath.dirname(posixpath.dirname(name))
    if not target:
        return posixpath.join(owner_dir, posixpath.basename(name)[:-5])
    return posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(owner_dir, target))


def package(path, *, prune_unused_tags=False):
    if Path(path).stat().st_size > MAX_FILE:
        raise ValueError("原文件超过 128 MiB")
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 20000 or sum(i.file_size for i in infos) > 256 * 1024 * 1024:
            raise ValueError("Office 解压大小或条目数超过上限")
        names = [i.filename for i in infos]
        if len(names) != len(set(names)) or any(n.startswith("/") or ".." in n.split("/") or "\\" in n for n in names):
            raise ValueError("Office 包含重复或无效部件路径")
        data = {i.filename: archive.read(i) for i in infos if not i.is_dir()}
    if "[Content_Types].xml" not in data or "_rels/.rels" not in data:
        raise ValueError("缺少 Office 包元数据")
    removed_tags = 0
    if prune_unused_tags:
        # pptx-automizer removes p:custDataLst but, with cleanup disabled,
        # leaves its tag relationships behind without copying their parts.
        # Remove only these now-unused metadata relationships. Referenced tags
        # and all other missing targets must still fail normal validation.
        for name, content in list(data.items()):
            if not name.endswith(".rels"):
                continue
            owner = posixpath.join(posixpath.dirname(posixpath.dirname(name)), posixpath.basename(name)[:-5])
            if owner not in data or not owner.endswith(".xml"):
                continue
            used = {value for node in xml(data[owner]).iter() for key, value in node.attrib.items()
                    if key.startswith("{" + NS["r"] + "}")}
            relationships = xml(content)
            stale = [node for node in relationships if node.get("Type") == NS["r"] + "/tags"
                     and node.get("TargetMode") != "External" and node.get("Id") not in used]
            for node in stale:
                relationships.remove(node)
            if stale:
                data[name] = etree.tostring(relationships, encoding="UTF-8", xml_declaration=True)
                removed_tags += len(stale)
    # Check internal targets without fetching external resources or extracting archives.
    external = 0
    for name, content in data.items():
        if name.endswith(".xml") or name.endswith(".rels"):
            root = xml(content)
            if not name.endswith(".rels"):
                continue
            for relationship in root:
                if relationship.get("TargetMode") == "External":
                    external += 1
                    continue
                resolved = relationship_target(name, relationship.get("Target", ""))
                if resolved not in data:
                    raise ValueError("Office 内部引用缺失：" + name + " -> " + resolved)
    warnings = [f"包含 {external} 个外部关联，未访问或验证其内容"] if external else []
    if removed_tags:
        warnings.append(f"已清理 {removed_tags} 个未被页面引用的自定义标签关联；组装器不保留这些非视觉元数据")
    return data, warnings


def paragraphs(root, kind):
    return root.findall(".//" + ("w:p" if kind == "docx" else "a:p"), NS)


def paragraph_nodes(paragraph):
    # Text boxes can contain their own paragraphs inside a Word paragraph.
    # Those have separate locations in the inspection and must not be edited twice.
    pending = [paragraph]
    while pending:
        node = pending.pop()
        if node is not paragraph and node.tag in ("{" + NS["w"] + "}p", "{" + NS["a"] + "}p"):
            continue
        yield node
        pending.extend(reversed(node))


def visible_text(paragraph):
    pieces = []
    for node in paragraph_nodes(paragraph):
        if node.tag in ("{" + NS["w"] + "}t", "{" + NS["a"] + "}t"):
            pieces.append(node.text or "")
        elif node.tag == "{" + NS["w"] + "}tab":
            pieces.append("\t")
        elif node.tag in ("{" + NS["w"] + "}br", "{" + NS["w"] + "}cr", "{" + NS["a"] + "}br"):
            pieces.append("\n")
    return "".join(pieces)


def editable_part(name, kind):
    if kind == "docx":
        return re.fullmatch(r"word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml", name)
    return re.fullmatch(r"ppt/(slides/slide\d+|notesSlides/notesSlide\d+)\.xml", name)


def inspect(path, outline=False):
    kind = Path(path).suffix.lower().lstrip(".")
    if kind not in ("docx", "pptx"):
        raise ValueError("结构检查支持 DOCX/PPTX")
    data, warnings = package(path)
    result = {"format": kind, "digest": digest(path), "warnings": warnings, "paragraphs": [], "truncated": False}
    chars = 0
    for part in sorted(data):
        if not editable_part(part, kind):
            continue
        for index, paragraph in enumerate(paragraphs(xml(data[part]), kind), 1):
            text = visible_text(paragraph)
            if chars + len(text) > 500000 or len(result["paragraphs"]) >= 15000:
                result["truncated"] = True
                break
            result["paragraphs"].append({"part": part, "paragraph": index, "text": text})
            chars += len(text)
    if kind == "docx":
        body = xml(data["word/document.xml"]).find("w:body", NS)
        blocks = [n for n in body if n.tag != "{" + NS["w"] + "}sectPr"]
        result["blocks"] = [{"block": i, "type": etree.QName(n).localname,
                             "text": " ".join(n.xpath(".//w:t/text()", namespaces=NS))[:400]}
                            for i, n in enumerate(blocks, 1)]
    else:
        presentation = xml(data["ppt/presentation.xml"])
        rels = {r.get("Id"): r.get("Target") for r in xml(data["ppt/_rels/presentation.xml.rels"])}
        size = presentation.find("p:sldSz", NS)
        result["size"] = {"width": int(size.get("cx")), "height": int(size.get("cy"))}
        result["slides"] = []
        for i, item in enumerate(presentation.findall("p:sldIdLst/p:sldId", NS), 1):
            target = rels[item.get("{" + NS["r"] + "}id")]
            part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("ppt", target))
            match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", part)
            if not match:
                raise ValueError("不支持非标准幻灯片部件名：" + part)
            result["slides"].append({"page": i, "part": part, "slideFile": int(match.group(1)),
                                     "hidden": xml(data[part]).get("show") in ("0", "false")})
    if outline:
        import document_build
        result["outline"] = document_build.outline(path, kind)
    return result


def replace_text(paragraph, before, after):
    if not before or any(c in before + after for c in "\r\n\t"):
        raise ValueError("替换要求非空原文字，且不跨换行或制表符；结构修改请使用专门脚本")
    if paragraph.xpath(".//w:fldChar | .//w:fldSimple | .//w:instrText | .//w:ins | .//w:del | .//a:fld | ancestor::w:ins | ancestor::w:del | ancestor::w:fldSimple", namespaces=NS):
        raise ValueError("该段含域或修订记录，需专门处理")
    text = visible_text(paragraph)
    if text.count(before) != 1:
        raise ValueError("原文字必须在目标段落中恰好出现一次；请重新检查文件并精确定位")
    start, end = text.index(before), text.index(before) + len(before)
    position, inserted = 0, False
    for node in paragraph_nodes(paragraph):
        if node.tag in ("{" + NS["w"] + "}t", "{" + NS["a"] + "}t"):
            value = node.text or ""
            stop = position + len(value)
            if position < end and stop > start:
                left, right = max(0, start - position), min(len(value), end - position)
                node.text = value[:left] + (after if not inserted else "") + value[right:]
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                inserted = True
            position = stop
        elif node.tag in ("{" + NS["w"] + "}tab", "{" + NS["w"] + "}br", "{" + NS["w"] + "}cr", "{" + NS["a"] + "}br"):
            position += 1
    if not inserted:
        raise ValueError("未找到可修改文字")


def patch(request):
    source, output = request["source"], request["output"]
    if digest(source) != request["digest"]:
        raise ValueError("原文件已变化，请重新 document_inspect")
    kind = Path(source).suffix.lower().lstrip(".")
    data, _ = package(source)
    changed = {}
    for edit in request["edits"]:
        part = edit["part"]
        if not editable_part(part, kind) or part not in data:
            raise ValueError("不支持该编辑部件")
        root = changed.setdefault(part, xml(data[part]))
        items = paragraphs(root, kind)
        number = edit["paragraph"]
        if not isinstance(number, int) or not 1 <= number <= len(items):
            raise ValueError("段落位置不存在")
        replace_text(items[number - 1], edit["before"], edit["after"])
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "x", compression=zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            payload = etree.tostring(changed[item.filename], encoding="UTF-8", xml_declaration=True, standalone=True) if item.filename in changed else original.read(item)
            target.writestr(item, payload)
    result = inspect(output)
    result["changedParts"] = list(changed)
    return result


def compose_word(request):
    from docx import Document
    from docxcompose.composer import Composer
    composer = None
    for item in request["items"]:
        package(item["source"])
        document = Document(item["source"])
        body = document.element.body
        blocks = [n for n in body if n.tag != "{" + NS["w"] + "}sectPr"]
        start, end = item.get("start", 1), item.get("end", len(blocks))
        if not blocks or not 1 <= start <= end <= len(blocks):
            raise ValueError("Word 章节范围不存在")
        for index, block in enumerate(blocks, 1):
            if not start <= index <= end:
                body.remove(block)
        if composer is None:
            composer = Composer(document)
        else:
            composer.append(document)
    composer.save(request["output"])
    return inspect(request["output"])


def finalize_slides(request):
    """Keep the assembled deck's reachable parts; preserve chart/media relations."""
    data, warnings = package(request["source"], prune_unused_tags=True)
    presentation = xml(data["ppt/presentation.xml"])
    active_ids = {n.get("{" + NS["r"] + "}id") for n in presentation.findall("p:sldIdLst/p:sldId", NS)}
    relname = "ppt/_rels/presentation.xml.rels"
    relationships = xml(data[relname])
    active_parts = set()
    for node in list(relationships):
        if node.get("Type", "").endswith("/slide"):
            if node.get("Id") in active_ids:
                active_parts.add(relationship_target(relname, node.get("Target")))
            else:
                relationships.remove(node)
    data[relname] = etree.tostring(relationships, encoding="UTF-8", xml_declaration=True)
    keep, pending = {"[Content_Types].xml"}, ["_rels/.rels"]
    while pending:
        part = pending.pop()
        if part in keep:
            continue
        keep.add(part)
        if part.endswith(".rels"):
            for node in xml(data[part]):
                if node.get("TargetMode") != "External":
                    pending.append(relationship_target(part, node.get("Target", "")))
        else:
            rel = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
            if rel in data:
                pending.append(rel)
    extra_slides = {p for p in keep if re.fullmatch(r"ppt/slides/slide\d+\.xml", p)} - active_parts
    if extra_slides:
        raise ValueError("选中页面仍引用未选取页面；请处理跨页链接后重新组装")
    types = xml(data["[Content_Types].xml"])
    for node in list(types):
        if node.get("PartName") and unquote(node.get("PartName")).lstrip("/") not in keep:
            types.remove(node)
    data["[Content_Types].xml"] = etree.tostring(types, encoding="UTF-8", xml_declaration=True)
    with zipfile.ZipFile(request["output"], "x", compression=zipfile.ZIP_DEFLATED) as archive:
        for part in sorted(keep):
            archive.writestr(part, data[part])
    result = inspect(request["output"])
    result["warnings"] = list(dict.fromkeys(result["warnings"] + warnings))
    return result


def office_binary():
    candidates = [shutil.which("soffice"), shutil.which("libreoffice")]
    if os.name == "nt":
        for folder in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
            if folder:
                candidates.append(str(Path(folder) / "LibreOffice/program/soffice.exe"))
    return next((p for p in candidates if p and Path(p).is_file()), None)


def convert_office(source, folder):
    package(source)
    binary = office_binary()
    if not binary:
        raise ValueError("缺少 LibreOffice，无法渲染 Office 预览；安装后将 soffice 加入 PATH，Windows 也支持标准安装目录。文件编辑能力仍可用，不能声称已完成视觉检查。")
    profile = folder / "office-profile"
    (profile / "user").mkdir(parents=True)
    # Dedicated profile isolates concurrent conversions and disables macro execution.
    (profile / "user/registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>', encoding="utf-8")
    env = {**os.environ}
    for key, name in (("HOME", "office-home"), ("XDG_CONFIG_HOME", "office-config"),
                      ("XDG_CACHE_HOME", "office-cache"), ("XDG_DATA_HOME", "office-data")):
        path = folder / name
        path.mkdir()
        env[key] = str(path)
    env.update(TMPDIR=str(folder), TMP=str(folder), TEMP=str(folder))
    # Unix IPC still requires writable /tmp (TMPDIR does not redirect it).
    # Keep stderr/stdout: CalledProcessError alone hides the actual Office failure.
    result = subprocess.run([binary, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo", "--nodefault",
                             "--norestore", "--convert-to", "pdf", "--outdir", str(folder), str(source)],
                            timeout=150, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env,
                            **({"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {}))
    diagnostic = (result.stdout + b"\n" + result.stderr).decode("utf-8", errors="replace").strip()[-1400:]
    if result.returncode != 0:
        raise ValueError(f"LibreOffice 转换失败（退出码 {result.returncode}）：{diagnostic or '未返回诊断输出'}")
    pdf = folder / (source.stem + ".pdf")
    if not pdf.is_file():
        raise ValueError("LibreOffice 没有生成 PDF，视觉检查未完成：" + (diagnostic or "未返回诊断输出"))
    return pdf


def render(request):
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    source = Path(request["source"])
    folder = Path(request["directory"])
    pdf = source if source.suffix.lower() == ".pdf" else convert_office(source, Path(request.get("workdir", folder)))
    images, contacts = [], []
    with closing(pdfium.PdfDocument(str(pdf))) as document:
        total = len(document)
        selected = request.get("pages") or list(range(1, min(total, 20) + 1))
        if not selected or len(selected) > 50 or any(not isinstance(i, int) or not 1 <= i <= total for i in selected):
            raise ValueError("预览页码越界，单次最多 50 页")
        selected = list(dict.fromkeys(selected))
        thumbs = []
        for number in selected:
            with closing(document[number - 1]) as page:
                scale = min(1.5, 1800 / max(page.get_size()))
                with closing(page.render(scale=scale)) as bitmap:
                    picture = bitmap.to_pil().convert("RGB")
                    path = folder / f"page-{number:04}.png"
                    picture.save(path)
                    images.append({"page": number, "path": str(path)})
                    picture.thumbnail((440, 300))
                    tile = Image.new("RGB", (460, 335), "#eef1f5")
                    tile.paste(picture, ((460 - picture.width) // 2, 25))
                    ImageDraw.Draw(tile).text((10, 7), str(number), fill="black")
                    thumbs.append(tile)
                    picture.close()
        for start in range(0, len(thumbs), 12):
            batch = thumbs[start:start + 12]
            contact = Image.new("RGB", (460 * min(3, len(batch)), 335 * math.ceil(len(batch) / 3)), "white")
            for offset, thumb in enumerate(batch):
                contact.paste(thumb, ((offset % 3) * 460, (offset // 3) * 335))
                thumb.close()
            path = folder / f"contact-{start // 12 + 1}.jpg"
            contact.save(path, quality=88)
            contact.close()
            contacts.append(str(path))
    return {"pages": total, "images": images, "contacts": contacts,
            "unrenderedPages": [i for i in range(1, total + 1) if i not in selected], "visuallyReviewed": False}


def build(request):
    import document_build
    for path in [request.get("template"), *request.get("assets", {}).values()]:
        if path and Path(path).suffix.lower() in (".docx", ".pptx"):
            package(path)
    result = document_build.build(request)
    inspection = inspect(request["output"])
    inspection["build"] = result
    inspection["warnings"] = list(dict.fromkeys(inspection["warnings"] + result.pop("warnings", [])))
    return inspection


RASTER_TYPES = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}
MAX_IMAGE_EDGE = 2000


def images(request):
    """Extract embedded raster images (and optional cropped page regions) as reusable assets."""
    import hashlib
    import io
    from PIL import Image
    source = Path(request["source"])
    folder = Path(request["directory"]) / "images"
    folder.mkdir()
    kind = source.suffix.lower()
    min_side = max(16, int(request.get("minSize") or 96))
    crops = request.get("crops") or []
    if len(crops) > 20:
        raise ValueError("单次最多截取 20 个区域")
    results, warnings, unsupported = [], [], []
    skipped = {"small": 0, "repeated": 0, "unsupported": 0}
    seen = {}

    def clamp(value):
        return round(min(1.0, max(0.0, value)), 4)

    def record(picture, page, box, name, extra=None):
        # Pixel-size and page-share filters apply to every source; icons drawn at a few percent of the page
        # are noise for reuse even when their pixel size is large.
        if min(picture.size) < min_side or (box and (box[2] - box[0]) < 0.03 and (box[3] - box[1]) < 0.05):
            skipped["small"] += 1
            picture.close()
            return
        # Exact content match: a thumbnail signature would merge different shapes on transparent backgrounds.
        signature = (picture.size, hashlib.sha1(picture.convert("RGBA").tobytes()).hexdigest())
        if signature in seen:
            entry = seen[signature]
            if page not in entry["pages"]:
                entry["pages"].append(page)
            skipped["repeated"] += 1
            return
        if len(results) >= 60:
            picture.close()
            return
        if max(picture.size) > MAX_IMAGE_EDGE:
            picture.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
        path = folder / name
        if picture.mode not in ("RGB", "RGBA", "L"):
            picture = picture.convert("RGBA" if "A" in picture.mode or "transparency" in picture.info else "RGB")
        picture.save(path)
        entry = {"page": page, "pages": [page], "path": str(path), "width": picture.width, "height": picture.height, "box": box,
                 "fullPage": bool(box and (box[2] - box[0]) * (box[3] - box[1]) >= 0.85), **(extra or {})}
        results.append(entry)
        seen[signature] = entry
        picture.close()

    def page_list(total, cap):
        selected = request.get("pages")
        if not selected:
            selected = list(range(1, min(total, cap) + 1))
            if total > cap:
                warnings.append(f"共 {total} 页，默认只处理前 {cap} 页；其余页用 pages 指定，每次最多 50 页")
        if len(selected) > 50 or any(not isinstance(i, int) or not 1 <= i <= total for i in selected):
            raise ValueError("页码越界，单次最多 50 页")
        return list(dict.fromkeys(selected))

    total = 0
    if kind == ".pdf":
        import pypdfium2 as pdfium
        from pypdfium2 import raw
        with closing(pdfium.PdfDocument(str(source))) as document:
            total = len(document)
            selected = page_list(total, 20)
            for number in selected:
                with closing(document[number - 1]) as page:
                    width, height = page.get_size()
                    for index, obj in enumerate(page.get_objects(filter=(raw.FPDF_PAGEOBJ_IMAGE,), max_depth=4), 1):
                        px_w, px_h = obj.get_px_size()
                        if min(px_w, px_h) < min_side:
                            skipped["small"] += 1
                            continue
                        left, bottom, right, top = obj.get_bounds()
                        box = [clamp(left / width), clamp(1 - top / height), clamp(right / width), clamp(1 - bottom / height)]
                        try:
                            bitmap = obj.get_bitmap(render=True)
                        except Exception as error:  # noqa: BLE001 - unusual colour spaces or broken streams
                            skipped["unsupported"] += 1
                            warnings.append(f"第 {number} 页第 {index} 张图片无法解码：{type(error).__name__}")
                            continue
                        with closing(bitmap):
                            record(bitmap.to_pil(), number, box, f"p{number:03}-{index:02}.png")
    else:
        package(source)
        if kind == ".pptx":
            from pptx import Presentation
            presentation = Presentation(str(source))
            total = len(presentation.slides)
            selected = page_list(total, 50)
            width, height = presentation.slide_width, presentation.slide_height

            def walk(shapes):
                for shape in shapes:
                    if shape.shape_type == 6:
                        yield from walk(shape.shapes)
                    elif shape.shape_type == 13:
                        yield shape

            for number in selected:
                for index, shape in enumerate(walk(presentation.slides[number - 1].shapes), 1):
                    image = shape.image
                    extension = image.ext.lower()
                    if extension not in RASTER_TYPES:
                        skipped["unsupported"] += 1
                        unsupported.append({"page": number, "name": shape.name, "format": extension})
                        continue
                    picture = Image.open(io.BytesIO(image.blob))
                    picture.load()
                    crop = (shape.crop_left or 0, shape.crop_top or 0, shape.crop_right or 0, shape.crop_bottom or 0)
                    if any(value > 0 for value in crop):
                        w, h = picture.size
                        picture = picture.crop((int(w * crop[0]), int(h * crop[1]), int(w * (1 - crop[2])), int(h * (1 - crop[3]))))
                    box = None
                    if shape.left is not None and shape.width:
                        box = [clamp(shape.left / width), clamp(shape.top / height), clamp((shape.left + shape.width) / width), clamp((shape.top + shape.height) / height)]
                    record(picture, number, box, f"s{number:03}-{index:02}.{'png' if picture.mode in ('RGBA', 'P', 'LA') else 'jpg' if extension in ('jpg', 'jpeg') else 'png'}", {"name": shape.name})
        else:
            from docx import Document
            document = Document(str(source))
            total = 0
            selected = []
            # Walk blips in body order so anchored (floating) pictures are included, not only inline ones.
            index = 0
            for blip in document.element.body.iter("{%s}blip" % NS["a"]):
                embed = blip.get("{%s}embed" % NS["r"])
                part = document.part.related_parts.get(embed) if embed else None
                if part is None:
                    continue
                index += 1
                extension = Path(str(part.partname)).suffix.lstrip(".").lower()
                if extension not in RASTER_TYPES:
                    skipped["unsupported"] += 1
                    unsupported.append({"index": index, "format": extension})
                    continue
                picture = Image.open(io.BytesIO(part.blob))
                picture.load()
                record(picture, index, None, f"w{index:03}.{'jpg' if extension in ('jpg', 'jpeg') else 'png'}")
    crop_results = []
    if crops:
        import pypdfium2 as pdfium
        pdf = source if kind == ".pdf" else convert_office(source, Path(request.get("workdir", request["directory"])))
        with closing(pdfium.PdfDocument(str(pdf))) as document:
            for index, crop in enumerate(crops, 1):
                number, box = crop.get("page"), crop.get("box")
                if not isinstance(number, int) or not 1 <= number <= len(document) or not isinstance(box, list) or len(box) != 4:
                    raise ValueError("crops 需要 page 与 box=[x0,y0,x1,y1]（页面比例，左上角为原点）")
                x0, y0, x1, y1 = (min(1.0, max(0.0, float(v))) for v in box)
                if x1 - x0 < 0.02 or y1 - y0 < 0.02:
                    raise ValueError("截取区域太小")
                with closing(document[number - 1]) as page:
                    width, height = page.get_size()
                    scale = min(3.0, 2400 / max(width, height))
                    with closing(page.render(scale=scale, crop=(x0 * width, (1 - y1) * height, (1 - x1) * width, y0 * height))) as bitmap:
                        picture = bitmap.to_pil().convert("RGB")
                        path = folder / f"crop-p{number:03}-{index:02}.png"
                        picture.save(path)
                        crop_results.append({"page": number, "box": [x0, y0, x1, y1], "path": str(path), "width": picture.width, "height": picture.height})
                        picture.close()
    if unsupported:
        warnings.append("跳过了 %d 张矢量或不支持格式的图片（EMF/WMF/SVG 等），可用 crops 从渲染页面截取" % len(unsupported))
    results.sort(key=lambda item: (item["page"], -(item["width"] * item["height"])))
    return {"pages": total, "selectedPages": selected, "images": results, "crops": crop_results, "skipped": skipped,
            "unsupported": unsupported[:40], "warnings": warnings, "directory": str(folder)}


def main():
    operation, request_path, result_path = sys.argv[1:]
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    request = json.loads(Path(request_path).read_text(encoding="utf-8"))
    handlers = {"inspect": lambda r: inspect(r["source"], r.get("outline", False)), "patch": patch, "compose_word": compose_word,
                "finalize_slides": finalize_slides, "render": render, "build": build, "images": images}
    result = handlers[operation](request)
    Path(result_path).write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
