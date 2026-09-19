"""Branching flowcharts from a Mermaid subset: deterministic layered layout, native PowerPoint shapes or a PNG.

Run in the fixed document environment (python-pptx, python-docx and Pillow are available):

  uv run --no-project --python "$PI_PYTHON" flowchart.py flow.mmd --pptx 成稿.pptx --slide 5 --out 成稿-调整.pptx
  uv run --no-project --python "$PI_PYTHON" flowchart.py flow.mmd --png flow.png
  uv run --no-project --python "$PI_PYTHON" flowchart.py flow.mmd --docx 成稿.docx --after 处理流程 --out 成稿-调整.docx
  uv run --no-project --python "$PI_PYTHON" flowchart.py flow.mmd --check

document_build.py imports this module for ```mermaid code blocks inside Markdown content.
"""
import argparse
import json
import math
import os
import re
import sys
import unicodedata
from pathlib import Path

EMU_PER_PT = 12700
EMU_PER_INCH = 914400
DIAGRAM_LANGS = {"mermaid", "flowchart", "flow", "graph"}
MAX_NODES, MAX_EDGES = 40, 80
DARK, WHITE = "1F2937", "FFFFFF"
WHITE_RGB = (255, 255, 255)

# ---------------------------------------------------------------- parsing (Mermaid flowchart subset)

SHAPES = [("([", "])", "terminator"), ("[[", "]]", "subroutine"), ("[(", ")]", "database"), ("((", "))", "circle"),
          ("{{", "}}", "hexagon"), ("[/", "/]", "data"), ("[\\", "\\]", "data"), ("[/", "\\]", "data"), ("[\\", "/]", "data"),
          ("[", "]", "process"), ("(", ")", "rounded"), ("{", "}", "decision"), (">", "]", "flag")]
IDENT = re.compile(r"[A-Za-z0-9_\u00c0-\u024f\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uac00-\ud7af·]+")
EDGE_MID = re.compile(r"(?P<open>--|==|-\.)\s*(?P<label>[^\s\-=.>|][^>]*?)\s*(?P<close>-->|---|--|==>|===|==|\.->|\.-|->|=>)")
EDGE = re.compile(r"(?P<line>-{2,}|={2,}|-\.+-|-|=)(?P<head>>)?(?:\s*\|(?P<label>[^|]*)\|)?")
HEADER = re.compile(r"^(?:graph|flowchart)\b\s*(?P<dir>TB|TD|BT|LR|RL)?\s*$", re.I)
IGNORED = re.compile(r"^(?:subgraph\b|end\b|style\b|classDef\b|class\b|click\b|linkStyle\b|%%\{)", re.I)
TERMINAL = re.compile(r"^(开始|结束|完成|流程开始|流程结束|start|end|begin|stop|finish|done)$", re.I)


class Graph:
    def __init__(self):
        self.nodes, self.order, self.edges, self.warnings = {}, [], [], []
        self.direction = None

    def node(self, ident, label=None, kind=None):
        node = self.nodes.get(ident)
        if node is None:
            node = {"id": ident, "label": ident if label is None else label, "kind": kind or "process", "explicit": kind is not None}
            self.nodes[ident] = node
            self.order.append(ident)
        else:
            if label is not None:
                node["label"] = label
            if kind:
                node["kind"], node["explicit"] = kind, True
        return node


def parse_node(graph, line, pos):
    while pos < len(line) and line[pos] == " ":
        pos += 1
    match = IDENT.match(line, pos)
    if not match:
        raise ValueError(f"期望节点，得到“{line[pos:pos + 12]}”")
    ident, pos = match.group(), match.end()
    label = kind = None
    for opener, closer, shape in SHAPES:
        if line.startswith(opener, pos):
            start = pos + len(opener)
            if line.startswith('"', start):
                quote = line.find('"', start + 1)
                if quote < 0 or not line.startswith(closer, quote + 1):
                    raise ValueError(f"节点 {ident} 的引号或括号不匹配")
                label, pos = line[start + 1:quote], quote + 1 + len(closer)
            else:
                end = line.find(closer, start)
                if end < 0:
                    raise ValueError(f"节点 {ident} 缺少 {closer}")
                label, pos = line[start:end], end + len(closer)
            kind = shape
            break
    if label is not None:
        label = re.sub(r"<br\s*/?>", "\n", label).strip().strip('"').strip()
    graph.node(ident, label, kind)
    return ident, pos


def parse_group(graph, line, pos):
    ids = []
    ident, pos = parse_node(graph, line, pos)
    ids.append(ident)
    while True:
        cursor = pos
        while cursor < len(line) and line[cursor] == " ":
            cursor += 1
        if cursor < len(line) and line[cursor] == "&":
            ident, pos = parse_node(graph, line, cursor + 1)
            ids.append(ident)
        else:
            return ids, cursor


def parse_edge(line, pos):
    match = EDGE_MID.match(line, pos)
    if match:
        token, label = match.group("open") + match.group("close"), match.group("label")
    else:
        match = EDGE.match(line, pos)
        if not match or match.group("line") in ("-", "=") and not match.group("head"):
            raise ValueError(f"期望连线（如 -->），得到“{line[pos:pos + 12]}”")
        token, label = match.group("line") + (match.group("head") or ""), match.group("label") or ""
    return {"label": label.strip().strip('"'), "head": ">" in token, "dashed": "." in token, "thick": "=" in token}, match.end()


def parse_chain(graph, line):
    left, pos = parse_group(graph, line, 0)
    while pos < len(line):
        edge, pos = parse_edge(line, pos)
        right, pos = parse_group(graph, line, pos)
        for src in left:
            for dst in right:
                graph.edges.append({"src": src, "dst": dst, **edge})
        left = right


def parse(text):
    graph = Graph()
    for number, raw in enumerate(text.splitlines(), 1):
        line = raw.strip().rstrip(";").strip()
        if not line or line.startswith("%%"):
            continue
        header = HEADER.match(line)
        if header:
            if header.group("dir"):
                graph.direction = header.group("dir").upper().replace("TD", "TB")
            continue
        if IGNORED.match(line):
            if line.lower().startswith("subgraph"):
                graph.warnings.append("已忽略 subgraph 分组，节点按同一张图处理")
            continue
        try:
            parse_chain(graph, line)
        except ValueError as error:
            raise ValueError(f"第 {number} 行无法解析：{error}") from None
    if not graph.nodes:
        raise ValueError("流程图没有节点")
    if len(graph.nodes) > MAX_NODES or len(graph.edges) > MAX_EDGES:
        raise ValueError(f"流程图过大（最多 {MAX_NODES} 个节点、{MAX_EDGES} 条边），请拆成多张")
    incoming = {e["dst"] for e in graph.edges}
    outgoing = {e["src"] for e in graph.edges}
    for node in graph.nodes.values():
        if not node["explicit"] and TERMINAL.match(node["label"].strip()) and (node["id"] not in incoming or node["id"] not in outgoing):
            node["kind"] = "terminator"
    return graph


# ---------------------------------------------------------------- text metrics

def char_units(text, size):
    """Approximate rendered width in points: CJK ≈ 1em, Latin ≈ 0.55em."""
    width = 0.0
    for char in text:
        if unicodedata.east_asian_width(char) in ("W", "F"):
            width += size
        elif char == " ":
            width += size * 0.3
        else:
            width += size * 0.55
    return width


def wrap_lines(text, limit, measure):
    """Greedy wrap keeping Latin words whole; explicit newlines are kept."""
    lines = []
    for paragraph in text.split("\n"):
        current = ""
        for token in re.findall(r"[A-Za-z0-9_\-/.%+]+|\s+|.", paragraph):
            if current and measure(current + token.rstrip()) > limit:
                lines.append(current.rstrip())
                current = token.lstrip()
                while current and measure(current) > limit:
                    cut = max(1, next((i for i in range(len(current), 0, -1) if measure(current[:i]) <= limit), 1))
                    lines.append(current[:cut])
                    current = current[cut:]
            else:
                current += token
        lines.append(current.rstrip())
    return [line for line in lines if line] or [""]


def size_node(node, size, vertical=True):
    """Node box in points; everything scales with the font size so fitting is a pure zoom."""
    kind = node["kind"]
    em = size
    limits = {"decision": 8.5 * em, "circle": 6.5 * em, "terminator": 11 * em}
    limit = limits.get(kind, (12 if vertical else 9) * em)
    lines = wrap_lines(node["label"], limit, lambda s: char_units(s, size))
    text_w = max(char_units(line, size) for line in lines)
    text_h = len(lines) * size * 1.25
    node["lines"] = lines
    if kind == "decision":
        w, h = max(7.5 * em, text_w * 2 + 1.4 * em), max(4 * em, text_h * 2 + 0.6 * em)
    elif kind == "circle":
        w = h = max(4.3 * em, max(text_w, text_h) + 1.7 * em)
    elif kind == "terminator":
        w, h = max(6.3 * em, text_w + 2.6 * em), max(2.4 * em, text_h + em)
    elif kind in ("data", "hexagon"):
        w, h = max(7 * em, text_w * 1.3 + 3.2 * em), max(2.6 * em, text_h + 1.3 * em)
    elif kind == "database":
        w, h = max(6 * em, text_w + 2 * em), max(3.6 * em, text_h + 2.2 * em)
    else:
        w, h = max(6 * em, text_w + 2 * em), max(2.6 * em, text_h + 1.3 * em)
    node["w"], node["h"] = w, h


# ---------------------------------------------------------------- layout

def dedupe(points):
    result = []
    for point in points:
        if not result or abs(point[0] - result[-1][0]) > 0.01 or abs(point[1] - result[-1][1]) > 0.01:
            result.append(point)
    cleaned = []
    for point in result:  # drop collinear middle points
        if len(cleaned) >= 2 and ((abs(cleaned[-2][0] - cleaned[-1][0]) < 0.01 and abs(cleaned[-1][0] - point[0]) < 0.01)
                                  or (abs(cleaned[-2][1] - cleaned[-1][1]) < 0.01 and abs(cleaned[-1][1] - point[1]) < 0.01)):
            cleaned[-1] = point
        else:
            cleaned.append(point)
    return cleaned


def layout(graph, direction, size, compact=False):
    """Positions in points. Main axis = flow direction; cross axis = siblings. Returns xy geometry."""
    vertical = direction in ("TB", "BT")
    gap_main = size * (2.1 if compact else 3.2)
    gap_cross = size * (1.7 if compact else 2.6)
    ids = list(graph.order)
    nodes = {i: dict(graph.nodes[i]) for i in ids}
    for node in nodes.values():
        size_node(node, size, vertical)
    edges = [dict(e) for e in graph.edges]

    def main_size(i):
        return nodes[i]["h"] if vertical else nodes[i]["w"]

    def cross_size(i):
        return nodes[i]["w"] if vertical else nodes[i]["h"]

    outgoing = {i: [] for i in ids}
    for edge in edges:
        outgoing[edge["src"]].append(edge)
    state = {}

    def visit(u):
        state[u] = 1
        for edge in outgoing[u]:
            v = edge["dst"]
            if v == u:
                edge["loop"] = True
            elif state.get(v) == 1:
                edge["back"] = True
            elif v not in state:
                visit(v)
        state[u] = 2

    incoming_count = {i: 0 for i in ids}
    for edge in edges:
        incoming_count[edge["dst"]] += 1
    for start in [i for i in ids if incoming_count[i] == 0] + ids:
        if start not in state:
            visit(start)
    forward = [e for e in edges if not e.get("back") and not e.get("loop")]
    layer = {i: 0 for i in ids}
    indeg = {i: 0 for i in ids}
    for edge in forward:
        indeg[edge["dst"]] += 1
    queue = [i for i in ids if indeg[i] == 0]
    while queue:
        u = queue.pop(0)
        for edge in forward:
            if edge["src"] == u:
                v = edge["dst"]
                layer[v] = max(layer[v], layer[u] + 1)
                indeg[v] -= 1
                if indeg[v] == 0:
                    queue.append(v)
    links, dummies = [], 0
    for edge in forward:
        chain = [edge["src"]]
        for level in range(layer[edge["src"]] + 1, layer[edge["dst"]]):
            dummies += 1
            ident = f"\x00{dummies}"
            nodes[ident] = {"id": ident, "dummy": True, "w": 8, "h": 8, "kind": "dummy", "label": "", "lines": []}
            layer[ident] = level
            ids.append(ident)
            chain.append(ident)
        chain.append(edge["dst"])
        edge["chain"] = chain
        links.extend(zip(chain, chain[1:]))
    depth = max(layer.values()) + 1
    layers = [[] for _ in range(depth)]
    for i in ids:
        layers[layer[i]].append(i)
    pos = {i: k for row in layers for k, i in enumerate(row)}
    preds, succs = {i: [] for i in ids}, {i: [] for i in ids}
    for a, b in links:
        succs[a].append(b)
        preds[b].append(a)

    def resort(row, neighbours):
        row.sort(key=lambda i: (sum(pos[n] for n in neighbours[i]) / len(neighbours[i]) if neighbours[i] else pos[i], pos[i]))
        for k, i in enumerate(row):
            pos[i] = k

    for _ in range(4):
        for row in layers[1:]:
            resort(row, preds)
        for row in reversed(layers[:-1]):
            resort(row, succs)
    cross = {}
    for row in layers:
        total = sum(cross_size(i) for i in row) + gap_cross * (len(row) - 1)
        cursor = -total / 2
        for i in row:
            cross[i] = cursor + cross_size(i) / 2
            cursor += cross_size(i) + gap_cross
    # Straighten bypass edges: a dummy on the outside of its row may move to the chain's outermost lane.
    for edge in forward:
        chain = edge["chain"][1:-1]
        if not chain:
            continue
        outer = max((cross[d] for d in chain), key=abs)
        for d in chain:
            row = layers[layer[d]]
            if (row[0] == d and outer < cross[d]) or (row[-1] == d and outer > cross[d]):
                cross[d] = outer
    main = {}
    cursor = 0.0
    for index, row in enumerate(layers):
        extent = max(main_size(i) for i in row)
        for i in row:
            main[i] = cursor + extent / 2
        labelled = any(e["label"] and layer[e["src"]] == index and nodes[e["src"]]["kind"] != "decision" for e in forward)
        cursor += extent + gap_main + (size * 1.1 if labelled else 0)

    def top(i):
        return main[i] - main_size(i) / 2

    def bottom(i):
        return main[i] + main_size(i) / 2

    used_sides = {}
    for edge in forward:
        chain, u = edge["chain"], edge["src"]
        first = chain[1]
        points = []
        span = (min(cross[u], cross[first]), max(cross[u], cross[first]))
        blocked = any(cross[i] - cross_size(i) / 2 < span[1] and cross[i] + cross_size(i) / 2 > span[0]
                      for i in layers[layer[u]] if i != u and not nodes[i].get("dummy"))
        if nodes[u]["kind"] == "decision" and abs(cross[first] - cross[u]) > cross_size(u) / 2 + 6 and not blocked:
            side = 1 if cross[first] > cross[u] else -1
            used_sides.setdefault(u, set()).add(side)
            points.append((main[u], cross[u] + side * cross_size(u) / 2))
            points.append((main[u], cross[first]))
            current = cross[first]
        else:
            points.append((bottom(u), cross[u]))
            current = cross[u]
        for a, b in zip(chain, chain[1:]):
            if abs(cross[b] - current) > 0.5:
                mid = (bottom(a) + top(b)) / 2
                points.append((mid, current))
                points.append((mid, cross[b]))
                current = cross[b]
        points.append((top(chain[-1]), cross[chain[-1]]))
        edge["points"] = dedupe(points)
    real = [i for i in ids if not nodes[i].get("dummy")]
    extent_left = min(cross[i] - cross_size(i) / 2 for i in real)
    extent_right = max(cross[i] + cross_size(i) / 2 for i in real)
    lanes = {1: 0, -1: 0}
    for edge in edges:
        u, v = edge["src"], edge["dst"]
        if edge.get("loop"):
            side = -1 if 1 in used_sides.get(u, ()) else 1
            edge_c = cross[u] + side * cross_size(u) / 2
            edge["points"] = [(main[u] - 8, edge_c), (main[u] - 8, edge_c + side * 22), (main[u] + 8, edge_c + side * 22), (main[u] + 8, edge_c)]
        elif edge.get("back"):
            side = 1 if (cross[u] + cross[v]) / 2 >= 0 else -1
            if side in used_sides.get(u, ()):
                side = -side
            lanes[side] += 1
            lane = (extent_right + 24 + lanes[side] * 16) if side > 0 else (extent_left - 24 - lanes[side] * 16)
            edge["points"] = [(main[u], cross[u] + side * cross_size(u) / 2), (main[u], lane), (main[v], lane), (main[v], cross[v] + side * cross_size(v) / 2)]
    flip = -1 if direction in ("BT", "RL") else 1

    def xy(m, c):
        return (c, m * flip) if vertical else (m * flip, c)

    out_nodes = []
    for i in real:
        x, y = xy(main[i], cross[i])
        out_nodes.append({**nodes[i], "x": x, "y": y})
    out_edges = []
    label_size = max(9, size - 2)
    for edge in edges:
        points = [xy(m, c) for m, c in edge["points"]]
        item = {"src": edge["src"], "dst": edge["dst"], "label": edge["label"], "head": edge["head"], "dashed": edge["dashed"],
                "thick": edge["thick"], "points": points, "back": bool(edge.get("back") or edge.get("loop"))}
        if edge["label"]:
            segments = list(zip(points, points[1:]))
            w, h = char_units(edge["label"], label_size) + 8, label_size * 1.35

            def horizontal(s):
                return abs(s[0][1] - s[1][1]) < 0.01 and abs(s[0][0] - s[1][0]) > 0.01

            def length(s):
                return math.hypot(s[1][0] - s[0][0], s[1][1] - s[0][1])

            # Prefer the first cross-direction segment (the jog next to the source) when it can hold the label;
            # a tiny jog would push the label into the next node, so fall back to the longest segment.
            preferred = [s for s in segments if horizontal(s) == vertical and length(s) >= (w if horizontal(s) else h) + 4]
            segment = preferred[0] if preferred else max(segments, key=length)
            mx, my = (segment[0][0] + segment[1][0]) / 2, (segment[0][1] + segment[1][1]) / 2
            if horizontal(segment):
                item["label_box"] = (mx - w / 2, my - h - 3, w, h)
            elif mx < 0:
                item["label_box"] = (mx - 5 - w, my - h / 2, w, h)
            else:
                item["label_box"] = (mx + 5, my - h / 2, w, h)
        out_edges.append(item)
    xs = [n["x"] - n["w"] / 2 for n in out_nodes] + [n["x"] + n["w"] / 2 for n in out_nodes]
    ys = [n["y"] - n["h"] / 2 for n in out_nodes] + [n["y"] + n["h"] / 2 for n in out_nodes]
    for edge in out_edges:
        xs.extend(p[0] for p in edge["points"])
        ys.extend(p[1] for p in edge["points"])
        if "label_box" in edge:
            bx, by, bw, bh = edge["label_box"]
            xs.extend((bx, bx + bw))
            ys.extend((by, by + bh))
    return {"direction": direction, "size": size, "label_size": label_size, "nodes": out_nodes, "edges": out_edges,
            "bbox": (min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)), "layers": depth, "warnings": list(graph.warnings)}


def plan(graph, region_w, region_h, direction="auto", base=14, minimum=10):
    """Pick direction and spacing so the drawing fits region (points). Geometry scales with the font, so fitting is a
    zoom: normal spacing first, compact spacing when the font would drop below `minimum`, then the other orientation
    if it is clearly better. A header direction is only overridden when the alternative is at least 15% larger."""
    if direction != "auto":
        candidates = [direction]
    elif graph.direction:
        candidates = [graph.direction, "LR" if graph.direction in ("TB", "BT") else "TB"]
    else:
        candidates = ["TB", "LR"]
    results = []
    for candidate in candidates:
        options = []
        for compact in (False, True):
            lay = layout(graph, candidate, base, compact)
            scale = min(1.0, region_w / lay["bbox"][2], region_h / lay["bbox"][3])
            options.append((base * scale, scale, lay))
            if scale >= 1 or options[-1][0] >= minimum:
                break
        results.append(max(options, key=lambda item: item[0]))
    chosen = results[0]
    if len(results) > 1 and results[1][0] > chosen[0] * 1.15 and chosen[0] < minimum + 1:
        chosen = results[1]
    effective, scale, lay = chosen
    lay["scale"] = scale
    if effective < minimum:
        other = "LR" if lay["direction"] in ("TB", "BT") else "TB"
        lay["warnings"].append(f"流程图超出可用区域，已整体缩小到 {int(scale * 100)}%，字号约 {effective:.0f}pt；建议拆成两张、精简节点文字或改用 {other} 方向")
    return lay


def place(lay, region):
    """Scale (if needed) and centre the layout inside region = (x, y, w, h) in points; mutates and returns lay."""
    rx, ry, rw, rh = region
    bx, by, bw, bh = lay["bbox"]
    scale = lay.get("scale", 1.0)
    ox = rx + (rw - bw * scale) / 2 - bx * scale
    oy = ry + (rh - bh * scale) / 2 - by * scale

    def move(x, y):
        return (ox + x * scale, oy + y * scale)

    for node in lay["nodes"]:
        node["x"], node["y"] = move(node["x"], node["y"])
        node["w"], node["h"] = node["w"] * scale, node["h"] * scale
    for edge in lay["edges"]:
        edge["points"] = [move(*p) for p in edge["points"]]
        if "label_box" in edge:
            x, y, w, h = edge["label_box"]
            x, y = move(x, y)
            edge["label_box"] = (x, y, w * scale, h * scale)
    lay["bbox"] = (*move(bx, by), bw * scale, bh * scale)
    lay["size"], lay["label_size"] = lay["size"] * scale, lay["label_size"] * scale
    lay["placed"] = True
    return lay


# ---------------------------------------------------------------- PowerPoint output

SHAPE_TYPES = {"process": "RECTANGLE", "rounded": "ROUNDED_RECTANGLE", "decision": "FLOWCHART_DECISION", "terminator": "FLOWCHART_TERMINATOR",
               "data": "FLOWCHART_DATA", "database": "FLOWCHART_MAGNETIC_DISK", "circle": "OVAL", "hexagon": "HEXAGON",
               "subroutine": "FLOWCHART_PREDEFINED_PROCESS", "flag": "FLOWCHART_OFF_PAGE_CONNECTOR"}
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def tint(hex_color, amount):
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{int(c + (255 - c) * amount):02X}" for c in (r, g, b))


def paint(color_format, color, brightness=0.0):
    from pptx.dml.color import RGBColor
    if color[0] == "theme":
        color_format.theme_color = color[1]
        if brightness:
            color_format.brightness = brightness
    else:
        color_format.rgb = RGBColor.from_string(tint(color[1], brightness) if brightness > 0 else color[1])


def strip_style(shape):
    style = shape._element.find("{%s}style" % NS_P)
    if style is not None:
        shape._element.remove(style)


def write_text(frame, text, size, color, font=None, bold=False, align_center=True):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    paragraph = frame.paragraphs[0]
    if align_center:
        paragraph.alignment = 2
    for index, piece in enumerate(text.split("\n")):
        if index:
            paragraph.add_line_break()
        if not piece:
            continue
        run = paragraph.add_run()
        run.text = piece
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if font:
            run.font.name = font


def draw_on_slide(slide, lay, accent=("theme", 5), font=None, group=True):
    """Draw a placed layout (points) as native shapes; returns the shapes created."""
    from lxml import etree
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Emu, Pt
    if not lay.get("placed"):
        raise ValueError("layout must be placed before drawing")
    size = lay["size"]
    created = []

    def emu(value):
        return Emu(int(value * EMU_PER_PT))

    for edge in lay["edges"]:
        points = [(emu(x), emu(y)) for x, y in edge["points"]]
        if len(points) == 2:
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, points[0][0], points[0][1], points[1][0], points[1][1])
        else:
            builder = slide.shapes.build_freeform(points[0][0], points[0][1], scale=1.0)
            builder.add_line_segments(points[1:], close=False)
            shape = builder.convert_to_shape()
            shape.fill.background()
        strip_style(shape)
        shape.name = f"Flow edge {edge['src']}->{edge['dst']}"
        shape.line.width = Pt(2.25 if edge["thick"] else 1.5)
        paint(shape.line.color, accent)
        line = shape.line._get_or_add_ln()
        if edge["dashed"]:
            dash = etree.SubElement(line, "{%s}prstDash" % NS_A)
            dash.set("val", "dash")
        if edge["head"]:
            tail = etree.SubElement(line, "{%s}tailEnd" % NS_A)
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")
        created.append(shape)
    for node in lay["nodes"]:
        kind = node["kind"]
        shape = slide.shapes.add_shape(getattr(MSO_SHAPE, SHAPE_TYPES.get(kind, "RECTANGLE")),
                                       emu(node["x"] - node["w"] / 2), emu(node["y"] - node["h"] / 2), emu(node["w"]), emu(node["h"]))
        strip_style(shape)
        shape.name = "Flow node " + node["id"]
        shape.fill.solid()
        if kind == "terminator":
            paint(shape.fill.fore_color, accent)
            shape.line.fill.background()
            text_color, bold = WHITE, True
        elif kind == "decision":
            paint(shape.fill.fore_color, accent, 0.72)
            paint(shape.line.color, accent)
            shape.line.width = Pt(1.25)
            text_color, bold = DARK, True
        else:
            paint(shape.fill.fore_color, accent, 0.88)
            paint(shape.line.color, accent)
            shape.line.width = Pt(0.75)
            text_color, bold = DARK, False
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = 3
        inset = Pt(2 if kind == "decision" else 4)
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = inset
        write_text(frame, node["label"], size, text_color, font, bold)
        created.append(shape)
    for edge in lay["edges"]:
        if "label_box" not in edge:
            continue
        x, y, w, h = edge["label_box"]
        box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        box.name = f"Flow label {edge['src']}->{edge['dst']}"
        frame = box.text_frame
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = 3
        write_text(frame, edge["label"], lay["label_size"], DARK, font)
        created.append(box)
    if group and len(created) > 1:
        try:
            grouped = slide.shapes.add_group_shape(created)
            grouped.name = "Flowchart"
        except Exception:  # noqa: BLE001 - grouping is cosmetic; older python-pptx lacks it
            pass
    return created


def text_bottom(shape):
    """Bottom of the text actually in a box (generated pages use full-height boxes for short intros); EMU."""
    size = 18.0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                size = run.font.size.pt
                break
        else:
            continue
        break
    available = max(40.0, shape.width / EMU_PER_PT - 14)
    lines = 0
    for paragraph in shape.text_frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs)
        lines += max(1, math.ceil(char_units(text, size) / available)) if text.strip() else 1
    estimate = int((lines * size * 1.3 + 10) * EMU_PER_PT)
    return shape.top + min(shape.height, estimate)


def slide_region(slide, presentation):
    """Below the title and any intro text in the upper part of the slide, with side margins; in EMU."""
    width, height = presentation.slide_width, presentation.slide_height
    top = int(height * 0.18)
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() and shape.top is not None and shape.top < height * 0.45:
            top = max(top, text_bottom(shape) + int(height * 0.02))
    left, right, bottom = int(width * 0.06), int(width * 0.94), int(height * 0.92)
    if bottom - top < height * 0.3:
        top = int(height * 0.18)
    return left, top, right - left, bottom - top


def render_on_slide(slide, presentation, text, region=None, accent=("theme", 5), font=None, direction="auto", base=14, minimum=10):
    """Parse, plan and draw. Region in EMU (defaults to below the title). Returns a summary dict."""
    graph = parse(text)
    left, top, width, height = region or slide_region(slide, presentation)
    lay = plan(graph, width / EMU_PER_PT, height / EMU_PER_PT, direction, base, minimum)
    place(lay, (left / EMU_PER_PT, top / EMU_PER_PT, width / EMU_PER_PT, height / EMU_PER_PT))
    draw_on_slide(slide, lay, accent, font)
    return {"nodes": len(lay["nodes"]), "edges": len(lay["edges"]), "direction": lay["direction"], "fontSize": round(lay["size"], 1),
            "scale": round(lay["scale"], 2), "layers": lay["layers"], "warnings": lay["warnings"]}


# ---------------------------------------------------------------- PNG output (Word and previews)

FONT_CANDIDATES = ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc", "C:/Windows/Fonts/simhei.ttf", "C:/Windows/Fonts/simsun.ttc",
                   "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
                   "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
                   "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", "/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc"]


def find_cjk_font():
    candidates = [os.environ.get("FLOWCHART_FONT")] + FONT_CANDIDATES
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    for folder in ("/usr/share/fonts", "/usr/local/share/fonts"):
        for pattern in ("**/*CJK*.tt[cf]", "**/*CJK*.otf", "**/*wqy*.tt[cf]"):
            found = next(Path(folder).glob(pattern), None) if Path(folder).is_dir() else None
            if found:
                return str(found)
    return None


def load_font(path, size_px):
    from PIL import ImageFont
    if not path:
        return ImageFont.load_default()
    index = 2 if "NotoSansCJK-" in Path(path).name else 0
    try:
        return ImageFont.truetype(path, size_px, index=index)
    except OSError:
        return ImageFont.truetype(path, size_px)


def render_png(lay, path, width_px=1600, font_path=None, accent="2F6FBA", margin=20):
    """Rasterise a (non-placed) layout at 2x and downsample; returns (width, height) in pixels."""
    from PIL import Image, ImageDraw
    if font_path is None and any(ord(ch) > 127 for node in lay["nodes"] for ch in node["label"]):
        raise LookupError("未找到中文字体，无法输出 PNG（PPT 输出不受影响）；可设置环境变量 FLOWCHART_FONT 指向字体文件")
    bx, by, bw, bh = lay["bbox"]
    bw, bh = bw + 2 * margin, bh + 2 * margin
    scale = width_px / bw * 2
    image = Image.new("RGB", (int(bw * scale), int(bh * scale)), WHITE_RGB)
    draw = ImageDraw.Draw(image)

    def rgb(hex_color):
        return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))

    def point(x, y):
        return ((x - bx + margin) * scale, (y - by + margin) * scale)

    accent_rgb, dark = rgb(accent), rgb(DARK)
    font = load_font(font_path, int(lay["size"] * scale))
    label_font = load_font(font_path, int(lay["label_size"] * scale))
    line_h = lay["size"] * scale * 1.25

    def dashed(points, width):
        on, off = 6 * scale, 4 * scale
        for (x1, y1), (x2, y2) in zip(points, points[1:]):
            length = math.hypot(x2 - x1, y2 - y1)
            steps = max(1, int(length / (on + off)))
            for step in range(steps + 1):
                start = step * (on + off)
                end = min(length, start + on)
                if start >= length:
                    break
                t0, t1 = start / length, end / length
                draw.line([(x1 + (x2 - x1) * t0, y1 + (y2 - y1) * t0), (x1 + (x2 - x1) * t1, y1 + (y2 - y1) * t1)], fill=accent_rgb, width=width)

    for edge in lay["edges"]:
        points = [point(x, y) for x, y in edge["points"]]
        width = int((2.25 if edge["thick"] else 1.5) * scale)
        if edge["dashed"]:
            dashed(points, width)
        else:
            draw.line(points, fill=accent_rgb, width=width, joint="curve")
        if edge["head"]:
            (x1, y1), (x2, y2) = points[-2], points[-1]
            angle = math.atan2(y2 - y1, x2 - x1)
            length, half = 10 * scale, 5 * scale
            base = (x2 - length * math.cos(angle), y2 - length * math.sin(angle))
            draw.polygon([(x2, y2), (base[0] + half * math.sin(angle), base[1] - half * math.cos(angle)),
                          (base[0] - half * math.sin(angle), base[1] + half * math.cos(angle))], fill=accent_rgb)
    for node in lay["nodes"]:
        kind = node["kind"]
        x0, y0 = point(node["x"] - node["w"] / 2, node["y"] - node["h"] / 2)
        x1, y1 = point(node["x"] + node["w"] / 2, node["y"] + node["h"] / 2)
        cx, cy = point(node["x"], node["y"])
        if kind == "terminator":
            fill, outline, text_color, line_w = accent_rgb, accent_rgb, WHITE_RGB, 1
        elif kind == "decision":
            fill, outline, text_color, line_w = rgb(tint(accent, 0.72)), accent_rgb, dark, int(1.25 * scale)
        else:
            fill, outline, text_color, line_w = rgb(tint(accent, 0.88)), accent_rgb, dark, max(1, int(0.75 * scale))
        if kind == "decision":
            draw.polygon([(cx, y0), (x1, cy), (cx, y1), (x0, cy)], fill=fill, outline=outline, width=line_w)
        elif kind == "terminator":
            draw.rounded_rectangle((x0, y0, x1, y1), radius=(y1 - y0) / 2, fill=fill, outline=outline, width=line_w)
        elif kind == "rounded":
            draw.rounded_rectangle((x0, y0, x1, y1), radius=6 * scale, fill=fill, outline=outline, width=line_w)
        elif kind == "circle":
            draw.ellipse((x0, y0, x1, y1), fill=fill, outline=outline, width=line_w)
        elif kind == "data":
            skew = (x1 - x0) * 0.16
            draw.polygon([(x0 + skew, y0), (x1, y0), (x1 - skew, y1), (x0, y1)], fill=fill, outline=outline, width=line_w)
        elif kind == "hexagon":
            skew = (x1 - x0) * 0.14
            draw.polygon([(x0 + skew, y0), (x1 - skew, y0), (x1, cy), (x1 - skew, y1), (x0 + skew, y1), (x0, cy)], fill=fill, outline=outline, width=line_w)
        elif kind == "database":
            cap = (y1 - y0) * 0.22
            draw.rectangle((x0, y0 + cap / 2, x1, y1 - cap / 2), fill=fill)
            draw.line([(x0, y0 + cap / 2), (x0, y1 - cap / 2)], fill=outline, width=line_w)
            draw.line([(x1, y0 + cap / 2), (x1, y1 - cap / 2)], fill=outline, width=line_w)
            draw.ellipse((x0, y1 - cap, x1, y1), fill=fill, outline=outline, width=line_w)
            draw.rectangle((x0 + line_w, y1 - cap, x1 - line_w, y1 - cap / 2), fill=fill)
            draw.ellipse((x0, y0, x1, y0 + cap), fill=fill, outline=outline, width=line_w)
        else:
            draw.rectangle((x0, y0, x1, y1), fill=fill, outline=outline, width=line_w)
            if kind == "subroutine":
                inset = 6 * scale
                draw.line([(x0 + inset, y0), (x0 + inset, y1)], fill=outline, width=line_w)
                draw.line([(x1 - inset, y0), (x1 - inset, y1)], fill=outline, width=line_w)
        limit = (x1 - x0) * (0.5 if kind == "decision" else 0.86)
        lines = wrap_lines(node["label"], limit, font.getlength)
        total = len(lines) * line_h
        for index, line in enumerate(lines):
            width = font.getlength(line)
            draw.text((cx - width / 2, cy - total / 2 + index * line_h), line, font=font, fill=text_color)
    for edge in lay["edges"]:
        if "label_box" not in edge:
            continue
        x, y, w, h = edge["label_box"]
        x0, y0 = point(x, y)
        x1, y1 = point(x + w, y + h)
        draw.rectangle((x0, y0, x1, y1), fill=WHITE_RGB)
        width = label_font.getlength(edge["label"])
        draw.text(((x0 + x1) / 2 - width / 2, y0 + (y1 - y0 - lay["label_size"] * scale * 1.2) / 2), edge["label"], font=label_font, fill=dark)
    image = image.resize((width_px, max(1, round(image.height * width_px / image.width))), Image.LANCZOS)
    image.save(path)
    result = image.size
    image.close()
    return result


def render_png_from_text(text, path, width_px=1600, direction=None, base=14, accent="2F6FBA", font_path=None):
    graph = parse(text)
    lay = layout(graph, direction or graph.direction or "TB", base)
    size = render_png(lay, path, width_px, font_path if font_path is not None else find_cjk_font(), accent)
    return {"nodes": len(lay["nodes"]), "edges": len(lay["edges"]), "direction": lay["direction"], "fontSize": lay["size"],
            "widthPt": lay["bbox"][2], "heightPt": lay["bbox"][3], "pixels": list(size), "warnings": lay["warnings"]}


# ---------------------------------------------------------------- Word insertion

def insert_into_docx(docx_in, docx_out, png_path, width_pt, after=None, caption=None, width_cm=None):
    from docx import Document
    from docx.shared import Cm, Emu
    document = Document(docx_in)
    section = document.sections[-1]
    content = section.page_width - section.left_margin - section.right_margin
    wanted = Cm(width_cm) if width_cm else Emu(int(width_pt * EMU_PER_PT))
    wanted = min(wanted, content)
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    paragraph.paragraph_format.keep_with_next = bool(caption)
    paragraph.add_run().add_picture(str(png_path), width=wanted)
    elements = [paragraph._p]
    if caption:
        names = {s.name for s in document.styles}
        note = document.add_paragraph(caption, style="Caption" if "Caption" in names else None)
        note.alignment = 1
        elements.append(note._p)
    if after:
        anchor = next((p for p in document.paragraphs if after in p.text), None)
        if anchor is None:
            raise ValueError(f"没有找到包含“{after}”的段落")
        for element in reversed(elements):
            anchor._p.addnext(element)
    document.save(docx_out)
    return {"after": after, "caption": caption, "widthCm": round(wanted / 360000, 1)}


# ---------------------------------------------------------------- CLI

def parse_region(value):
    parts = [float(v) for v in value.split(",")]
    if len(parts) != 4:
        raise argparse.ArgumentTypeError("region 需要 x,y,w,h（英寸）")
    return tuple(int(v * EMU_PER_INCH) for v in parts)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Mermaid 子集流程图：PPT 原生形状、PNG 或插入 Word")
    parser.add_argument("spec", help="Mermaid 流程图文本文件，- 表示标准输入")
    parser.add_argument("--pptx", help="在这份 PPT 的副本上绘制")
    parser.add_argument("--slide", type=int, help="目标页码（从 1 起）")
    parser.add_argument("--region", type=parse_region, help="绘图区域 x,y,w,h（英寸）；默认标题与引导文字之下")
    parser.add_argument("--png", help="输出 PNG 路径")
    parser.add_argument("--width", type=int, default=1600, help="PNG 宽度像素")
    parser.add_argument("--docx", help="把 PNG 插入这份 Word 的副本")
    parser.add_argument("--after", help="Word：插入到包含该文字的段落之后，默认文末")
    parser.add_argument("--caption", help="Word：图注文字")
    parser.add_argument("--width-cm", type=float, help="Word：图片宽度厘米")
    parser.add_argument("--out", help="输出文件路径（pptx/docx）")
    parser.add_argument("--direction", default="auto", choices=["auto", "TB", "LR", "BT", "RL"], help="PPT 默认 auto：TB 与 LR 中取更合适者")
    parser.add_argument("--size", type=int, default=14, help="起始字号，装不下时逐步减到 10")
    parser.add_argument("--accent", default="theme", help="强调色 RRGGBB，PPT 默认沿用主题强调色")
    parser.add_argument("--font", help="PPT 字体名；PNG/Word 用 --font-file")
    parser.add_argument("--font-file", help="PNG 字体文件路径（默认自动查找中文字体）")
    parser.add_argument("--check", action="store_true", help="只解析与布局，输出摘要")
    args = parser.parse_args(argv)
    text = sys.stdin.read() if args.spec == "-" else Path(args.spec).read_text(encoding="utf-8")
    accent_hex = None if args.accent == "theme" else args.accent.lstrip("#").upper()
    if accent_hex is not None and not re.fullmatch(r"[0-9A-F]{6}", accent_hex):
        parser.error("--accent 需要 RRGGBB")
    summary = {}
    try:
        if args.check or not (args.pptx or args.png or args.docx):
            graph = parse(text)
            lay = layout(graph, graph.direction or "TB", args.size)
            summary = {"nodes": [{"id": n["id"], "label": n["label"], "kind": n["kind"]} for n in lay["nodes"]],
                       "edges": [{"src": e["src"], "dst": e["dst"], "label": e["label"], "back": e["back"]} for e in lay["edges"]],
                       "direction": lay["direction"], "layers": lay["layers"], "widthPt": round(lay["bbox"][2]), "heightPt": round(lay["bbox"][3]),
                       "warnings": lay["warnings"]}
        if args.pptx:
            from pptx import Presentation
            if not args.out or not args.slide:
                parser.error("--pptx 需要 --slide 与 --out")
            presentation = Presentation(args.pptx)
            if not 1 <= args.slide <= len(presentation.slides):
                parser.error(f"页码越界，共 {len(presentation.slides)} 页")
            slide = presentation.slides[args.slide - 1]
            accent = ("theme", 5) if accent_hex is None else ("rgb", accent_hex)
            summary["pptx"] = render_on_slide(slide, presentation, text, args.region, accent, args.font, args.direction, args.size)
            presentation.save(args.out)
            summary["output"] = args.out
        if args.png or args.docx:
            png = args.png or str(Path(args.out or args.docx).with_suffix("")) + "-flow.png"
            direction = None if args.direction == "auto" else args.direction
            summary["png"] = render_png_from_text(text, png, args.width, direction, args.size, accent_hex or "2F6FBA", args.font_file)
            summary["png"]["path"] = png
        if args.docx:
            if not args.out:
                parser.error("--docx 需要 --out")
            summary["docx"] = insert_into_docx(args.docx, args.out, png, summary["png"]["widthPt"], args.after, args.caption, args.width_cm)
            summary["output"] = args.out
    except LookupError as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2
    except ValueError as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
