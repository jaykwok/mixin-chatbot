"""Real editable Office fixtures and assertions for the document tool integration."""
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Inches as WordInches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from pypdf import PdfWriter
from lxml import etree
from docx.oxml.ns import qn


def check_edit_guards():
    script = Path(__file__).resolve().parents[2] / "src/agent/modules/document-work/scripts/document_ops.py"
    spec = importlib.util.spec_from_file_location("document_ops", script)
    operations = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(operations)
    w, a = operations.NS["w"], operations.NS["a"]
    paragraph = operations.xml(f'<w:p xmlns:w="{w}"><w:r><w:t>外层</w:t></w:r><w:p><w:r><w:t>文本框</w:t></w:r></w:p></w:p>'.encode())
    assert operations.visible_text(paragraph) == "外层"
    operations.replace_text(paragraph, "外层", "修改")
    assert operations.visible_text(paragraph[1]) == "文本框"
    paragraph = operations.xml(f'<a:p xmlns:a="{a}"><a:pPr><a:tabLst><a:tab pos="100"/></a:tabLst></a:pPr><!--comment--><a:r><a:t>名称</a:t></a:r></a:p>'.encode())
    assert operations.visible_text(paragraph) == "名称"
    operations.replace_text(paragraph, "名称", "客户名称")
    assert operations.visible_text(paragraph) == "客户名称"
    paragraph = operations.xml(f'<w:p xmlns:w="{w}"><w:fldSimple w:instr="DATE"><w:r><w:t>日期</w:t></w:r></w:fldSimple></w:p>'.encode())
    try:
        operations.replace_text(paragraph, "日期", "不应替换")
        raise AssertionError("field edits must be rejected")
    except ValueError:
        assert operations.visible_text(paragraph) == "日期"


def add_tag_fixture(path):
    """A non-root slide's metadata reproduces automizer's dangling-tag output."""
    with zipfile.ZipFile(path) as archive:
        data = {n: archive.read(n) for n in archive.namelist()}
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct = "http://schemas.openxmlformats.org/package/2006/content-types"
    slide = etree.fromstring(data["ppt/slides/slide1.xml"])
    tags = etree.SubElement(slide.find("{" + p + "}cSld"), "{" + p + "}custDataLst")
    etree.SubElement(tags, "{" + p + "}tags", {"{" + r + "}id": "rIdTrialTag"})
    data["ppt/slides/slide1.xml"] = etree.tostring(slide)
    relationships = etree.fromstring(data["ppt/slides/_rels/slide1.xml.rels"])
    etree.SubElement(relationships, "{" + rel + "}Relationship", Id="rIdTrialTag", Type=r + "/tags", Target="../tags/tag777.xml")
    data["ppt/slides/_rels/slide1.xml.rels"] = etree.tostring(relationships)
    types = etree.fromstring(data["[Content_Types].xml"])
    etree.SubElement(types, "{" + ct + "}Override", PartName="/ppt/tags/tag777.xml", ContentType="application/vnd.openxmlformats-officedocument.presentationml.tags+xml")
    data["[Content_Types].xml"] = etree.tostring(types)
    data["ppt/tags/tag777.xml"] = f'<p:tagLst xmlns:p="{p}"><p:tag name="TRIAL" val="SOURCE-ONLY"/></p:tagLst>'.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in data.items():
            archive.writestr(name, content)
    # A still-referenced missing tag must never be silently dropped by cleanup.
    broken = path.with_name("broken-referenced-tag.pptx")
    with zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in data.items():
            if name != "ppt/tags/tag777.xml":
                archive.writestr(name, content)
    spec = importlib.util.spec_from_file_location("document_ops", Path(__file__).resolve().parents[2] / "src/agent/modules/document-work/scripts/document_ops.py")
    operations = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(operations)
    original, _ = operations.package(path, prune_unused_tags=True)
    assert b"rIdTrialTag" in original["ppt/slides/_rels/slide1.xml.rels"]
    try:
        operations.package(broken, prune_unused_tags=True)
        raise AssertionError("referenced missing tag must be rejected")
    except ValueError as error:
        assert "tag777.xml" in str(error)


def add_edge_fixtures(root):
    """Templates for the build edge cases: inherited headers across sections, slide links to dropped pages,
    look-alike images on transparent backgrounds and a deck above the per-call page limit."""
    import io
    from docx.enum.section import WD_SECTION
    document = Document()
    document.add_paragraph("第一节正文")
    document.sections[0].header.paragraphs[0].text = "共享页眉"
    document.sections[0].footer.paragraphs[0].text = "共享页脚"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("第二节正文")
    assert document.sections[1].header.is_linked_to_previous
    document.save(root / "sections.docx")
    prs = Presentation()
    for label in ("首页", "目录", "封底"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = label
    link = prs.slides[0].shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    link.text_frame.text = "跳到目录"
    link.click_action.target_slide = prs.slides[1]
    # Links from the master and from notes keep the dropped slide reachable without any slide referencing it.
    prs.slide_masters[0].placeholders[0].click_action.target_slide = prs.slides[1]
    notes = prs.slides[2].notes_slide
    notes.notes_text_frame.text = "备注"
    notes.notes_placeholder.click_action.target_slide = prs.slides[1]
    prs.save(root / "linked.pptx")
    prs = Presentation()
    for draw in ("ellipse", "rectangle"):
        picture = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
        getattr(ImageDraw.Draw(picture), draw)((20, 20, 236, 236), fill="black")
        buffer = io.BytesIO()
        picture.save(buffer, "PNG")
        buffer.seek(0)
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(buffer, Inches(1), Inches(1), Inches(3), Inches(3))
    prs.save(root / "shapes.pptx")
    prs = Presentation()
    for _ in range(51):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(root / "big.pptx")


def check_edge_cases(root, results):
    sections = Document(results["sectionsBuild"]["output"])
    assert sections.sections[0].header.paragraphs[0].text == "共享页眉", "inherited header must survive body clearing"
    assert sections.sections[0].footer.paragraphs[0].text == "共享页脚"
    assert not any("第一节正文" in p.text for p in sections.paragraphs)
    linked = results["linkedBuild"]
    deck = Presentation(linked["output"])
    assert [first_text(s) for s in deck.slides] == ["首页", "新页", "封底"], [first_text(s) for s in deck.slides]
    assert any("指向未保留页面的链接" in w for w in linked["warnings"]), linked["warnings"]
    assert any("母版、版式或备注" in w for w in linked["warnings"]), linked["warnings"]
    with zipfile.ZipFile(linked["output"]) as archive:
        parts = [n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        assert len(parts) == len(set(parts)) == 3, parts
        assert not any(b"hlinksldjump" in archive.read(n) for n in archive.namelist() if n.endswith(".xml")), "slide-jump links to dropped pages must be gone"
    assert not any(s.click_action.hyperlink.address or s.click_action.action for s in deck.slides[0].shapes if s.has_text_frame and "跳到目录" in s.text_frame.text)
    overflow = results["overflowBuild"]
    deck = Presentation(overflow["output"])
    titles = [first_text(s) for s in deck.slides]
    assert titles.count("逐行文字") + titles.count("逐行文字（续）") >= 2, "40 explicit lines must continue on a new slide"
    assert sum(1 for t in titles if t.startswith("长表格")) >= 2, "tall rows must page the table"
    column_widths = []
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= deck.slide_height * 0.93, (first_text(slide), shape.name)
        table = next((s.table for s in slide.shapes if s.has_table), None)
        if table:
            assert table.rows[0].cells[0].text == "项目", "header repeats on continuation"
            column_widths.append([c.width for c in table.columns])
    assert len(column_widths) >= 2 and all(w == column_widths[0] for w in column_widths), column_widths
    shapes = results["shapeImages"]
    assert len(shapes["images"]) == 2 and shapes["skipped"]["repeated"] == 0, "circle and square on transparency are distinct"
    filtered = results["filteredImages"]
    assert not filtered["images"] and filtered["skipped"]["small"] == 2, "minSize applies to PPT sources too"
    big = results["bigImages"]
    assert len(big["selectedPages"]) == 50 and big["pages"] == 51 and any("前 50 页" in w for w in big["warnings"]), big["warnings"]


def check_office_errors(root):
    script = Path(__file__).resolve().parents[2] / "src/agent/modules/document-work/scripts/document_ops.py"
    spec = importlib.util.spec_from_file_location("document_ops", script)
    operations = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(operations)
    source = root / "source.docx"
    for code, diagnostic in ((1, b"Read-only file system: /tmp/OSL_PIPE_fixture"), (0, b"Error: source file could not be loaded")):
        folder = root / f"office-error-{code}"
        folder.mkdir()
        result = subprocess.CompletedProcess([], code, stdout=b"", stderr=diagnostic)
        with patch.object(operations, "office_binary", return_value="fixture-soffice"), patch.object(operations.subprocess, "run", return_value=result) as run:
            try:
                operations.convert_office(source, folder)
                raise AssertionError("conversion failure must be reported")
            except ValueError as error:
                assert diagnostic.decode() in str(error), str(error)
            env = run.call_args.kwargs["env"]
            for key in ("HOME", "XDG_CONFIG_HOME", "XDG_CACHE_HOME", "XDG_DATA_HOME", "TMPDIR"):
                target = Path(env[key])
                assert target.is_dir() and target.is_relative_to(folder)


def prepare(root):
    check_edit_guards()
    root.mkdir(parents=True, exist_ok=True)
    picture = Image.new("RGB", (480, 240), "#16324F")
    ImageDraw.Draw(picture).rectangle((40, 40, 440, 200), fill="#58C1B2")
    picture.save(root / "product.png")
    document = Document()
    document.add_heading("客户方案", 0)
    p = document.add_paragraph()
    p.add_run("客户").bold = True
    p.add_run("A").italic = True
    p.add_run(" 使用正式产品资料。")
    document.sections[0].header.paragraphs[0].text = "客户 A 内部资料"
    document.add_picture(str(root / "product.png"), width=WordInches(3))
    table = document.add_table(rows=2, cols=2)
    table.style = "Light Shading Accent 1"
    table.cell(0, 0).text, table.cell(0, 1).text = "项目", "内容"
    table.cell(1, 0).text, table.cell(1, 1).text = "部署", "本地部署"
    document.save(root / "source.docx")
    check_office_errors(root)
    supplement = Document()
    supplement.add_heading("实施计划", 1)
    supplement.add_paragraph("准备、实施、验证。")
    supplement.add_picture(str(root / "product.png"), width=WordInches(2))
    supplement.save(root / "supplement.docx")
    for name, labels in [("source.pptx", ["第一页：能力", "第二页：部署", "第三页：实施"]), ("supplement.pptx", ["补充：验证结果"])]:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
        color = "F1F6FA" if name == "source.pptx" else "EFF7F0"
        background = etree.fromstring(
            ('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
             'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
             '<p:bgPr><a:solidFill><a:srgbClr val="' + color + '"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>').encode())
        prs.slide_masters[0].element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld").insert(0, background)
        for label in labels:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = label
            frame = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5), Inches(1)).text_frame
            p = frame.paragraphs[0]
            run = p.add_run(); run.text = "客户"; run.font.bold = True; run.font.size = Pt(22)
            run = p.add_run(); run.text = "A"; run.font.italic = True; run.font.size = Pt(22)
            slide.shapes.add_picture(str(root / "product.png"), Inches(0.8), Inches(3), width=Inches(4))
            chart = CategoryChartData(); chart.categories = ["准备", "验证"]; chart.add_series("进度", (30, 80))
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(2), Inches(6), Inches(4), chart)
            slide.notes_slide.notes_text_frame.text = "测试备注：" + label
        if name == "source.pptx":
            # Logical order differs from slide1.xml/slide2.xml order.
            ids = prs.slides._sldIdLst
            first = ids[0]; ids.remove(first); ids.append(first)
        prs.save(root / name)
    add_tag_fixture(root / "supplement.pptx")
    add_edge_fixtures(root)
    pdf = PdfWriter()
    for _ in range(22):
        pdf.add_blank_page(width=600, height=400)
    pdf.write(root / "pages.pdf")
    # Two distinct raster pages so image extraction can prove per-page results and de-duplication.
    second = Image.new("RGB", (480, 240), "#7A1F1F")
    ImageDraw.Draw(second).ellipse((60, 40, 420, 200), fill="#F2C14E")
    with Image.open(root / "product.png") as first:
        first.save(root / "figure.pdf", save_all=True, append_images=[second, second])


def run_guide_examples(root, results):
    """The skill's adjustment snippets must run against real build outputs in the fixed environment."""
    project = Path(__file__).resolve().parents[2]
    scratch = root / "guide-scratch"
    scratch.mkdir(exist_ok=True)
    shutil.copy(results["slideBuild"]["output"], scratch / "成稿.pptx")
    shutil.copy(results["wordBuild"]["output"], scratch / "成稿.docx")
    shutil.copy(root / "product.png", scratch / "flow.png")
    for guide in ("word", "slides"):
        content = (project / "src/agent/modules/document-work/skills/document-work/references" / (guide + ".md")).read_text(encoding="utf-8")
        code = re.search(r"```python\n(.*?)\n```", content, re.S).group(1)
        script = scratch / (guide + "-starter.py")
        script.write_text(code, encoding="utf-8")
        subprocess.run([sys.executable, str(script)], env={**os.environ, "PI_USER_TMP": str(scratch)}, check=True, timeout=30)
    assert len(Presentation(scratch / "成稿-调整.pptx").slides) == len(Presentation(scratch / "成稿.pptx").slides)
    adjusted = Document(scratch / "成稿-调整.docx")
    assert any("TOC" in (node.get(qn("w:instr")) or "") for node in adjusted.element.iter(qn("w:fldSimple")))
    run_flowchart_script(root, results, scratch)


def check(root, results):
    patched = Document(results["wordPatch"]["output"])
    p = next(p for p in patched.paragraphs if "客户B" in p.text)
    assert p.runs[0].bold and p.runs[1].italic
    assert len(patched.inline_shapes) == 1
    assert patched.tables[0].cell(1, 1).text == "私有化部署"
    with zipfile.ZipFile(root / "source.docx") as before, zipfile.ZipFile(results["wordPatch"]["output"]) as after:
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name), name
    merged = Document(results["wordCompose"]["output"])
    assert any(p.text == "实施计划" for p in merged.paragraphs)
    assert len(merged.inline_shapes) == 2
    assert merged.sections[0].header.paragraphs[0].text == "客户 A 内部资料"
    selected = Document(results["wordSelection"]["output"])
    assert len(selected.paragraphs) == 2 and len(selected.tables) == 0
    deck = Presentation(results["pptCompose"]["output"])
    assert any("自定义标签关联" in warning for warning in results["pptCompose"]["warnings"])
    assert [s.shapes.title.text for s in deck.slides] == ["第三页：实施", "第二页：部署", "补充：验证结果", "第三页：实施"]
    with zipfile.ZipFile(results["pptCompose"]["output"]) as archive:
        parts = [n for n in archive.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        assert len(parts) == 4, "unselected slide parts must not remain in the output archive"
        assert all("第一页：能力".encode() not in archive.read(n) for n in parts)
    for slide in deck.slides:
        assert sum(s.shape_type == 13 for s in slide.shapes) == 1
        charts = [s.chart for s in slide.shapes if s.has_chart]
        assert len(charts) == 1 and tuple(charts[0].series[0].values) == (30.0, 80.0)
        assert "测试备注" in slide.notes_slide.notes_text_frame.text
        expected = "EFF7F0" if "补充" in slide.shapes.title.text else "F1F6FA"
        assert str(slide.slide_layout.slide_master.background.fill.fore_color.rgb) == expected
    changed = Presentation(results["pptPatch"]["output"])
    assert any("客户B" in s.text for s in changed.slides[0].shapes if s.has_text_frame)
    with zipfile.ZipFile(results["pptCompose"]["output"]) as before, zipfile.ZipFile(results["pptPatch"]["output"]) as after:
        for name in before.namelist():
            if name != results["pptPatchPart"]:
                assert before.read(name) == after.read(name), name
    assert results["preview"]["pages"] == 22
    assert results["preview"]["unrenderedPages"] == list(range(3, 22))
    assert [i["page"] for i in results["preview"]["images"]] == [22, 1, 2]
    for item in results["preview"]["images"]:
        with Image.open(item["path"]) as picture:
            assert max(picture.size) <= 1800
    for name, stamp in results["sourceDigests"].items():
        assert hashlib.sha256((root / name).read_bytes()).hexdigest() == stamp
    for preview in results.get("officePreviews", []):
        assert preview["images"] and not preview["unrenderedPages"]
        assert preview["visuallyReviewed"] is False
    check_builds(root, results)
    run_guide_examples(root, results)
    print("DOCUMENT_ARTIFACTS_PASSED")


def first_text(slide):
    if slide.shapes.title is not None and slide.shapes.title.text.strip():
        return slide.shapes.title.text
    return next((s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()), "")


def check_builds(root, results):
    outline = results["outline"]["outline"]
    assert [s["title"] for s in outline["slides"]] == ["第二页：部署", "第三页：实施", "第一页：能力"], outline["slides"]
    assert all(s["pictures"] == 1 and s["charts"] == 1 for s in outline["slides"])
    assert outline["titleStyle"]["samples"] == 3 and outline["layouts"]
    word_outline = results["wordOutline"]["outline"]
    assert [h["text"] for h in word_outline["headings"]] == ["客户方案"] and word_outline["tables"] == 1
    assert "Heading 1" in word_outline["styles"] and word_outline["inlineImages"] == 1

    built = Document(results["wordBuild"]["output"])
    assert results["wordBuild"]["build"]["templated"] is True
    assert built.sections[0].header.paragraphs[0].text == "客户 A 内部资料", "template header must survive"
    styles = [p.style.name for p in built.paragraphs if p.text.strip()]
    assert styles[:3] == ["Title", "Subtitle", "Heading 1"], styles
    assert "List Bullet" in styles and "Heading 2" in styles and "Quote" in styles and "Caption" in styles
    bold = next(p for p in built.paragraphs if "OpenClaw" in p.text)
    assert any(r.bold and "OpenClaw" in r.text for r in bold.runs), "inline bold must map to runs"
    assert built.tables[0].rows[0].cells[0].text == "场景" and len(built.tables[0].rows) == 3
    assert built.tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert len(built.inline_shapes) == 1 and abs(built.inline_shapes[0].width - 6 * 360000) < 2000
    assert not any("客户A" in p.text for p in built.paragraphs), "template body must be cleared"
    default = Document(results["wordDefault"]["output"])
    assert default.paragraphs[0].style.name == "Title" and len(default.tables) == 1

    deck = Presentation(results["slideBuild"]["output"])
    build = results["slideBuild"]["build"]
    titles = [first_text(slide) for slide in deck.slides]
    assert titles[0] == "第二页：部署" and titles[-1] == "第一页：能力", titles
    assert build["keptPages"] == [1, len(deck.slides)] and build["generatedPages"] == list(range(2, len(deck.slides)))
    assert titles[1] == "交流方案" and titles[2] == "客户交流目标", titles
    assert any(t.startswith("附录：长文（续）") for t in titles), "long section must continue on a new slide"
    assert any(item["reason"].startswith("“附录：长文”") for item in build["attention"])
    assert build["titleStyle"]["samples"] >= 1 and build["layout"]
    generated = [deck.slides[i - 1] for i in build["generatedPages"]]
    table_slide = next(s for s in generated if any(sh.has_table for sh in s.shapes))
    table = next(sh.table for sh in table_slide.shapes if sh.has_table)
    assert table.cell(0, 0).text == "场景" and len(table.rows) == 3 and table.cell(2, 2).text == "独立资源池"
    picture_slide = next(s for s in generated if any(sh.shape_type == 13 for sh in s.shapes))
    assert any(sh.has_text_frame and "说明文字" in sh.text_frame.text for sh in picture_slide.shapes)
    assert "试点先行" in picture_slide.notes_slide.notes_text_frame.text
    bullet_slide = next(s for s in generated if any(sh.has_text_frame and "现状与风险" in sh.text_frame.text for sh in s.shapes))
    frame = next(sh.text_frame for sh in bullet_slide.shapes if sh.has_text_frame and "现状与风险" in sh.text_frame.text)
    levels = {p.text: p.level for p in frame.paragraphs}
    assert levels["现状与风险"] == 0 and levels["二级要点"] == 1
    assert all(str(s.slide_layout.slide_master.background.fill.fore_color.rgb) == "F1F6FA" for s in deck.slides), "template master must be reused"
    with zipfile.ZipFile(results["slideBuild"]["output"]) as archive:
        parts = [n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        assert len(parts) == len(deck.slides), "dropped template slides must not linger"
    default_deck = Presentation(results["slideDefault"]["output"])
    assert len(default_deck.slides) >= 4 and abs(default_deck.slide_width - 12192000) < 10

    inline = Presentation(results["slideInline"]["output"])
    inline_titles = [first_text(s) for s in inline.slides]
    assert inline_titles == ["第三页：实施", "新增页面", "补充：验证结果"], inline_titles
    assert any(sh.has_table for sh in inline.slides[1].shapes)
    assert str(inline.slides[1].slide_layout.slide_master.background.fill.fore_color.rgb) == "F1F6FA"
    assert str(inline.slides[2].slide_layout.slide_master.background.fill.fore_color.rgb) == "EFF7F0"
    merged = Document(results["wordInline"]["output"])
    assert [p.text for p in merged.paragraphs if p.style.name.startswith("Heading")] == ["补充章节"]
    assert merged.sections[0].header.paragraphs[0].text == "客户 A 内部资料"
    assert [p.text for p in merged.paragraphs if p.style.name == "List Number"] == ["第一步", "第二步"]
    assert len(merged.inline_shapes) == 1
    check_layouts_and_images(root, results)


def check_layouts_and_images(root, results):
    deck = Presentation(results["layoutBuild"]["output"])
    build = results["layoutBuild"]["build"]
    by_title = {first_text(slide): slide for slide in deck.slides}
    cards = [s for s in by_title["三层防护"].shapes if s.name.startswith("Card ")]
    assert [c.name for c in cards] == ["Card 终端侧", "Card 网络侧", "Card 平台侧"], [c.name for c in cards]
    assert cards[0].text_frame.paragraphs[0].text == "终端侧" and cards[0].text_frame.paragraphs[1].text == "客户端"
    assert any("围绕三个层面" in s.text_frame.text for s in by_title["三层防护"].shapes if s.has_text_frame), "intro text must stay above the cards"
    nodes = [s for s in by_title["实施安排"].shapes if s.name.startswith("Timeline node")]
    assert [n.text_frame.text for n in nodes] == ["1", "2", "3"], "stage-labelled lists become timelines automatically"
    assert any(s.has_text_frame and s.text_frame.text == "第一阶段" for s in by_title["实施安排"].shapes)
    assert any(s.has_text_frame and "调研与试点范围确认" in s.text_frame.text for s in by_title["实施安排"].shapes)
    explicit = [s.name for s in by_title["关键能力"].shapes if s.name.startswith("Card ")]
    assert explicit == ["Card 资产识别", "Card 风险检测", "Card 审计溯源"], explicit
    assert not any(s.name.startswith(("Card ", "Timeline")) for s in by_title["保持列表"].shapes), "<!-- plain --> keeps the list"
    assert not any(s.name.startswith(("Card ", "Timeline")) for s in by_title["无法分段"].shapes)
    assert any("timeline" in w and "分段" in w for w in results["layoutBuild"]["warnings"]), results["layoutBuild"]["warnings"]
    reasons = [item["reason"] for item in build["attention"]]
    assert any("3 段卡片" in r for r in reasons) and any("3 段时间轴" in r for r in reasons), reasons
    names = lambda title, prefix: [s.name for s in by_title[title].shapes if s.name.startswith(prefix)]
    assert names("实施流程", "Flow step") == ["Flow step 1", "Flow step 2", "Flow step 3", "Flow step 4"], "arrow chains become flows automatically"
    assert names("带图标的流程", "Icon ") == ["Icon 🔍", "Icon 🛡️", "Icon ⚙️"], names("带图标的流程", "Icon ")
    flow = [s for s in by_title["带图标的流程"].shapes if s.name == "Flow step 1"][0]
    assert flow.text_frame.text == "资产识别", "glyph must be stripped from the step label"
    assert [n for n in names("防护体系架构", "Layer ") if n != "Layer item"] == ["Layer 应用层", "Layer 平台层", "Layer 基础设施层"], "### names ending in 层 become layers"
    assert len(names("防护体系架构", "Layer item")) == 5
    assert names("关键指标", "Stat ") == ["Stat 135000+", "Stat 512", "Stat 24h"], "numeric labels become stats automatically"
    assert names("运营闭环", "Cycle node") == ["Cycle node 1", "Cycle node 2", "Cycle node 3", "Cycle node 4"]
    assert len(names("运营闭环", "Cycle arrow")) == 4
    with zipfile.ZipFile(results["layoutBuild"]["output"]) as archive:
        cycle_part = next(n for n in archive.namelist() if n.startswith("ppt/slides/slide") and "Cycle arrow".encode() in archive.read(n))
        assert b'tailEnd type="triangle"' in archive.read(cycle_part), "cycle arrows need arrowheads"
    assert names("能力成熟度", "Pyramid ") == ["Pyramid 智能防护", "Pyramid 基础防护"]
    widths = [s.width for s in by_title["能力成熟度"].shapes if s.name.startswith("Pyramid ")]
    assert widths[0] < widths[1], "pyramid narrows towards the top"
    icon_cards = by_title["带图标的卡片"]
    assert names("带图标的卡片", "Card ") == ["Card 资产识别", "Card 风险检测"]
    assert names("带图标的卡片", "Icon") == ["Icon ⭐", "Icon"], "glyph badge and image icon"
    assert any(s.shape_type == 13 for s in icon_cards.shapes), "image icon must be a picture"

    pdf_images = results["pdfImages"]
    assert pdf_images["pages"] == 3 and [i["page"] for i in pdf_images["images"]] == [1, 2], pdf_images["images"]
    assert pdf_images["images"][1]["pages"] == [2, 3] and pdf_images["skipped"]["repeated"] == 1
    assert pdf_images["images"][0]["fullPage"] is True and pdf_images["images"][0]["box"] == [0, 0, 1, 1]
    with Image.open(pdf_images["images"][0]["path"]) as picture:
        assert picture.size == (480, 240)
    crop = pdf_images["crops"][0]
    with Image.open(crop["path"]) as picture:
        assert abs(picture.width / picture.height - (0.5 * 480) / (0.8 * 240)) < 0.05
    deck_images = results["deckImages"]
    assert len(deck_images["images"]) == 1 and deck_images["images"][0]["pages"] == [1, 2, 3], deck_images["images"]
    assert deck_images["images"][0]["box"][0] > 0 and deck_images["images"][0]["fullPage"] is False
    assert deck_images["skipped"]["repeated"] == 2
    word_images = results["wordImages"]
    assert len(word_images["images"]) == 1 and word_images["images"][0]["width"] == 480
    for entry in pdf_images["images"] + deck_images["images"] + word_images["images"]:
        assert Path(entry["path"]).is_file()
    pictured = by_title["图片分层"]
    assert any("layers" in w and "图片图标" in w for w in results["layoutBuild"]["warnings"]), results["layoutBuild"]["warnings"]
    assert any(s.shape_type == 13 for s in pictured.shapes) and not names("图片分层", "Layer "), "pictures fall back to the plain layout instead of vanishing"
    assert names("指标单位", "Stat ") == ["Stat 135,000+", "Stat 82 个国家", "Stat 14.6%"], "numbers with short units are stats"
    assert names("小数指标", "Stat ") == ["Stat 1024.50 GB", "Stat 2048.25 GB", "Stat 4096.75 GB"], "decimals with units are stats, not dates"
    for title, labels in (
            ("年月形小数指标", ["2000.01 GB", "2000.02 GB", "2000.03 GB"]),
            ("混合小数指标", ["1024.50 GB", "2048.01 GB", "4096.75 GB"]),
            ("小数指标单位", ["2000.01GiB", "2000.02 ms", "2000.03 万元", "2000.04%"])):
        assert names(title, "Stat ") == ["Stat " + label for label in labels], f"{title}: units keep date-like decimals as stats"
        assert not names(title, "Timeline"), title
        page = deck.slides.index(by_title[title]) + 1
        assert any(item["page"] == page and item["mode"] == "stats" for item in build["layouts"]), title
    assert [n.text_frame.text for n in by_title["日期节点"].shapes if n.name.startswith("Timeline node")] == ["1", "2", "3"], "pure dates form a timeline"
    assert [n.text_frame.text for n in by_title["日期事件"].shapes if n.name.startswith("Timeline node")] == ["1", "2", "3"], "date + event labels stay a timeline, never stats"
    assert not names("日期事件", "Stat ") and any(s.has_text_frame and s.text_frame.text == "2026年3月立项" for s in by_title["日期事件"].shapes)
    assert [n.text_frame.text for n in by_title["数字日期事件"].shapes if n.name.startswith("Timeline node")] == ["1", "2", "3"], "numeric dates with event names stay a timeline"
    assert not names("数字日期事件", "Stat ") and any(s.has_text_frame and s.text_frame.text == "2026.09 UAT" for s in by_title["数字日期事件"].shapes)
    assert [n.text_frame.text for n in by_title["分期安排"].shapes if n.name.startswith("Timeline node")] == ["1", "2", "3"], "long parenthesised stage labels still form a timeline"
    assert any(s.has_text_frame and s.text_frame.text == "第一阶段（第 1 个月）风险摸底与试点" for s in by_title["分期安排"].shapes)
    wide = by_title["宽图说明"]
    picture = next(s for s in wide.shapes if s.shape_type == 13)
    caption = next(s for s in wide.shapes if s.has_text_frame and "整体架构" in s.text_frame.text)
    assert picture.width >= deck.slide_width * 0.6 and picture.top >= caption.top + caption.height, "wide figure with a short caption is stacked full-width"
    missed = next(item for item in build["attention"] if item["page"] == deck.slides.index(by_title["未识别列表"]) + 1)
    assert "未自动排成图示" in missed["reason"] and "<!-- cards -->" in missed["reason"], missed
    layouts = {item["page"]: item["mode"] for item in build["layouts"]}
    assert layouts[deck.slides.index(by_title["三层防护"]) + 1] == "cards" and layouts[deck.slides.index(by_title["实施安排"]) + 1] == "timeline"
    assert layouts[deck.slides.index(by_title["处理流程"]) + 1] == "flowchart" and layouts[deck.slides.index(by_title["指标单位"]) + 1] == "stats"
    assert deck.slides.index(by_title["未识别列表"]) + 1 not in layouts and deck.slides.index(by_title["保持列表"]) + 1 not in layouts
    assert not any(item["page"] == deck.slides.index(by_title["保持列表"]) + 1 for item in build["attention"]), "<!-- plain --> pages are not flagged"
    kept_build = results["slideBuild"]["build"]
    assert all(item["page"] in kept_build["generatedPages"] for item in kept_build["layouts"]) and kept_build["keptPages"], "layouts only describe generated pages"
    check_edge_cases(root, results)
    check_flowcharts(root, results)


def walk_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:
            yield from walk_shapes(shape.shapes)
        else:
            yield shape


def diamonds(shapes):
    return [s for s in shapes if s.shape_type == 1 and s.name.startswith("Flow node") and s.auto_shape_type == 63]  # FLOWCHART_DECISION


FLOW_SPEC = """flowchart TB
  S([收到告警]) --> A[自动归类与初判]
  A --> B{是否高危?}
  B -- 是 --> C[立即阻断并通知客户]
  B -- 否 --> D[进入工单队列]
  C --> E[安全大脑研判]
  D --> E
  E --> F{处置是否有效?}
  F -->|否| A
  F -->|是| G([结束])
  A -.-> L[(留存日志)]
"""


def check_flowcharts(root, results):
    deck = Presentation(results["layoutBuild"]["output"])
    build = results["layoutBuild"]["build"]
    by_title = {first_text(slide): slide for slide in deck.slides}
    flow = by_title["处理流程"]
    assert any(s.shape_type == 6 and s.name == "Flowchart" for s in flow.shapes), "flowchart shapes are grouped"
    shapes = list(walk_shapes(flow.shapes))
    assert [d.text_frame.text for d in diamonds(shapes)] == ["是否高危?"]
    assert len([s for s in shapes if s.name.startswith("Flow edge")]) == 7
    labels = {s.name: s.text_frame.text for s in shapes if s.name.startswith("Flow label")}
    assert labels == {"Flow label B->C": "是", "Flow label B->D": "否"}, labels
    assert any(s.has_text_frame and "收到告警后按下图处置" in s.text_frame.text for s in flow.shapes), "intro stays above the diagram"
    for shape in shapes:
        assert 0 <= shape.left and shape.left + shape.width <= deck.slide_width and shape.top + shape.height <= deck.slide_height, shape.name
    with zipfile.ZipFile(results["layoutBuild"]["output"]) as archive:
        part = next(n for n in archive.namelist() if n.startswith("ppt/slides/slide") and "Flow node S".encode() in archive.read(n))
        assert b'tailEnd type="triangle"' in archive.read(part) and b'prstDash val="dash"' in archive.read(part)
    assert any("6 节点流程图" in item["reason"] for item in build["attention"]), build["attention"]
    assert any("流程图未生成" in w for w in results["layoutBuild"]["warnings"]), results["layoutBuild"]["warnings"]
    assert any(s.has_text_frame and "A -->" in s.text_frame.text for s in by_title["无法解析"].shapes), "unparseable source is kept as code"
    word_flow = Document(results["wordFlow"]["output"])
    if any("流程图未生成" in w for w in results["wordFlow"]["warnings"]):
        assert any("字体" in w for w in results["wordFlow"]["warnings"]), results["wordFlow"]["warnings"]
    else:
        assert len(word_flow.inline_shapes) == 1 and word_flow.sections[0].header.paragraphs[0].text == "客户 A 内部资料"
        assert not any("flowchart" in p.text for p in word_flow.paragraphs), "source must not remain once the picture is inserted"


def run_flowchart_script(root, results, scratch):
    """The skill's flowchart.py must work as a CLI on real build outputs: check, PPT, PNG and Word insertion."""
    project = Path(__file__).resolve().parents[2]
    script = project / "src/agent/modules/document-work/skills/document-work/scripts/flowchart.py"
    spec = scratch / "flow.mmd"
    spec.write_text(FLOW_SPEC, encoding="utf-8")
    env = {**os.environ, "PYTHONUTF8": "1"}

    def run(*args, spec_path=spec):
        return subprocess.run([sys.executable, str(script), str(spec_path), *args], env=env, capture_output=True, text=True, encoding="utf-8", timeout=90)

    check = run("--check")
    assert check.returncode == 0, check.stderr
    summary = json.loads(check.stdout)
    assert summary["layers"] >= 6 and sum(1 for e in summary["edges"] if e["back"]) == 1 and summary["nodes"][0]["kind"] == "terminator"
    bad = scratch / "bad.mmd"
    bad.write_text("A -- B -- \n", encoding="utf-8")
    failed = run("--check", spec_path=bad)
    assert failed.returncode == 1 and "无法解析" in json.loads(failed.stderr)["error"], failed.stderr
    before = Presentation(scratch / "成稿.pptx")
    target = next(i for i, s in enumerate(before.slides) if first_text(s) == "客户交流目标")
    drawn = run("--pptx", str(scratch / "成稿.pptx"), "--slide", str(target + 1), "--out", str(scratch / "成稿-流程.pptx"))
    assert drawn.returncode == 0, drawn.stderr
    info = json.loads(drawn.stdout)["pptx"]
    assert info["nodes"] == 9 and info["edges"] == 10 and info["direction"] in ("TB", "LR"), info
    after = Presentation(scratch / "成稿-流程.pptx")
    assert len(after.slides) == len(before.slides)
    slide = after.slides[target]
    shapes = list(walk_shapes(slide.shapes))
    assert [d.text_frame.text for d in diamonds(shapes)] == ["是否高危?", "处置是否有效?"]
    assert any(s.name == "Flowchart" for s in slide.shapes)
    intro = next(s for s in before.slides[target].shapes if s.has_text_frame and "现状与风险" in s.text_frame.text)
    title = before.slides[target].shapes.title
    for shape in [s for s in shapes if s.name.startswith("Flow ")]:
        assert shape.top + shape.height <= after.slide_height and shape.left >= 0 and shape.left + shape.width <= after.slide_width, shape.name
        assert shape.top >= title.top + title.height and shape.top > intro.top, "diagram must start below the title and intro text, not over them"
    png = run("--png", str(scratch / "flow.png"))
    if png.returncode == 2:
        assert "中文字体" in json.loads(png.stderr)["error"], png.stderr
        return
    assert png.returncode == 0, png.stderr
    with Image.open(scratch / "flow.png") as picture:
        assert picture.width == 1600 and picture.height > picture.width, picture.size
    inserted = run("--docx", str(scratch / "成稿.docx"), "--after", "客户交流目标", "--caption", "图 1 处理流程", "--out", str(scratch / "成稿-流程.docx"))
    assert inserted.returncode == 0, inserted.stderr
    original = Document(scratch / "成稿.docx")
    document = Document(scratch / "成稿-流程.docx")
    assert len(document.inline_shapes) == len(original.inline_shapes) + 1
    anchor = next(i for i, p in enumerate(document.paragraphs) if "客户交流目标" in p.text)
    assert document.paragraphs[anchor + 1]._p.xpath(".//w:drawing"), "picture follows the anchor paragraph"
    assert document.paragraphs[anchor + 2].text == "图 1 处理流程"


if __name__ == "__main__":
    mode, directory = sys.argv[1:3]
    if mode == "prepare":
        prepare(Path(directory))
    else:
        check(Path(directory), json.loads(Path(sys.argv[3]).read_text(encoding="utf-8")))
