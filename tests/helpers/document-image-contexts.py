"""Offline regression and timing probe; requires the repository's document Python environment."""
import importlib.util
import io
import json
import sys
import time
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

root = Path(sys.argv[1]).resolve()
root.mkdir(parents=True, exist_ok=True)
module_path = Path(__file__).resolve().parents[2] / "src/agent/modules/document-work/scripts/document_ops.py"
spec = importlib.util.spec_from_file_location("document_ops", module_path)
ops = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ops)
document = Document()
document.add_heading("安全功能", 1)
last_picture = None
for n in range(3200):
    document.add_paragraph("功能说明段落 %d：安全网关策略与截图证据。" % n)
    if n % 8 == 0:
        image = Image.new("RGB", (128, 128), (n % 256, n // 16 % 256, n // 256))
        data = io.BytesIO()
        image.save(data, "PNG")
        image.close()
        data.seek(0)
        last_picture = document.add_paragraph().add_run().add_picture(data)
        document.add_paragraph("图 %d 策略截图" % (n // 8 + 1), style="Caption")

# A repeated image in a table, and a floating picture in a textbox, must still get paragraph context.
paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
paragraph.add_run("表格内截图")
paragraph._p.append(deepcopy(last_picture._inline.getparent()))
outer = document.add_paragraph("文本框外")
textbox = OxmlElement("w:txbxContent")
inner = deepcopy(paragraph._p)
wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
for inline in inner.iter("{%s}inline" % wp):
    inline.tag = "{%s}anchor" % wp
textbox.append(inner)
outer._p.append(textbox)
source = root / "manual.docx"
document.save(source)

loaded = Document(source)
start = time.perf_counter()
contexts = ops.word_image_contexts(loaded)
elapsed = time.perf_counter() - start
assert len(contexts) == 402, len(contexts)
assert contexts[0]["headingPath"] == ["安全功能"]
assert contexts[0]["caption"] == "图 1 策略截图"
assert "功能说明" in contexts[0]["before"]
assert contexts[-1]["paragraph"] > contexts[-2]["paragraph"]
assert "表格内截图" in contexts[-1]["before"] and "表格内截图" in contexts[-2]["before"]
assert elapsed < 10, elapsed

def call(name, **kwargs):
    directory = root / name
    directory.mkdir()
    return ops.images({"source": str(source), "directory": str(directory), **kwargs})

first = call("metadata", extract=False)
assert 0 < len(first["images"]) <= 60 and first["next"] == first["images"][-1]["order"] and first["totalMatched"] == 402
assert not any("path" in image for image in first["images"])
assert not list((root / "metadata/images").iterdir())
second = call("next", extract=False, after=first["next"])
assert second["images"][0]["order"] == first["next"] + 1
matched = call("matched", extract=False, match=["图 200 策略截图"])
assert matched["images"][0]["order"] == 200
repeated = call("repeated", orders=[400, 401, 402])
assert len(repeated["images"]) == 1 and len(repeated["images"][0]["occurrences"]) == 3
assert Path(repeated["images"][0]["path"]).is_file()
short = call("short-context", extract=False, contextChars=17)
assert all(len(image["before"]) <= 17 and len(image["after"]) <= 17 for image in short["images"])
for paragraph in loaded.paragraphs:
    if paragraph.text.startswith("功能说明"):
        paragraph.text += "长上下文" * 90
wide_source = root / "long-context.docx"
loaded.save(wide_source)
bounded = call("bounded-context", source=str(wide_source), extract=False, contextChars=300)
assert len(json.dumps(bounded, ensure_ascii=False).encode("utf-8")) < 64 * 1024
assert len(bounded["images"]) < 60 and bounded["skipped"]["limit"] > 0 and bounded["next"]
continued = call("bounded-next", source=str(wide_source), extract=False, contextChars=300, after=bounded["next"])
assert continued["images"][0]["order"] == bounded["next"] + 1
print(json.dumps({"paragraphs": len(loaded.paragraphs), "images": len(contexts), "contextSeconds": elapsed,
                  "metadataCount": len(first["images"]), "repeatedOccurrences": 3, "boundedContextCount": len(bounded["images"])}, ensure_ascii=False))
