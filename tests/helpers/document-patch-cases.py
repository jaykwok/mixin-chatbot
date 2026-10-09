"""Deterministic document_patch regressions on Word and PPT XML: only the characters that differ change, every other character
stays in its run (formatting, link, structure), the existing guards still apply, and formatting that cannot be determined fails
instead of being guessed. Usage: python document-patch-cases.py <work directory> (also run by document-fixtures.py prepare)."""
import copy
import importlib.util
import json
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree

SCRIPT = Path(__file__).resolve().parents[2] / "src/agent/modules/document-work/scripts/document_ops.py"
_spec = importlib.util.spec_from_file_location("document_ops_patch_cases", SCRIPT)
ops = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ops)
W, A, R = ops.NS["w"], ops.NS["a"], ops.NS["r"]
SPACE = "{http://www.w3.org/XML/1998/namespace}space"
RUNS = ("{" + W + "}r", "{" + A + "}r")

# label -> (Word w:rPr content, Word w:r attributes, PPT a:rPr attributes, PPT a:rPr children)
FORMATS = {
    "plain": ("", "", "", ""),
    "bold": ("<w:b/>", "", ' b="1"', ""),
    # Same look, different editor state: Word splits runs by rsid, PowerPoint keeps spelling/edit flags on a:rPr.
    "bold-state": ("<w:b/>", ' w:rsidR="00AB12CD" w:rsidRPr="00EF3456"', ' b="1" dirty="0" err="1"', ""),
    "italic": ("<w:i/>", "", ' i="1"', ""),
    "font": ('<w:rFonts w:ascii="Arial Black" w:hAnsi="Arial Black" w:eastAsia="SimHei"/>', "", "", '<a:latin typeface="Arial Black"/><a:ea typeface="SimHei"/>'),
    "size": ('<w:sz w:val="44"/><w:szCs w:val="44"/>', "", ' sz="3200"', ""),
    "colour": ('<w:color w:val="C00000"/>', "", "", '<a:solidFill><a:srgbClr val="C00000"/></a:solidFill>'),
    "super": ('<w:vertAlign w:val="superscript"/>', "", ' baseline="30000"', ""),
    "sub": ('<w:vertAlign w:val="subscript"/>', "", ' baseline="-25000"', ""),
    "english": ('<w:lang w:val="en-US"/>', "", ' lang="en-US"', ""),
    "chinese": ('<w:lang w:val="zh-CN" w:eastAsia="zh-CN"/>', "", ' lang="zh-CN"', ""),
    "linked": ('<w:color w:val="0563C1"/><w:u w:val="single"/>', "", "", ""),
}


def run(text, label="plain"):
    return ("run", text, label)


def link(target, *runs):
    """Word: one w:hyperlink holding the runs; PPT: the link is part of each run's a:rPr."""
    return ("link", target, runs)


def raw(word, ppt=None):
    return ("raw", word, ppt)


def word_run(text, label):
    content, attributes, _, _ = FORMATS[label]
    properties = f"<w:rPr>{content}</w:rPr>" if content else ""
    space = ' xml:space="preserve"' if text != text.strip() else ""
    return f"<w:r{attributes}>{properties}<w:t{space}>{escape(text)}</w:t></w:r>" if text else f"<w:r{attributes}>{properties}<w:t/></w:r>"


def ppt_run(text, label, target=None):
    _, _, attributes, children = FORMATS[label]
    children += f'<a:hlinkClick r:id="{target}"/>' if target else ""
    properties = f"<a:rPr{attributes}>{children}</a:rPr>" if attributes or children else ""
    return f"<a:r>{properties}<a:t>{escape(text)}</a:t></a:r>"


def markup(kind, items, links):
    pieces = []
    for item in items:
        if item[0] == "run":
            pieces.append(word_run(*item[1:]) if kind == "word" else ppt_run(*item[1:]))
        elif item[0] == "link":
            target = links.get(item[1], item[1])
            if kind == "word":
                pieces.append(f'<w:hyperlink r:id="{target}" w:history="1">' + "".join(word_run(*r[1:]) for r in item[2]) + "</w:hyperlink>")
            else:
                pieces.extend(ppt_run(*r[1:], target=target) for r in item[2])
        else:
            pieces.append(item[1] if kind == "word" else item[2])
    if kind == "word":
        return f'<w:p xmlns:w="{W}" xmlns:r="{R}"><w:pPr><w:spacing w:after="240"/></w:pPr>{"".join(pieces)}</w:p>'
    return f'<a:p xmlns:a="{A}" xmlns:r="{R}"><a:pPr algn="l"/>{"".join(pieces)}<a:endParaRPr lang="zh-CN"/></a:p>'


def build(kind, items, links=None):
    return ops.xml(markup(kind, items, links or {}).encode())


def skeleton(element):
    """Canonical XML without the characters of w:t/a:t (and their xml:space): equal skeletons mean every run, property,
    link and other node is unchanged."""
    copied = copy.deepcopy(element)
    for node in copied.iter(*ops.TEXT):
        node.text = None
        node.attrib.pop(SPACE, None)
    return etree.tostring(copied, method="c14n")


def characters(paragraph):
    """Each visible character with the ordinal of the run holding it (runs counted in document order, nested paragraphs excluded)."""
    runs, result = {}, []
    for node in ops.paragraph_nodes(paragraph):
        if node.tag in RUNS:
            runs[node] = len(runs)
        elif node.tag in ops.TEXT:
            result.extend((c, runs[node.getparent()]) for c in node.text or "")
    return result


def expand(segments):
    return [(c, ordinal) for text, ordinal in segments for c in text]


def check_paragraph(kind, name, items, edits, expected):
    """expected: [(text, run ordinal)...] for the whole paragraph, "noop", or ("error", message fragment)."""
    paragraph = build(kind, items)
    original, shape = etree.tostring(paragraph), skeleton(paragraph)
    label = f"{kind}: {name}"
    if isinstance(expected, tuple):
        try:
            for before, after in edits:
                ops.replace_text(paragraph, before, after)
            raise AssertionError(f"{label}: expected failure “{expected[1]}”")
        except ValueError as error:
            assert expected[1] in str(error), (label, str(error))
        assert etree.tostring(paragraph) == original, f"{label}: a failed edit changed the paragraph"
        return
    changed = [ops.replace_text(paragraph, before, after) for before, after in edits]
    if expected == "noop":
        assert changed == [False] * len(edits) and etree.tostring(paragraph) == original, f"{label}: no-op changed the XML"
        return
    assert all(changed), (label, changed)
    assert skeleton(paragraph) == shape, f"{label}: runs, properties or structure changed\n{etree.tostring(paragraph, encoding='unicode')}"
    assert characters(paragraph) == expand(expected), (label, characters(paragraph), expand(expected))
    if kind == "ppt":
        assert not any(SPACE in node.attrib for node in paragraph.iter(*ops.TEXT)), f"{label}: xml:space added to a:t"


CROSS_FORMAT = "跨越格式或所属结构"
INSERT_BOUNDARY = "插入位置两侧"
AMBIGUOUS = "修改位置不唯一"
UNIQUE = "恰好出现一次"
P = [run("客户", "bold"), run("A", "italic"), run(" 保持原文")]
SPLIT = [run("客", "bold"), run("户", "bold-state"), run("A", "italic")]
SUFFIX = [run("客户", "bold"), run("A", "italic"), run("公司", "colour")]
LINKED = [run("访问"), link("rIdLink", run("官网", "linked"), run("首页", "linked")), run("了解")]
FORMATTED = [run("H"), run("2", "sub"), run("O 与 x"), run("2", "super"), run("；"), run("红字", "colour"), run("大号", "size"),
             run("Arial", "font"), run("English", "english"), run("中文", "chinese")]

# (name, items, edits, expected) for both Word and PPT.
SHARED = [
    ("single run", [run("客户A公司")], [("客户A", "客户B")], [("客户B公司", 0)]),
    ("single run, different length", [run("客户A公司")], [("客户A", "甲方客户")], [("甲方客户公司", 0)]),
    ("cross-run whole match (reported defect)", P, [("客户A", "客户B")], [("客户", 0), ("B", 1), (" 保持原文", 2)]),
    ("minimal match control", P, [("A", "B")], [("客户", 0), ("B", 1), (" 保持原文", 2)]),
    ("no-op whole match", P, [("客户A", "客户A")], "noop"),
    ("no-op single run", [run("客户A公司")], [("A公", "A公")], "noop"),
    ("same length across formats keeps each position", P, [("客户A", "顾问B")], [("顾问", 0), ("B", 1), (" 保持原文", 2)]),
    ("different length across formats fails", P, [("客户A", "甲方")], ("error", CROSS_FORMAT)),
    ("split same-format runs, replacement", SPLIT, [("客户", "顾客们")], [("顾客们", 0), ("A", 2)]),
    ("split same-format runs, insertion at the split", SPLIT, [("客户", "客全户")], [("客全", 0), ("户", 1), ("A", 2)]),
    ("split same-format runs, insertion at match start", SPLIT, [("客户", "甲方客户")], [("甲方客", 0), ("户", 1), ("A", 2)]),
    ("whole match unique, shortened text repeats later", [run("客户A", "bold"), run("，A 类", "italic")], [("客户A", "客户B")],
     [("客户B", 0), ("，A 类", 1)]),
    ("whole match unique, shortened text repeats earlier", [run("A 类：", "italic"), run("客户A", "bold")], [("客户A", "客户B")],
     [("A 类：", 0), ("客户B", 1)]),
    ("common prefix and suffix kept", SUFFIX, [("客户A公司", "客户B公司")], [("客户", 0), ("B", 1), ("公司", 2)]),
    ("insertion between different formats fails", SUFFIX, [("客户A公司", "客户AB公司")], ("error", INSERT_BOUNDARY)),
    ("narrowed insertion takes the side that remains", SUFFIX, [("A", "AB")], [("客户", 0), ("AB", 1), ("公司", 2)]),
    ("insertion inside a run", P, [("保持原文", "保持全部原文")], [("客户", 0), ("A", 1), (" 保持全部原文", 2)]),
    ("deletion inside a run", P, [(" 保持原文", " 原文")], [("客户", 0), ("A", 1), (" 原文", 2)]),
    ("deletion of a whole run", P, [("客户A", "客户")], [("客户", 0), (" 保持原文", 2)]),
    ("deletion across formats", P, [("户A 保", "户保")], [("客户", 0), ("保持原文", 2)]),
    ("ambiguous deletion across formats fails", [run("好", "bold"), run("好", "italic"), run("的")], [("好好的", "好的")], ("error", AMBIGUOUS)),
    ("ambiguous deletion in one format", [run("好好", "bold"), run("的")], [("好好的", "好的")], [("好", 0), ("的", 1)]),
    ("ambiguous insertion across formats fails", [run("哈", "bold"), run("哈", "italic")], [("哈哈", "哈哈哈")], ("error", AMBIGUOUS)),
    ("ambiguous insertion in one format", [run("哈", "bold"), run("哈", "bold-state")], [("哈哈", "哈哈哈")], [("哈", 0), ("哈哈", 1)]),
    ("empty run between same-format runs", [run("客户", "bold"), run("", "italic"), run("A", "bold")], [("客户A", "甲方")], [("甲方", 0)]),
    ("empty run at an insertion between same formats", [run("客户", "bold"), run("", "italic"), run("A", "bold")], [("客户A", "客户的A")],
     [("客户的", 0), ("A", 2)]),
    ("empty run at an insertion between different formats fails", [run("客户", "bold"), run("", "bold"), run("A", "italic")],
     [("客户A", "客户的A")], ("error", INSERT_BOUNDARY)),
    ("empty run before the match", [run("", "italic"), run("客户A", "bold")], [("客户", "新客户")], [("新客户A", 1)]),
    ("multiple edits in one paragraph", P, [("客户A", "客户B"), ("保持原文", "保持译文")], [("客户", 0), ("B", 1), (" 保持译文", 2)]),
    ("later edit sees the earlier one", P, [("A", "B"), ("客户B", "客户C")], [("客户", 0), ("C", 1), (" 保持原文", 2)]),
    ("subscript and superscript", FORMATTED, [("H2O", "H3O"), ("x2", "x3")],
     [("H", 0), ("3", 1), ("O 与 x", 2), ("3", 3), ("；", 4), ("红字", 5), ("大号", 6), ("Arial", 7), ("English", 8), ("中文", 9)]),
    ("colour, size, font and language", FORMATTED, [("红字大号", "蓝字大号"), ("Arial", "Arial Black"), ("English中文", "Englisch中文"), ("中文", "汉语")],
     [("H", 0), ("2", 1), ("O 与 x", 2), ("2", 3), ("；", 4), ("蓝字", 5), ("大号", 6), ("Arial Black", 7), ("Englisch", 8), ("汉语", 9)]),
    ("different sizes fail", [run("大号", "size"), run("小号")], [("大号小号", "中号")], ("error", CROSS_FORMAT)),
    ("different languages fail", [run("客户", "chinese"), run("A", "english")], [("客户A", "甲方")], ("error", CROSS_FORMAT)),
    ("link text, same length", LINKED, [("官网首页", "官网主页")], [("访问", 0), ("官网", 1), ("主页", 2), ("了解", 3)]),
    ("link text across the link's own runs", LINKED, [("官网首页", "官方首页面")], [("访问", 0), ("官方首页面", 1), ("了解", 3)]),
    ("insertion at a link boundary fails", LINKED, [("问官", "问一下官")], ("error", INSERT_BOUNDARY)),
    ("replacement across a link boundary fails", LINKED, [("访问官网", "浏览网")], ("error", CROSS_FORMAT)),
    ("replacement next to a link", LINKED, [("访问官网", "去官网")], [("去", 0), ("官网", 1), ("首页", 2), ("了解", 3)]),
    ("two different links fail", [link("rIdLink", run("官网", "linked")), link("rIdOther", run("首页", "linked"))], [("官网首页", "主站")],
     ("error", CROSS_FORMAT)),
    # Same look inside and after the link: only the link itself separates them.
    ("insertion at a link end fails", [link("rIdLink", run("官网")), run("首页")], [("官网首页", "官网的首页")], ("error", INSERT_BOUNDARY)),
    ("replacement across a link end fails", [link("rIdLink", run("官网")), run("首页")], [("官网首页", "主页")], ("error", CROSS_FORMAT)),
    # Guards that existed before the fix.
    ("tab or line break in the text", P, [("客户A", "客户\tB")], ("error", "不跨换行或制表符")),
    ("not unique", [run("A 与 A")], [("A", "B")], ("error", UNIQUE)),
    ("overlapping occurrences are not unique", [run("哈哈哈")], [("哈哈", "嘿嘿")], ("error", UNIQUE)),
    ("not found", P, [("不存在", "x")], ("error", UNIQUE)),
]

BOOKMARK = [run("客户", "bold"), raw('<w:bookmarkStart w:id="0" w:name="mark"/>'), run("名称", "bold"), raw('<w:bookmarkEnd w:id="0"/>')]
PICTURE = [run("客户", "bold"), raw("<w:r><w:rPr><w:noProof/></w:rPr><w:drawing/></w:r>"), run("A", "bold")]
WORD_ONLY = [
    ("insertion at a bookmark boundary fails", BOOKMARK, [("客户名称", "客户之名称")], ("error", INSERT_BOUNDARY)),
    ("same length next to a bookmark", BOOKMARK, [("客户名称", "客户全称")], [("客户", 0), ("全称", 1)]),
    ("deletion next to a bookmark", BOOKMARK, [("客户名称", "客户称")], [("客户", 0), ("称", 1)]),
    ("replacement across a bookmark fails", BOOKMARK, [("客户名称", "甲方")], ("error", CROSS_FORMAT)),
    ("proofing marks between same-format runs", [run("客", "bold"), raw('<w:proofErr w:type="spellStart"/>'), run("户", "bold"),
                                                 raw('<w:proofErr w:type="spellEnd"/>')], [("客户", "顾客们")], [("顾客们", 0)]),
    ("replacement across a picture fails", PICTURE, [("客户A", "甲方")], ("error", CROSS_FORMAT)),
    ("same length around a picture", PICTURE, [("客户A", "客户B")], [("客户", 0), ("B", 2)]),
    ("insertion at a comment range fails", [run("客户", "bold"), raw('<w:commentRangeStart w:id="0"/>'), run("A", "bold"),
                                           raw('<w:commentRangeEnd w:id="0"/>'), raw('<w:r><w:commentReference w:id="0"/></w:r>')],
     [("客户A", "客户的A")], ("error", INSERT_BOUNDARY)),
    ("replacement across a footnote reference fails", [run("客户", "bold"), raw('<w:r><w:footnoteReference w:id="1"/></w:r>'), run("A", "bold")],
     [("客户A", "甲方")], ("error", CROSS_FORMAT)),
    ("tab between runs is kept", [run("客户", "bold"), raw("<w:r><w:tab/></w:r>"), run("A", "italic")], [("A", "B")], [("客户", 0), ("B", 2)]),
    ("simple field", [raw('<w:fldSimple w:instr="DATE"><w:r><w:t>日期</w:t></w:r></w:fldSimple>'), run("之后")], [("之后", "以后")], ("error", "域或修订")),
    ("complex field", [run("页码"), raw('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'), raw('<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'),
                       raw('<w:r><w:fldChar w:fldCharType="separate"/></w:r>'), run("1"), raw('<w:r><w:fldChar w:fldCharType="end"/></w:r>')],
     [("页码", "页")], ("error", "域或修订")),
    ("inserted revision", [run("客户"), raw('<w:ins w:id="1" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r><w:t>A</w:t></w:r></w:ins>')],
     [("客户", "甲方")], ("error", "域或修订")),
    ("deleted revision", [run("客户"), raw('<w:del w:id="1" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>旧</w:delText></w:r></w:del>')],
     [("客户", "甲方")], ("error", "域或修订")),
]
PPT_ONLY = [
    ("slide number field", [run("第 "), raw(None, '<a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum"><a:t>1</a:t></a:fld>'), run(" 页")],
     [(" 页", "页")], ("error", "域或修订")),
    ("line break between runs is kept", [run("第一行", "bold"), raw(None, "<a:br/>"), run("第二行", "italic")], [("第二", "次二")], [("第一行", 0), ("次二行", 1)]),
    ("line break in the text", [run("第一行", "bold"), raw(None, "<a:br/>"), run("第二行", "italic")], [("一行\n第二", "一二")], ("error", "不跨换行或制表符")),
]


def check_nested_and_whitespace():
    # A text box paragraph inside a Word paragraph is its own location and is not edited through the outer one.
    paragraph = ops.xml(f'<w:p xmlns:w="{W}"><w:r><w:t>外层</w:t></w:r><w:r><w:pict><w:p><w:r><w:t>外层文本框</w:t></w:r></w:p></w:pict></w:r></w:p>'.encode())
    assert ops.replace_text(paragraph, "外层", "修改") and ops.visible_text(paragraph) == "修改"
    assert paragraph.find(".//w:pict/w:p", ops.NS) is not None and ops.visible_text(paragraph.find(".//w:pict/w:p", ops.NS)) == "外层文本框"
    # Word drops edge spaces of w:t without xml:space; the attribute is added only where the new text needs it.
    paragraph = build("word", [run("客户", "bold"), run("A", "italic")])
    ops.replace_text(paragraph, "客户", "客户 ")
    first, second = paragraph.findall(".//w:t", ops.NS)
    assert first.text == "客户 " and first.get(SPACE) == "preserve" and second.get(SPACE) is None


def check_part(kind, source, output, expected):
    """Compares one part before and after a patch: every paragraph not in expected is unchanged including its text, the
    structure of the whole part is unchanged, and the paragraphs in expected ({number: [(text, run ordinal)...]}) hold exactly
    those characters in those runs."""
    old_root, new_root = ops.xml(source), ops.xml(output)
    assert skeleton(old_root) == skeleton(new_root), "structure or run properties of the part changed"
    old, new = ops.paragraphs(old_root, kind), ops.paragraphs(new_root, kind)
    assert len(old) == len(new)
    for number, (before, after) in enumerate(zip(old, new), 1):
        if number in expected:
            assert characters(after) == expand(expected[number]), (number, characters(after))
        else:
            assert etree.tostring(before, method="c14n") == etree.tostring(after, method="c14n"), number


def number_of(kind, data, part, text):
    return next(i for i, p in enumerate(ops.paragraphs(ops.xml(data[part]), kind), 1) if ops.visible_text(p) == text)


def contents(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def word_source(path):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import parse_xml
    document = Document()
    links = {name: document.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
             for name, url in (("rIdLink", "https://example.com/"), ("rIdOther", "https://example.org/"))}
    body = document.element.body
    for items in (P, FORMATTED, LINKED, SUFFIX):
        body.insert(len(body) - 1, parse_xml(markup("word", items, links)))
    header = document.sections[0].header.paragraphs[0]._p
    header.addnext(parse_xml(markup("word", [run("页眉：客户", "bold"), run("A", "italic")], links)))
    header.getparent().remove(header)
    document.save(path)


def slide_source(path):
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE
    from pptx.oxml import parse_xml
    from pptx.util import Inches
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    links = {name: slide.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
             for name, url in (("rIdLink", "https://example.com/"), ("rIdOther", "https://example.org/"))}
    frame = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(6)).text_frame
    body = frame._txBody
    for paragraph in body.findall("{" + A + "}p"):
        body.remove(paragraph)
    for items in (P, FORMATTED, LINKED, SUFFIX):
        body.append(parse_xml(markup("ppt", items, links)))
    slide.notes_slide.notes_text_frame.text = "备注：客户A"
    deck.save(path)


def check_packages(directory):
    """patch() on real packages: several parts and edits, unchanged parts keep their bytes, no-op patches keep every byte,
    and a failing edit leaves no output file."""
    directory.mkdir(parents=True, exist_ok=True)
    results = {}
    for kind, suffix, make, other_part in (("docx", "docx", word_source, "word/header1.xml"), ("pptx", "pptx", slide_source, "ppt/notesSlides/notesSlide1.xml")):
        source = directory / f"patch-source.{suffix}"
        if source.exists():
            source.unlink()
        make(source)
        data = contents(source)
        main = "word/document.xml" if kind == "docx" else "ppt/slides/slide1.xml"
        base = {"source": str(source), "digest": ops.digest(source)}
        number = lambda part, text: number_of(kind, data, part, text)
        p, formatted, linked = number(main, "客户A 保持原文"), number(main, "H2O 与 x2；红字大号ArialEnglish中文"), number(main, "访问官网首页了解")
        other_text = "页眉：客户A" if kind == "docx" else "备注：客户A"
        other = number(other_part, other_text)
        edits = [
            {"part": main, "paragraph": p, "before": "客户A", "after": "客户B"},
            {"part": main, "paragraph": formatted, "before": "H2O", "after": "H3O"},
            {"part": main, "paragraph": formatted, "before": "x2", "after": "x3"},
            {"part": main, "paragraph": formatted, "before": "红字大号", "after": "蓝字大号"},
            {"part": main, "paragraph": linked, "before": "官网首页", "after": "官方首页面"},
            {"part": other_part, "paragraph": other, "before": other_text, "after": other_text[:-1] + "B"},
        ]
        output = directory / f"patch-output.{suffix}"
        for stale in directory.glob("patch-*-output*." + suffix):
            stale.unlink()
        if output.exists():
            output.unlink()
        result = ops.patch({**base, "output": str(output), "edits": edits})
        assert result["changedParts"] == [main, other_part], result["changedParts"]
        assert not any("相同" in w for w in result["warnings"]), result["warnings"]
        patched = contents(output)
        assert set(patched) == set(data)
        for name in data:
            if name not in (main, other_part):
                assert patched[name] == data[name], name
        check_part(kind, data[main], patched[main], {
            p: [("客户", 0), ("B", 1), (" 保持原文", 2)],
            formatted: [("H", 0), ("3", 1), ("O 与 x", 2), ("3", 3), ("；", 4), ("蓝字", 5), ("大号", 6), ("Arial", 7), ("English", 8), ("中文", 9)],
            linked: [("访问", 0), ("官方首页面", 1), ("了解", 3)],
        })
        expected_other = [("页眉：客户", 0), ("B", 1)] if kind == "docx" else [("备注：客户B", 0)]
        check_part(kind, data[other_part], patched[other_part], {other: expected_other})
        # Only no-op edits: every part keeps its bytes.
        noop = directory / f"patch-noop-output.{suffix}"
        result = ops.patch({**base, "output": str(noop), "edits": [{"part": main, "paragraph": p, "before": "客户A", "after": "客户A"}]})
        assert result["changedParts"] == [] and any("第 1 项编辑的 before 与 after 相同" in w for w in result["warnings"]), result
        assert contents(noop) == data
        # A no-op next to a real edit is reported; the real edit still applies.
        mixed = directory / f"patch-mixed-output.{suffix}"
        result = ops.patch({**base, "output": str(mixed), "edits": [edits[0], {"part": main, "paragraph": linked, "before": "了解", "after": "了解"}]})
        assert result["changedParts"] == [main] and any("第 2 项" in w for w in result["warnings"]), result
        # A failing edit after a valid one publishes nothing.
        for failing in ({"part": main, "paragraph": p, "before": "客户B", "after": "甲方"},
                        {"part": main, "paragraph": 999, "before": "客户", "after": "甲方"}):
            failed = directory / f"patch-failed-output.{suffix}"
            try:
                ops.patch({**base, "output": str(failed), "edits": [edits[0], failing]})
                raise AssertionError("failing edit must stop the patch")
            except ValueError:
                assert not failed.exists(), "a failed patch left an output file"
        results[kind] = {"source": str(source), "output": str(output)}
    return results


def main(directory):
    count = 0
    for kind, extra in (("word", WORD_ONLY), ("ppt", PPT_ONLY)):
        for name, items, edits, expected in SHARED + extra:
            check_paragraph(kind, name, items, edits, expected)
            count += 1
    check_nested_and_whitespace()
    outputs = check_packages(Path(directory))
    print(json.dumps({"paragraphCases": count, "packages": outputs}, ensure_ascii=False))
    print("DOCUMENT_PATCH_CASES_PASSED")


if __name__ == "__main__":
    main(sys.argv[1])
